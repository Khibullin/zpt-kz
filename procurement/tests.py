from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone

from .analysis import evaluate
from .importing import import_offers, import_demand
from .models import Membership, Supplier, Offer, Preferences, Selection, Activity, Favorite


class WorkspaceTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user('buyer', password='test-password-123')
        cls.other = get_user_model().objects.create_user('other', password='test-password-123')
        cls.sales = get_user_model().objects.create_user('manager', password='test-password-123')
        cls.admin = get_user_model().objects.create_user('workspace-admin', password='test-password-123')
        cls.outsider = get_user_model().objects.create_user('outsider', password='test-password-123')
        for user, role in [(cls.user, 'user'), (cls.other, 'user'), (cls.sales, 'sales'), (cls.admin, 'admin')]:
            Membership.objects.create(user=user, role=role)
        cls.supplier = Supplier.objects.create(name='Omega')
        cls.offer = Offer.objects.create(supplier=cls.supplier, article='00123', brand='EEP', name='Деталь', cost=1000,
            kaspi_price=3000, checked_on=timezone.localdate(), evidence_date=timezone.localdate(),
            recent_reviews=1, stock='>20 шт.', match_verified=True, kaspi_url='https://kaspi.kz/shop/p/test-123/')

    def url(self, name, **kwargs):
        return reverse('procurement:' + name, kwargs=kwargs or None)

    def login(self, user=None):
        self.client.force_login(user or self.user)

    def test_anonymous_redirects_and_login_page_renders(self):
        self.assertEqual(self.client.get(self.url('dashboard')).status_code, 302)
        self.assertContains(self.client.get(self.url('login')), 'Вход в кабинет закупок')

    def test_regular_site_account_is_not_a_workspace_member(self):
        self.login(self.outsider)
        self.assertEqual(self.client.get(self.url('dashboard')).status_code, 403)

    def test_inactive_member_denied(self):
        Membership.objects.filter(user=self.user).update(active=False)
        self.login()
        self.assertEqual(self.client.get(self.url('dashboard')).status_code, 403)

    def test_dashboard_detail_and_forms_render(self):
        self.login()
        self.assertContains(self.client.get(self.url('dashboard')), '00123')
        self.assertContains(self.client.get(self.url('detail', pk=self.offer.pk)), 'Сравнение поставщиков')
        self.assertEqual(self.client.get(self.url('settings')).status_code, 200)
        self.assertEqual(self.client.get(self.url('basket')).status_code, 200)

    def test_user_cannot_upload_verify_or_edit_drafts(self):
        self.login()
        for name, kwargs in [('upload', {}), ('users', {}), ('draft', {'pk': self.offer.pk}), ('verify', {'pk': self.offer.pk})]:
            self.assertEqual(self.client.post(self.url(name, **kwargs), {}).status_code, 403)

    def test_manager_can_edit_draft_but_cannot_grant_roles(self):
        self.login(self.sales)
        self.assertEqual(self.client.get(self.url('draft', pk=self.offer.pk)).status_code, 200)
        self.assertEqual(self.client.get(self.url('users')).status_code, 403)
        self.assertEqual(self.client.get(self.url('upload')).status_code, 200)

    def test_administrator_role_assignment_and_audit(self):
        self.login(self.admin)
        self.assertEqual(self.client.get(self.url('users')).status_code, 200)
        self.client.post(self.url('users'), {'user_id': self.outsider.pk, 'role': 'sales'})
        self.assertEqual(Membership.objects.get(user=self.outsider).role, 'sales')
        self.outsider.refresh_from_db()
        self.assertFalse(self.outsider.is_staff)
        self.assertTrue(Activity.objects.filter(action='role_changed', user=self.admin).exists())
        self.assertEqual(self.client.post(self.url('users'), {'user_id': self.admin.pk, 'role': 'user'}).status_code, 400)

    def test_basket_is_private_and_post_only(self):
        self.login()
        self.assertEqual(self.client.get(self.url('select', pk=self.offer.pk)).status_code, 405)
        self.client.post(self.url('select', pk=self.offer.pk), {'quantity': 2})
        selected = Selection.objects.get(user=self.user)
        self.login(self.other)
        self.assertNotContains(self.client.get(self.url('basket')), '00123')
        self.client.post(self.url('remove', pk=selected.pk))
        self.assertTrue(Selection.objects.filter(pk=selected.pk).exists())
        self.assertNotContains(self.client.get(self.url('export')), '00123')

    def test_invalid_quantity_rejected(self):
        self.login()
        for value in ['0', '-1', '1.5', 'abc', '100001']:
            self.assertEqual(self.client.post(self.url('select', pk=self.offer.pk), {'quantity': value}).status_code, 400)

    def test_csrf_protected(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(client.post(self.url('select', pk=self.offer.pk), {'quantity': 1}).status_code, 403)

    def test_unknown_expenses_are_not_zero(self):
        preferences = Preferences(user=self.user)
        result = evaluate(self.offer, preferences)
        self.assertEqual(result['difference'], 2000)
        self.assertIsNone(result['profit'])
        self.assertFalse(result['qualified'])

    def test_financial_calculation_and_qualification(self):
        preferences = Preferences(user=self.user, commission=Decimal(10), expenses=Decimal(100), min_profit=Decimal(1000))
        result = evaluate(self.offer, preferences)
        self.assertEqual(result['profit'], Decimal(1600))
        self.assertEqual(result['markup'], Decimal(200))
        self.assertTrue(result['qualified'])
        self.assertEqual(result['signal'], 'Слабый свежий сигнал')

    def test_stale_and_future_evidence_excluded(self):
        preferences = Preferences(user=self.user, commission=0, expenses=0)
        for date in [timezone.localdate()-timedelta(days=8), timezone.localdate()+timedelta(days=1)]:
            self.offer.evidence_date = date
            self.assertFalse(evaluate(self.offer, preferences)['qualified'])

    def test_no_sales_claim_from_reviews(self):
        self.offer.review_count = 20
        self.offer.recent_reviews = None
        result = evaluate(self.offer, Preferences(user=self.user, commission=0, expenses=0))
        self.assertFalse(result['qualified'])
        self.assertEqual(result['signal'], 'Данных недостаточно')

    def test_unverified_match_excluded(self):
        self.offer.match_verified = False
        self.assertFalse(evaluate(self.offer, Preferences(user=self.user, commission=0, expenses=0))['qualified'])

    def test_zero_kaspi_and_cost_do_not_divide_by_zero(self):
        self.offer.kaspi_price = 0
        self.assertIsNone(evaluate(self.offer, Preferences(user=self.user, commission=0, expenses=0))['margin'])
        self.offer.cost = 0
        self.assertIsNone(evaluate(self.offer, Preferences(user=self.user))['markup'])

    def csv(self, content):
        return SimpleUploadedFile('prices.csv', content.encode('utf-8-sig'))

    def test_csv_import_leading_zeros_and_unknown_prices(self):
        count = import_offers(self.csv('article;brand;name;cost;kaspi_price\n00001;TEST;Деталь;1200;Не найдено\n'), 'New supplier', timezone.localdate(), self.sales)
        self.assertEqual(count, 1)
        offer = Offer.objects.get(article='00001')
        self.assertIsNone(offer.kaspi_price)
        self.assertFalse(offer.match_verified)
        self.assertIsNone(offer.sales_30)

    def test_invalid_row_rolls_back_entire_import(self):
        with self.assertRaises(ValueError):
            import_offers(self.csv('article,brand,name,cost\nA,TEST,Good,100\nB,TEST,Bad,-1\n'), 'Atomic', timezone.localdate(), self.sales)
        self.assertFalse(Supplier.objects.filter(name='Atomic').exists())

    def test_duplicate_rows_are_not_silently_overwritten(self):
        with self.assertRaises(ValueError):
            import_offers(self.csv('article,brand,name,cost\nA,TEST,Good,100\nA,TEST,Again,200\n'), 'Duplicate', timezone.localdate(), self.sales)

    def test_new_card_invalidates_old_match_and_demand(self):
        import_offers(self.csv('article,brand,name,cost,kaspi_url\n00123,EEP,Деталь,1000,https://kaspi.kz/shop/p/other-123/\n'), 'Omega', timezone.localdate(), self.sales)
        self.offer.refresh_from_db()
        self.assertFalse(self.offer.match_verified)
        self.assertIsNone(self.offer.recent_reviews)

    def test_formula_injection_is_escaped_in_export(self):
        self.offer.name = '=HYPERLINK("x")'
        self.offer.save()
        Selection.objects.create(user=self.user, offer=self.offer)
        self.login()
        response = self.client.get(self.url('export'))
        self.assertIn("'=HYPERLINK", response.content.decode('utf-8-sig'))

    def test_unknown_or_zero_stock_not_recommended(self):
        for stock in ['', '0 шт.', 'Нет данных']:
            self.offer.stock = stock
            self.assertFalse(evaluate(self.offer, Preferences(user=self.user, commission=0, expenses=0))['qualified'])

    def test_demand_import_requires_explicit_sales_window(self):
        with self.assertRaises(ValueError):
            import_demand(self.csv('article,brand,sales_30\n00123,EEP,10\n'), 'Omega', timezone.localdate(), self.sales)
        self.offer.refresh_from_db()
        self.assertIsNone(self.offer.sales_30)

    def test_demand_import_zero_vs_unknown_and_resets_old_values(self):
        date = timezone.localdate()
        start = (date - timedelta(days=29)).isoformat()
        import_demand(self.csv(f'article,brand,sales_30,period_start,period_end\n00123,EEP,0,{start},{date.isoformat()}\n'), 'Omega', date, self.sales)
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.sales_30, 0)
        self.assertIsNone(self.offer.recent_reviews)
        self.assertTrue(self.offer.evidence_source)
        import_demand(self.csv('article,brand,review_count,recent_reviews\n00123,EEP,4,1\n'), 'Omega', date, self.sales)
        self.offer.refresh_from_db()
        self.assertIsNone(self.offer.sales_30)
        self.assertEqual(self.offer.recent_reviews, 1)

    def test_xlsx_import_original_omega_headers(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Оптовой каталог'
        sheet.append(['Omega price list'])
        for _ in range(6):
            sheet.append([])
        sheet.append(['№', 'Артикул', 'Бренд', 'Наименование', 'Категория', 'Оптовая цена, ₸', 'Остаток на сайте', 'Мин. заказ, шт.'])
        sheet.append([1, '00099', 'TEST', 'Деталь', 'Двигатель', 450, '>10 шт.', 1])
        file = BytesIO()
        workbook.save(file)
        upload = SimpleUploadedFile('omega.xlsx', file.getvalue())
        self.assertEqual(import_offers(upload, 'Excel supplier', timezone.localdate(), self.sales), 1)
        self.assertEqual(Offer.objects.get(article='00099').cost, 450)

    def test_business_admin_can_revoke_workspace_access_only(self):
        self.login(self.admin)
        self.client.post(self.url('users'), {'user_id': self.user.pk, 'role': 'disabled'})
        self.user.refresh_from_db()
        self.assertTrue(self.user.is_active)
        self.assertFalse(Membership.objects.get(user=self.user).active)

    def test_best_mode_recalculates_after_expense_change(self):
        self.login()
        Preferences.objects.create(user=self.user, commission=10, expenses=100)
        self.assertContains(self.client.get(self.url('dashboard')+'?mode=best'), '00123')
        self.client.post(self.url('settings'), {'commission': 10, 'expenses': 2000, 'min_profit': 1000, 'min_markup': 30})
        self.assertNotContains(self.client.get(self.url('dashboard')+'?mode=best'), '00123')

    def test_photo_import_rejects_script_url(self):
        with self.assertRaises(ValueError):
            import_offers(self.csv('article,brand,name,cost,photo_url\nA,TEST,Good,100,javascript:alert(1)\n'), 'Photo', timezone.localdate(), self.sales)
        self.assertFalse(Supplier.objects.filter(name='Photo').exists())

    def test_category_stock_and_cost_filters(self):
        self.login()
        self.offer.category = 'Двигатель'
        self.offer.save()
        self.assertContains(self.client.get(self.url('dashboard'), {'category': 'Двигатель', 'availability': 'in_stock', 'max_cost': 1000}), '00123')
        self.assertEqual(self.client.get(self.url('dashboard'), {'max_cost': 999}).context['total'], 0)
        self.assertEqual(self.client.get(self.url('dashboard'), {'availability': 'unknown'}).context['total'], 0)

    def test_unknown_profit_never_matches_minimum(self):
        self.login()
        response = self.client.get(self.url('dashboard'), {'min_profit': 0})
        self.assertEqual(response.context['total'], 0)
        Preferences.objects.filter(user=self.user).update(commission=10, expenses=100)
        self.assertEqual(self.client.get(self.url('dashboard'), {'min_profit': 1600}).context['total'], 1)
        self.assertEqual(self.client.get(self.url('dashboard'), {'min_profit': 1601}).context['total'], 0)

    def test_demand_filter_separates_reviews_and_sales(self):
        self.login()
        self.assertEqual(self.client.get(self.url('dashboard'), {'demand': 'reviews'}).context['total'], 1)
        self.assertEqual(self.client.get(self.url('dashboard'), {'demand': 'sales'}).context['total'], 0)
        self.offer.evidence_date = timezone.localdate()-timedelta(days=8)
        self.offer.save()
        self.assertEqual(self.client.get(self.url('dashboard'), {'demand': 'reviews'}).context['total'], 0)
        self.assertEqual(self.client.get(self.url('dashboard'), {'demand': 'unknown'}).context['total'], 1)

    def test_favorites_are_private_and_idempotent(self):
        self.login()
        url = self.url('favorite', pk=self.offer.pk)
        self.assertEqual(self.client.get(url).status_code, 405)
        for _ in range(2):
            self.client.post(url, {'action': 'add'})
        self.assertEqual(Favorite.objects.filter(user=self.user).count(), 1)
        self.assertEqual(self.client.get(self.url('dashboard'), {'favorites': 1}).context['total'], 1)
        self.login(self.other)
        self.assertEqual(self.client.get(self.url('dashboard'), {'favorites': 1}).context['total'], 0)
        self.client.post(url, {'action': 'remove'})
        self.assertTrue(Favorite.objects.filter(user=self.user).exists())

    def test_favorite_redirect_does_not_leave_site(self):
        self.login()
        response = self.client.post(self.url('favorite', pk=self.offer.pk), {'action': 'add', 'next': 'https://evil.example/'})
        self.assertEqual(response.url, self.url('dashboard'))

    def test_current_selection_export_uses_same_filters(self):
        self.login()
        Favorite.objects.create(user=self.other, offer=self.offer)
        response = self.client.get(self.url('dashboard'), {'favorites': 1, 'download': 'csv'})
        self.assertNotIn('00123', response.content.decode('utf-8-sig'))
        response = self.client.get(self.url('dashboard'), {'demand': 'reviews', 'download': 'csv'})
        self.assertIn('00123', response.content.decode('utf-8-sig'))
        self.assertIn('attachment;', response['Content-Disposition'])

    def test_invalid_filter_visible_not_silently_ignored(self):
        self.login()
        response = self.client.get(self.url('dashboard'), {'min_profit': 'nan'})
        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.context['filter_form'].errors)

    def test_mode_links_keep_active_filters(self):
        self.login()
        response = self.client.get(self.url('dashboard'), {'availability': 'in_stock', 'q': '00123'})
        self.assertContains(response, 'availability=in_stock')
        self.assertContains(response, 'q=00123')
        self.assertContains(response, 'product-drawer')

    def test_quick_view_allows_only_same_origin_and_still_requires_role(self):
        self.login()
        response = self.client.get(self.url('detail', pk=self.offer.pk), {'compact': 1})
        self.assertEqual(response['X-Frame-Options'], 'SAMEORIGIN')
        self.assertContains(response, '<base target="_top">', html=True)
        self.assertNotContains(response, '<aside>')
        self.assertContains(response, 'margin-left:0;padding:16px')
        self.login(self.outsider)
        self.assertEqual(self.client.get(self.url('detail', pk=self.offer.pk), {'compact': 1}).status_code, 403)

    def test_account_entry_visible_only_to_active_workspace_members(self):
        account = reverse('interface_section', kwargs={'section': 'account'})
        self.assertNotContains(self.client.get(account), 'Кабинет закупок')
        self.login(self.outsider)
        self.assertNotContains(self.client.get(account), 'Кабинет закупок')
        self.login()
        response = self.client.get(account)
        self.assertContains(response, 'Кабинет закупок')
        self.assertContains(response, self.url('dashboard'))
        self.assertIn('no-store', response['Cache-Control'])
        self.assertIn('Cookie', response['Vary'])
        Membership.objects.filter(user=self.user).update(active=False)
        self.assertNotContains(self.client.get(account), 'Кабинет закупок')

    def test_anonymous_deep_link_returns_after_successful_login(self):
        target = self.url('detail', pk=self.offer.pk)
        response = self.client.get(target)
        self.assertRedirects(response, self.url('login')+'?next='+target)
        response = self.client.post(self.url('login'), {'username': 'buyer', 'password': 'test-password-123', 'next': target})
        self.assertRedirects(response, target)

    def test_successful_login_defaults_to_workspace(self):
        response = self.client.post(self.url('login'), {'username': 'buyer', 'password': 'test-password-123'})
        self.assertRedirects(response, self.url('dashboard'))

    def test_login_rejects_accounts_without_assigned_role(self):
        response = self.client.post(self.url('login'), {'username': 'outsider', 'password': 'test-password-123'})
        self.assertContains(response, 'Доступ к кабинету закупок не выдан')
        self.assertNotIn('_auth_user_id', self.client.session)
        self.assertFalse(Membership.objects.filter(user=self.outsider).exists())

    def test_already_signed_in_member_bypasses_login_form(self):
        self.login()
        self.assertRedirects(self.client.get(self.url('login')), self.url('dashboard'))
        self.login(self.outsider)
        response = self.client.get(self.url('login'))
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, 'Доступ к кабинету закупок не выдан', status_code=403)
        self.assertNotContains(response, '<aside>', status_code=403)

    def test_login_does_not_redirect_to_external_next(self):
        response = self.client.post(self.url('login'), {'username': 'buyer', 'password': 'test-password-123', 'next': 'https://evil.example/'})
        self.assertRedirects(response, self.url('dashboard'))

    def test_private_headers_on_pages_redirects_errors_and_exports(self):
        for user in [None, self.user, self.outsider]:
            self.client.logout()
            if user:
                self.login(user)
            for url in [self.url('dashboard'), self.url('export'), self.url('login'), self.url('detail', pk=999999)]:
                response = self.client.get(url)
                self.assertEqual(response['X-Robots-Tag'], 'noindex, nofollow, noarchive, nosnippet')
                self.assertIn('private', response['Cache-Control'])
                self.assertIn('no-store', response['Cache-Control'])
                self.assertIn('Cookie', response['Vary'])

    @patch.dict('os.environ', {'GOOGLE_TAG_MANAGER_ID': 'GTM-TEST123'})
    def test_external_analytics_not_injected_into_workspace(self):
        self.login()
        response = self.client.get(self.url('dashboard'))
        self.assertNotContains(response, 'googletagmanager.com')
        self.assertContains(response, 'noindex,nofollow,noarchive,nosnippet')

    def test_robots_excludes_workspace(self):
        self.assertContains(self.client.get('/robots.txt'), 'Disallow: /office/procurement/')

    def test_omega_categories_include_empty_sectors(self):
        self.login()
        response = self.client.get(self.url('dashboard'), {'category': 'Электрооборудование'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['total'], 0)
        self.assertContains(response, 'Тех обслуживание и сервис')

    def test_vehicle_pairs_do_not_cross_match_or_include_unverified(self):
        from .models import VehicleFitment
        for make, model, verified in [('Toyota', 'Camry', True), ('Lexus', 'RX', True), ('Toyota', 'RX', False)]:
            VehicleFitment.objects.create(offer=self.offer, make=make, model=model, verified=verified, source_url='https://example.com/fitment')
        self.login()
        response = self.client.get(self.url('dashboard'), {'vehicle_make': 'Toyota', 'vehicle_model': 'Camry'})
        self.assertEqual(response.context['total'], 1)
        self.assertContains(response, 'Toyota Camry')
        response = self.client.get(self.url('dashboard'), {'vehicle_make': 'Toyota', 'vehicle_model': 'RX'})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.context['total'], 0)
        self.assertEqual(self.client.get(self.url('dashboard'), {'fitment': 'unknown'}).context['total'], 0)
        self.assertEqual(self.client.get(self.url('dashboard'), {'fitment': 'verified'}).context['total'], 1)
        self.assertEqual(self.client.get(self.url('dashboard')).context['total'], 1)

    def test_text_applicability_does_not_create_verified_vehicle_match(self):
        self.offer.applicability = 'Toyota Camry'
        self.offer.save()
        self.login()
        self.assertEqual(self.client.get(self.url('dashboard'), {'fitment': 'verified'}).context['total'], 0)
        self.assertEqual(self.client.get(self.url('dashboard'), {'fitment': 'unknown'}).context['total'], 1)

    def test_pagination_last_and_jump_keep_filters(self):
        Offer.objects.bulk_create([Offer(supplier=self.supplier, brand='EEP', article=f'P{i}', name='Деталь', cost=100, category='Двигатель') for i in range(90)])
        self.login()
        response = self.client.get(self.url('dashboard'), {'category': 'Двигатель', 'page': 2})
        self.assertEqual(response.context['page'].number, 2)
        self.assertContains(response, 'Последняя →')
        self.assertContains(response, 'page=3')
        self.assertContains(response, 'name="category" value="Двигатель"')
        self.assertContains(response, 'aria-current="page">2')
        last = self.client.get(self.url('dashboard'), {'category': 'Двигатель', 'page': 3})
        self.assertEqual(len(last.context['page']), 10)
        self.assertContains(last, 'Первая')

    def enrichment_file(self, rows=None):
        import json
        return SimpleUploadedFile('fitments.json', json.dumps({'checked_on': timezone.localdate().isoformat(), 'offers': rows if rows is not None else [{'brand': self.offer.brand, 'article': self.offer.article, 'source_url': self.offer.supplier_url, 'category': 'Трансмиссия', 'fitments': [{'make': 'TOYOTA', 'model': 'CAMRY', 'specifications': {'Шасси': 'ASV50L', 'Годы': '2011–2017', 'Двигатель': '2AR-FE', 'Поколение (Omega)': 'v50-2011-2017'}}]}]}).encode())

    def test_enrichment_preserves_prices_demand_and_selection_and_is_idempotent(self):
        from .importing import import_enrichment
        from .models import VehicleFitment
        self.offer.supplier_url = 'https://omega-auto-parts.kz/transmissiya/test-part'
        self.offer.save()
        Selection.objects.create(user=self.user, offer=self.offer, quantity=3)
        original = (self.offer.cost, self.offer.kaspi_price, self.offer.checked_on, self.offer.evidence_date, self.offer.recent_reviews, self.offer.match_verified)
        for _ in range(2):
            self.assertEqual(import_enrichment(self.enrichment_file(), 'Omega', timezone.localdate(), self.sales), 1)
        self.offer.refresh_from_db()
        self.assertEqual(original, (self.offer.cost, self.offer.kaspi_price, self.offer.checked_on, self.offer.evidence_date, self.offer.recent_reviews, self.offer.match_verified))
        self.assertEqual(Selection.objects.get(user=self.user).quantity, 3)
        self.assertEqual(VehicleFitment.objects.filter(offer=self.offer).count(), 1)
        fit = self.offer.fitments.get()
        self.assertEqual(fit.configurations[0]['engine'], '2AR-FE')
        self.login()
        response = self.client.get(self.url('dashboard'), {'vehicle_make': 'TOYOTA', 'vehicle_model': 'CAMRY', 'category': 'Трансмиссия'})
        self.assertEqual(response.context['total'], 1)
        self.assertContains(self.client.get(self.url('detail', pk=self.offer.pk)), 'ASV50L')

    def test_enrichment_wrong_source_rolls_back_all_changes(self):
        from .importing import import_enrichment
        self.offer.supplier_url = 'https://omega-auto-parts.kz/transmissiya/test-part'
        self.offer.save()
        rows = [{'brand': self.offer.brand, 'article': self.offer.article, 'source_url': 'https://other.example/part', 'category': 'Двигатель', 'fitments': []}]
        with self.assertRaises(ValueError):
            import_enrichment(self.enrichment_file(rows), 'Omega', timezone.localdate(), self.sales)
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.category, '')
        self.assertFalse(Activity.objects.filter(action='fitment_import').exists())


    def test_https_forms_keep_csrf_and_do_not_send_external_referrers(self):
        import re
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        response = client.get(self.url('settings'), secure=True)
        self.assertEqual(response['Referrer-Policy'], 'same-origin')
        self.assertContains(response, 'name="referrer" content="same-origin"')
        token = re.search(r'name="csrfmiddlewaretoken" value="([^"]+)"', response.content.decode()).group(1)
        data = {'csrfmiddlewaretoken': token, 'commission': '10', 'expenses': '100', 'min_profit': '1000', 'min_markup': '30'}
        self.assertEqual(client.post(self.url('settings'), data, secure=True).status_code, 403)
        response = client.post(self.url('settings'), data, secure=True, HTTP_REFERER='https://testserver' + self.url('settings'))
        self.assertEqual(response.status_code, 302)

    def prepare_purchase(self, quantity=2):
        self.login()
        self.offer.supplier_url = 'https://omega-auto-parts.kz/transmissiya/test-part'
        self.offer.save()
        Selection.objects.create(user=self.user, offer=self.offer, quantity=quantity)
        response = self.client.get(self.url('purchase_confirm', supplier_id=self.supplier.pk))
        form = response.context['form']
        return {'supplier_order_number': '', 'comment': '', 'request_key': str(form['request_key'].value()), 'selection_token': form['selection_token'].value()}

    def test_basket_creates_internal_orders_without_external_actions(self):
        self.prepare_purchase()
        response = self.client.get(self.url('basket'))
        self.assertContains(response, 'Создать закупочный заказ')
        self.assertContains(response, 'Закупочные заказы ZPT.KZ')
        for text in ('Перейти к заказу у Omega', 'Скопировать заказ', 'Отметить как заказанное', 'Скачать CSV'):
            self.assertNotContains(response, text)

    def test_purchase_record_is_private_snapshot_and_repeat_is_idempotent(self):
        from .models import PurchaseOrder
        data = self.prepare_purchase()
        Selection.objects.create(user=self.other, offer=self.offer, quantity=5)
        response = self.client.post(self.url('purchase_confirm', supplier_id=self.supplier.pk), data)
        order = PurchaseOrder.objects.get(user=self.user)
        self.assertRedirects(response, self.url('purchase_detail', pk=order.pk))
        self.assertEqual(order.estimated_total, Decimal(2000))
        self.assertEqual(order.items[0]['quantity'], 2)
        self.assertFalse(Selection.objects.filter(user=self.user).exists())
        self.assertTrue(Selection.objects.filter(user=self.other).exists())
        self.assertEqual(self.client.post(self.url('purchase_confirm', supplier_id=self.supplier.pk), data).status_code, 302)
        self.assertEqual(PurchaseOrder.objects.count(), 1)
        self.offer.cost = 9999
        self.offer.name = 'Новое имя'
        self.offer.save()
        self.assertContains(self.client.get(self.url('purchase_detail', pk=order.pk)), 'Деталь')
        order.refresh_from_db()
        self.assertEqual(order.items[0]['cost'], '1000.00')
        self.login(self.other)
        self.assertEqual(self.client.get(self.url('purchase_detail', pk=order.pk)).status_code, 404)
        self.assertEqual(self.client.post(self.url('purchase_detail', pk=order.pk), {'status':'received'}).status_code, 404)
        self.assertNotContains(self.client.get(self.url('basket')), 'OMEGA-123')

    def test_changed_or_tampered_basket_cannot_be_recorded(self):
        from .models import PurchaseOrder
        data = self.prepare_purchase()
        Selection.objects.filter(user=self.user).update(quantity=3)
        response = self.client.post(self.url('purchase_confirm', supplier_id=self.supplier.pk), data)
        self.assertEqual(response.status_code, 400)
        self.assertFalse(PurchaseOrder.objects.exists())
        self.assertEqual(Selection.objects.get(user=self.user).quantity, 3)
        data['selection_token'] = 'tampered'
        self.assertEqual(self.client.post(self.url('purchase_confirm', supplier_id=self.supplier.pk), data).status_code, 400)

    def edit_purchase_data(self, order, **overrides):
        data = {field: getattr(order, field) for field in ('status', 'supplier_order_number', 'invoice_number', 'waybill_number', 'comment', 'version')}
        data.update({'lines-TOTAL_FORMS': str(len(order.items)), 'lines-INITIAL_FORMS': str(len(order.items)), 'lines-MIN_NUM_FORMS': '0', 'lines-MAX_NUM_FORMS': '20000'})
        for index, item in enumerate(order.items):
            data[f'lines-{index}-quantity'] = item['quantity']
            data[f'lines-{index}-cost'] = item['cost']
        data.update(overrides)
        return data

    def create_internal_purchase(self):
        from .models import PurchaseOrder
        data = self.prepare_purchase()
        self.client.post(self.url('purchase_confirm', supplier_id=self.supplier.pk), data)
        return PurchaseOrder.objects.get(user=self.user)

    def test_internal_number_draft_and_optional_supplier_metadata(self):
        order = self.create_internal_purchase()
        self.assertEqual(order.order_number, f'ЗК-{order.pk:06d}')
        self.assertEqual(order.status, 'draft')
        self.assertEqual(order.supplier_order_number, '')
        self.assertEqual(order.original_items, order.items)
        self.assertContains(self.client.get(self.url('purchase_detail', pk=order.pk)), 'Выгрузить заказ в Excel')
        response = self.client.post(self.url('purchase_detail', pk=order.pk), self.edit_purchase_data(order, status='confirmed', supplier_order_number='OMG-55', invoice_number='INV-17', comment='Подтвердили наличие'))
        self.assertEqual(response.status_code, 302)
        order.refresh_from_db()
        self.assertEqual(order.status, 'confirmed')
        self.assertEqual(order.supplier_order_number, 'OMG-55')
        self.assertEqual(order.invoice_number, 'INV-17')
        self.assertTrue(Activity.objects.filter(action='purchase_updated').exists())

    def test_quantity_price_edits_keep_originals_and_reject_stale_version(self):
        order = self.create_internal_purchase()
        old = self.edit_purchase_data(order, **{'lines-0-quantity': '3', 'lines-0-cost': '1500.00'})
        self.assertEqual(self.client.post(self.url('purchase_detail', pk=order.pk), old).status_code, 302)
        order.refresh_from_db()
        self.assertEqual(order.total, Decimal(4500))
        self.assertEqual(order.original_items[0]['quantity'], 2)
        self.assertEqual(order.original_items[0]['cost'], '1000.00')
        self.assertEqual(order.estimated_total, Decimal(2000))
        self.assertEqual(self.client.post(self.url('purchase_detail', pk=order.pk), old).status_code, 400)
        invalid = self.edit_purchase_data(order, **{'lines-0-cost': '-1'})
        self.assertEqual(self.client.post(self.url('purchase_detail', pk=order.pk), invalid).status_code, 400)
        order.refresh_from_db()
        self.assertEqual(order.total, Decimal(4500))

    def test_excel_is_private_uses_current_values_and_preserves_articles(self):
        from openpyxl import load_workbook
        order = self.create_internal_purchase()
        order.items[0]['name'] = '=HYPERLINK("https://example.com")'
        order.save(update_fields=['items'])
        response = self.client.get(self.url('purchase_excel', pk=order.pk))
        self.assertEqual(response.status_code, 200)
        self.assertIn('.xlsx', response['Content-Disposition'])
        book = load_workbook(BytesIO(b''.join(response.streaming_content)), data_only=False)
        sheet = book['Заказ']
        self.assertEqual(sheet['B8'].value, '00123')
        self.assertEqual(sheet['B8'].data_type, 's')
        self.assertEqual(sheet['D8'].data_type, 's')
        self.assertEqual(sheet['E8'].value, 2)
        self.assertEqual(sheet['F8'].value, 1000)
        self.assertEqual(sheet['G8'].value, '=E8*F8')
        self.assertIn(order.order_number, sheet['A1'].value)
        self.login(self.other)
        self.assertEqual(self.client.get(self.url('purchase_excel', pk=order.pk)).status_code, 404)

    def test_documents_are_private_and_invalid_upload_rolls_back_edits(self):
        from .models import PurchaseDocument
        order = self.create_internal_purchase()
        data = self.edit_purchase_data(order, invoice_file=SimpleUploadedFile('invoice.pdf', b'%PDF-test'))
        self.assertEqual(self.client.post(self.url('purchase_detail', pk=order.pk), data).status_code, 302)
        doc = PurchaseDocument.objects.get(order=order)
        response = self.client.get(self.url('purchase_document', pk=order.pk, document_id=doc.pk))
        self.assertEqual(b''.join(response.streaming_content), b'%PDF-test')
        self.assertIn('attachment', response['Content-Disposition'])
        order.refresh_from_db()
        bad = self.edit_purchase_data(order, status='paid', invoice_file=SimpleUploadedFile('unsafe.html', b'<script>'))
        self.assertEqual(self.client.post(self.url('purchase_detail', pk=order.pk), bad).status_code, 400)
        order.refresh_from_db()
        self.assertEqual(order.status, 'draft')
        self.login(self.other)
        self.assertEqual(self.client.get(self.url('purchase_document', pk=order.pk, document_id=doc.pk)).status_code, 404)

    def competitor_file(self, sellers=None, source=None):
        import json
        sellers = sellers if sellers is not None else [
            {'merchant_id':'111','name':'A','url':'https://kaspi.kz/shop/m/111/products','price':'2000','rating':None,'reviews':None},
            {'merchant_id':'222','name':'B','url':'https://kaspi.kz/shop/m/222/products','price':'2000','rating':'4.8','reviews':15},
            {'merchant_id':'333','name':'C','url':'https://kaspi.kz/shop/m/333/products','price':'2000','rating':'4.8','reviews':25},
            {'merchant_id':'444','name':'D','url':'https://kaspi.kz/shop/m/444/products','price':'1900','rating':'4.1','reviews':3}]
        return SimpleUploadedFile('competitors.json', json.dumps({'checked_on':timezone.localdate().isoformat(),'offers':[{'brand':self.offer.brand,'article':self.offer.article,'source_url':source or self.offer.kaspi_url,'city':'750000000','complete':False,'sellers':sellers}]}).encode())

    def test_competitor_sorting_and_import_do_not_change_cost_or_demand(self):
        from .importing import import_competitors
        from .models import KaspiSellerSnapshot
        for _ in range(2):
            self.assertEqual(import_competitors(self.competitor_file(), 'Omega', timezone.localdate(), self.sales), 1)
        snapshot = KaspiSellerSnapshot.objects.get(offer=self.offer)
        self.assertEqual([s['name'] for s in snapshot.ranked_sellers], ['D','C','B','A'])
        self.assertEqual(KaspiSellerSnapshot.objects.count(), 1)
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.cost, 1000)
        self.assertEqual(self.offer.kaspi_price, 3000)
        self.assertEqual(self.offer.recent_reviews, 1)
        self.assertTrue(self.offer.match_verified)
        self.login()
        response = self.client.get(self.url('dashboard'))
        self.assertContains(response, 'Отзывы магазина: 25')
        self.assertContains(response, 'Проверена часть предложений')
        self.assertContains(self.client.get(self.url('detail', pk=self.offer.pk)), 'Рейтинг магазина')

    def test_competitor_wrong_card_and_invalid_rating_are_rejected(self):
        from .importing import import_competitors
        from .models import KaspiSellerSnapshot
        with self.assertRaises(ValueError):
            import_competitors(self.competitor_file(source='https://kaspi.kz/shop/p/other-999/'), 'Omega', timezone.localdate(), self.sales)
        invalid=[{'merchant_id':'111','name':'A','url':'https://kaspi.kz/shop/m/111/products','price':'2000','rating':'6','reviews':1}]
        with self.assertRaises(ValueError):
            import_competitors(self.competitor_file(sellers=invalid), 'Omega', timezone.localdate(), self.sales)
        self.assertFalse(KaspiSellerSnapshot.objects.exists())

    def test_unknown_competitors_are_distinct_from_confirmed_empty(self):
        from .models import KaspiSellerSnapshot
        self.login()
        self.assertContains(self.client.get(self.url('dashboard')), 'Продавцы Kaspi ещё не проверены')
        snapshot=KaspiSellerSnapshot.objects.create(offer=self.offer, source_url=self.offer.kaspi_url, checked_on=timezone.localdate(), sellers=[], complete=False)
        self.assertContains(self.client.get(self.url('dashboard')), 'Данные продавцов не получены')
        snapshot.complete=True
        snapshot.save()
        self.assertContains(self.client.get(self.url('dashboard')), 'Нет предложений на дату проверки')
