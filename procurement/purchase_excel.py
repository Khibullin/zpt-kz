"""Runtime XLSX download for an internal purchase order."""
from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


def export_order(order):
    workbook = Workbook()
    workbook.properties.creator = 'ZPT.KZ'
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    sheet = workbook.active
    sheet.title = 'Заказ'
    def literal(row, col, value):
        cell = sheet.cell(row, col, str(value))
        cell.data_type = 's'  # Preserve leading zeros and disallow formula injection.
        return cell
    for row, text in [(1, f'Закупочный заказ {order.order_number}'), (3, f'Поставщик: {order.supplier.name}'),
                      (4, f'Номер у поставщика: {order.supplier_order_number or "—"}'),
                      (5, f'Статус: {order.get_status_display()}'), (6, f'Дата: {order.created_at:%d.%m.%Y}')]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        literal(row, 1, text)
    sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='142E4D')
    sheet.row_dimensions[1].height = 32
    headers = ['№', 'Артикул', 'Бренд', 'Наименование', 'Количество, шт.', 'Цена, ₸', 'Сумма, ₸']
    for col, name in enumerate(headers, 1):
        cell = literal(7, col, name)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176B9A')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[7].height = 30
    for index, item in enumerate(order.items, 1):
        row = index + 7
        sheet.cell(row, 1, index)
        for col, field in [(2, 'article'), (3, 'brand'), (4, 'name')]:
            literal(row, col, item[field])
        sheet.cell(row, 5, item['quantity'])
        sheet.cell(row, 6, Decimal(str(item['cost'])))
        sheet.cell(row, 7, f'=E{row}*F{row}')
        sheet.cell(row, 4).alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[row].height = 34
        for col in (6, 7):
            sheet.cell(row, col).number_format = '#,##0.00'
    end = 7 + len(order.items)
    total = end + 2
    literal(total, 4, 'ИТОГО')
    sheet.cell(total, 5, f'=SUM(E8:E{end})' if order.items else 0)
    sheet.cell(total, 7, f'=SUM(G8:G{end})' if order.items else 0).number_format = '#,##0.00'
    for cell in sheet[total]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E8F2FA')
    for col, width in enumerate([6, 24, 20, 58, 20, 20, 22], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = 'E8'
    sheet.auto_filter.ref = f'A7:G{max(7, end)}'
    sheet.print_title_rows = '1:7'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f'A1:G{total}'
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
