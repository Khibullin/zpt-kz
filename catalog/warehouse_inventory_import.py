"""Physical warehouse inventory import (PP1/PP2).

Exact Product.article match only. Never creates Product, never writes PP2 when
importing PP1, never touches Kaspi listings, prices, or Product.stock_qty.

Dry-run is the default. Writes require explicit apply=True.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from django.db import transaction

from catalog.ag_parts_import import cell_text, normalize_header
from catalog.models import Product, ProductWarehouseStock, StockMovement, Warehouse
from catalog.stock_service import StockServiceError, apply_stock_movement, resolve_warehouse
from catalog.warehouse_stock_sync import parse_stock_quantity

SOURCE_INVENTORY = 'warehouse_inventory'

STATUS_CREATE_OPENING = 'CREATE_OPENING'
STATUS_UNMATCHED = 'UNMATCHED'
STATUS_BLANK_QTY = 'BLANK_QTY'
STATUS_BLANK_QTY_MATCHED = 'BLANK_QTY_MATCHED'
STATUS_ERROR = 'ERROR'
STATUS_CONFLICT_EXISTING_BALANCE = 'CONFLICT_EXISTING_BALANCE'
STATUS_ALREADY_IMPORTED = 'ALREADY_IMPORTED'
STATUS_DUPLICATE_ARTICLE = 'DUPLICATE_ARTICLE'
STATUS_INVALID_ROW = 'INVALID_ROW'
STATUS_AMBIGUOUS_ARTICLE = 'AMBIGUOUS_ARTICLE'

SAFE_APPLY_STATUSES = frozenset({STATUS_CREATE_OPENING})

DEFAULT_ARTICLE_COLUMN = 'Артикул'
DEFAULT_QUANTITY_COLUMN = 'Наличие на складе АГ'


@dataclass
class InventoryRow:
    row_number: int
    article: str = ''
    raw_qty: object = None
    qty: int | None = None
    product_id: int | None = None
    product_title: str = ''
    status: str = ''
    reason: str = ''


@dataclass
class InventoryImportResult:
    rows: list[InventoryRow]
    warehouse_code: str
    warehouse_name: str
    inventory_date: date
    reference: str
    source_path: str
    apply: bool
    sheet_name: str = ''
    headers: list = field(default_factory=list)
    summary: dict = field(default_factory=dict)


def inventory_reference(warehouse_code: str, inventory_date: date) -> str:
    return f'{warehouse_code}-{inventory_date.isoformat()}'


def inventory_note(inventory_date: date, source_path: str) -> str:
    filename = Path(source_path).name if source_path else ''
    as_of = inventory_date.isoformat()
    if filename:
        return f'Physical inventory as of {as_of}\n{filename}'
    return f'Physical inventory as of {as_of}'


def parse_inventory_date(value) -> date:
    text = str(value or '').strip()
    try:
        parsed = date.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f'invalid_inventory_date:{text}') from exc
    if parsed.year < 2000:
        raise ValueError(f'invalid_inventory_date:{text}')
    return parsed


def parse_inventory_quantity(raw):
    """Return (qty, error). Blank is not zero."""
    qty, error = parse_stock_quantity(raw)
    if error == 'empty':
        return None, 'blank'
    if error:
        return None, error
    return qty, ''


def _column_index(headers, name: str):
    wanted = normalize_header(name)
    if not wanted:
        return None
    for index, header in enumerate(headers):
        if normalize_header(header) == wanted:
            return index
    return None


def _cell_at(values, index):
    if index is None or index >= len(values):
        return None
    return values[index]


def _load_inventory_table(path: Path, *, article_column: str, quantity_column: str):
    suffix = path.suffix.lower()
    if suffix not in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        raise ValueError(f'unsupported_file_type:{suffix or "missing"}')

    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    try:
        fallback = None
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            headers, article_idx, qty_idx, data_rows = _extract_sheet_table(
                worksheet,
                article_column=article_column,
                quantity_column=quantity_column,
            )
            candidate = (sheet_name, headers, article_idx, qty_idx, data_rows)
            if fallback is None:
                fallback = candidate
            if article_idx is not None and qty_idx is not None:
                return candidate
    finally:
        workbook.close()
    if fallback is None:
        raise ValueError('workbook_has_no_sheets')
    return fallback


def _extract_sheet_table(worksheet, *, article_column: str, quantity_column: str):
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = None
    headers = []
    article_idx = None
    qty_idx = None
    scan_limit = min(len(rows), 30)
    for index in range(scan_limit):
        candidate = [cell_text(value) for value in (rows[index] or ())]
        found_article = _column_index(candidate, article_column)
        found_qty = _column_index(candidate, quantity_column)
        if found_article is not None and found_qty is not None:
            header_index = index
            headers = candidate
            article_idx = found_article
            qty_idx = found_qty
            break
        if header_index is None and any(candidate):
            headers = candidate
    data_rows = []
    if header_index is None:
        return headers, None, None, data_rows
    for offset, raw in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = list(raw or ())
        if not any(cell_text(value) for value in values):
            continue
        data_rows.append((offset, values))
    return headers, article_idx, qty_idx, data_rows


def _find_product(article: str):
    matches = list(Product.objects.filter(article=article).order_by('pk'))
    if not matches:
        return None, 'product_not_found'
    if len(matches) > 1:
        return None, 'ambiguous_article'
    return matches[0], ''


def _existing_inventory_movement(*, product, warehouse, reference: str):
    return StockMovement.objects.filter(
        product=product,
        warehouse=warehouse,
        source=SOURCE_INVENTORY,
        reference=reference,
    ).exists()


def classify_inventory_rows(
    rows: list[InventoryRow],
    *,
    warehouse: Warehouse,
    reference: str,
) -> list[InventoryRow]:
    article_counts: dict[str, int] = {}
    for row in rows:
        if row.article:
            article_counts[row.article] = article_counts.get(row.article, 0) + 1

    for row in rows:
        if not row.article:
            row.status = STATUS_INVALID_ROW
            row.reason = 'empty_article'
            continue

        qty, qty_error = parse_inventory_quantity(row.raw_qty)
        if qty_error != 'blank':
            row.qty = qty
        product, match_error = _find_product(row.article)
        if product is not None:
            row.product_id = product.pk
            row.product_title = product.title

        if article_counts.get(row.article, 0) > 1:
            row.status = STATUS_DUPLICATE_ARTICLE
            row.reason = 'duplicate_source_article'
            continue
        if match_error == 'ambiguous_article':
            row.status = STATUS_AMBIGUOUS_ARTICLE
            row.reason = match_error
            continue
        if product is None:
            row.status = STATUS_UNMATCHED
            row.reason = match_error
            continue
        if qty_error == 'blank':
            row.status = STATUS_BLANK_QTY_MATCHED
            row.reason = 'blank_qty'
            continue
        if qty_error:
            row.status = STATUS_ERROR
            row.reason = qty_error
            continue
        if _existing_inventory_movement(
            product=product,
            warehouse=warehouse,
            reference=reference,
        ):
            row.status = STATUS_ALREADY_IMPORTED
            row.reason = 'same_reference_already_imported'
            continue
        existing = ProductWarehouseStock.objects.filter(
            product=product,
            warehouse=warehouse,
        ).first()
        if existing is not None:
            row.status = STATUS_CONFLICT_EXISTING_BALANCE
            row.reason = 'existing_balance'
            continue
        row.status = STATUS_CREATE_OPENING
        row.reason = 'create_opening'
    return rows


def summarize_inventory_rows(rows: list[InventoryRow]) -> dict:
    summary = {
        'source_rows': len(rows),
        'known_qty': 0,
        'blank_qty': 0,
        'matched_products': 0,
        'unmatched_products': 0,
        'safe_write_candidates': 0,
        'candidate_qty': 0,
        'conflicts': 0,
        'already_imported': 0,
        'errors': 0,
        'duplicates': 0,
        'ambiguous': 0,
        'invalid_rows': 0,
    }
    for row in rows:
        qty, qty_error = parse_inventory_quantity(row.raw_qty)
        if qty_error == 'blank':
            summary['blank_qty'] += 1
        elif qty is not None:
            summary['known_qty'] += 1
        if row.product_id is not None and row.status not in {
            STATUS_DUPLICATE_ARTICLE,
            STATUS_AMBIGUOUS_ARTICLE,
        }:
            summary['matched_products'] += 1
        if row.status == STATUS_UNMATCHED:
            summary['unmatched_products'] += 1
        if row.status == STATUS_CREATE_OPENING:
            summary['safe_write_candidates'] += 1
            summary['candidate_qty'] += int(row.qty or 0)
        elif row.status == STATUS_CONFLICT_EXISTING_BALANCE:
            summary['conflicts'] += 1
        elif row.status == STATUS_ALREADY_IMPORTED:
            summary['already_imported'] += 1
        elif row.status == STATUS_ERROR:
            summary['errors'] += 1
        elif row.status == STATUS_DUPLICATE_ARTICLE:
            summary['duplicates'] += 1
        elif row.status == STATUS_AMBIGUOUS_ARTICLE:
            summary['ambiguous'] += 1
        elif row.status == STATUS_INVALID_ROW:
            summary['invalid_rows'] += 1
    return summary


def apply_inventory_rows(
    rows: list[InventoryRow],
    *,
    warehouse: Warehouse,
    reference: str,
    note: str,
):
    with transaction.atomic():
        for row in rows:
            if row.status not in SAFE_APPLY_STATUSES:
                continue
            product = Product.objects.get(pk=row.product_id)
            stock_qty_before = product.stock_qty
            price_before = product.price
            cost_before = product.cost_price
            status_before = product.status
            apply_stock_movement(
                product=product,
                warehouse=warehouse,
                movement_type=StockMovement.MovementType.OPENING,
                quantity_delta=int(row.qty),
                source=SOURCE_INVENTORY,
                reference=reference,
                note=note,
            )
            product.refresh_from_db(
                fields=['stock_qty', 'price', 'cost_price', 'status']
            )
            if (
                product.stock_qty != stock_qty_before
                or product.price != price_before
                or product.cost_price != cost_before
                or product.status != status_before
            ):
                raise ValueError('importer_changed_product_fields')


def import_warehouse_inventory(
    *,
    path,
    warehouse_code: str,
    inventory_date,
    apply: bool = False,
    article_column: str = DEFAULT_ARTICLE_COLUMN,
    quantity_column: str = DEFAULT_QUANTITY_COLUMN,
) -> InventoryImportResult:
    source = Path(path)
    if not source.exists():
        raise ValueError(f'file_not_found:{source}')
    if isinstance(inventory_date, datetime):
        inventory_date = inventory_date.date()
    if not isinstance(inventory_date, date):
        inventory_date = parse_inventory_date(inventory_date)

    try:
        warehouse = resolve_warehouse(warehouse_code)
    except StockServiceError as exc:
        raise ValueError(str(exc)) from exc

    sheet_name, headers, article_idx, qty_idx, data_rows = _load_inventory_table(
        source,
        article_column=article_column,
        quantity_column=quantity_column,
    )
    if article_idx is None:
        raise ValueError(
            f'article_column_not_found:{article_column}. headers={headers}'
        )
    if qty_idx is None:
        raise ValueError(
            f'quantity_column_not_found:{quantity_column}. headers={headers}'
        )

    reference = inventory_reference(warehouse.code, inventory_date)
    note = inventory_note(inventory_date, str(source))
    rows = []
    for row_number, values in data_rows:
        article = cell_text(_cell_at(values, article_idx))
        rows.append(
            InventoryRow(
                row_number=row_number,
                article=article,
                raw_qty=_cell_at(values, qty_idx),
            )
        )
    classify_inventory_rows(rows, warehouse=warehouse, reference=reference)
    if apply:
        apply_inventory_rows(
            rows,
            warehouse=warehouse,
            reference=reference,
            note=note,
        )

    return InventoryImportResult(
        rows=rows,
        warehouse_code=warehouse.code,
        warehouse_name=warehouse.name,
        inventory_date=inventory_date,
        reference=reference,
        source_path=str(source),
        apply=apply,
        sheet_name=sheet_name,
        headers=list(headers),
        summary=summarize_inventory_rows(rows),
    )
