"""PREVIEW_ONLY Kaspi stock update from PP2 Rapido.

Does not write Kaspi qty or price. Duplicate physical SKUs are held.
Kaspi merchant native upload (SKU/price/PP1-PP5) is not emitted as
upload-ready because duplicate listings cannot safely receive the full
PP2 quantity.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from catalog.data.rapido_pp2_2026_09_18 import (
    HOLD_DUPLICATE_ARTICLES,
    NO_LISTING_ARTICLES,
    RAPIDO_PP2_QTY,
)
from catalog.models import Product, ProductKaspiListing
from catalog.stock_service import get_stock_quantity
from catalog.warehouses import WAREHOUSE_CODE_PP2

STATUS_READY = 'READY'
STATUS_HOLD_DUPLICATE = 'HOLD_DUPLICATE'
STATUS_NO_LISTING = 'NO_LISTING'

PREVIEW_HEADERS = (
    'Listing ID',
    'SKU Kaspi',
    'merchant SKU / article',
    'current price',
    'last known Kaspi qty',
    'PP2 target',
    'delta',
    'status',
    'note',
)


@dataclass
class KaspiStockPreviewRow:
    listing_id: int | None = None
    sku_kaspi: str = ''
    article: str = ''
    current_price: int | None = None
    last_known_kaspi_qty: int | None = None
    pp2_target: int | None = None
    delta: int | None = None
    status: str = ''
    note: str = ''


@dataclass
class KaspiStockPreviewResult:
    rows: list[KaspiStockPreviewRow]
    upload_ready: bool = False
    ready_count: int = 0
    hold_listing_count: int = 0
    duplicate_groups: list[str] = field(default_factory=list)
    no_listing: list[str] = field(default_factory=list)


def _compact_key(value: str) -> str:
    return re.sub(r'[^A-Za-z0-9]', '', value or '').upper()


def _article_in_model(model: str, article: str) -> bool:
    if not model or not article:
        return False
    pattern = r'(?<![A-Za-z0-9])' + re.escape(article) + r'(?![A-Za-z0-9])'
    return re.search(pattern, model, re.I) is not None


def _row_matches_article(sku: str, model: str, article: str) -> bool:
    if not article:
        return False
    compact_article = _compact_key(article)
    if compact_article and compact_article == _compact_key(sku):
        return True
    if _article_in_model(model, article):
        return True
    compact_model = _compact_key(model)
    if compact_article and compact_article in compact_model:
        return True
    return False


def _status_for_article(
    article: str,
    listing_count: int,
    *,
    declared_no_listing_only: bool = False,
) -> str:
    if article in NO_LISTING_ARTICLES:
        return STATUS_NO_LISTING
    if article in HOLD_DUPLICATE_ARTICLES or listing_count >= 2:
        return STATUS_HOLD_DUPLICATE
    if listing_count == 0 and not declared_no_listing_only:
        return STATUS_NO_LISTING
    return STATUS_READY


def _delta(target, last_known):
    if target is None or last_known is None:
        return None
    return int(target) - int(last_known)


def build_preview_from_listings(
    *,
    articles: dict[str, int] | None = None,
    use_live_pp2: bool = False,
) -> KaspiStockPreviewResult:
    targets = dict(articles or RAPIDO_PP2_QTY)
    listings = list(
        ProductKaspiListing.objects.filter(is_active=True)
        .select_related('product')
        .order_by('pk')
    )
    by_article: dict[str, list[ProductKaspiListing]] = {}
    for listing in listings:
        article = (listing.product.article or '').strip()
        if article not in targets:
            continue
        by_article.setdefault(article, []).append(listing)

    rows: list[KaspiStockPreviewRow] = []
    hold_groups = []
    no_listing = []
    for article, qty in targets.items():
        article_listings = by_article.get(article, [])
        status = _status_for_article(article, len(article_listings))
        live_qty = qty
        if use_live_pp2:
            product = Product.objects.filter(article=article).order_by('pk').first()
            if product is not None:
                live_qty = get_stock_quantity(product, WAREHOUSE_CODE_PP2)
        if status == STATUS_NO_LISTING:
            no_listing.append(article)
            rows.append(
                KaspiStockPreviewRow(
                    article=article,
                    pp2_target=live_qty,
                    status=STATUS_NO_LISTING,
                    note='No active Kaspi listing',
                )
            )
            continue
        if status == STATUS_HOLD_DUPLICATE:
            hold_groups.append(article)
            for listing in article_listings:
                rows.append(
                    KaspiStockPreviewRow(
                        listing_id=listing.pk,
                        sku_kaspi=listing.master_sku,
                        article=article,
                        current_price=listing.last_known_our_price,
                        last_known_kaspi_qty=listing.last_known_kaspi_qty,
                        pp2_target=None,
                        delta=None,
                        status=STATUS_HOLD_DUPLICATE,
                        note=(
                            f'PP2 physical qty={live_qty}; do not copy into '
                            f'each of {len(article_listings)} listings'
                        ),
                    )
                )
            continue
        listing = article_listings[0]
        rows.append(
            KaspiStockPreviewRow(
                listing_id=listing.pk,
                sku_kaspi=listing.master_sku,
                article=article,
                current_price=listing.last_known_our_price,
                last_known_kaspi_qty=listing.last_known_kaspi_qty,
                pp2_target=live_qty,
                delta=_delta(live_qty, listing.last_known_kaspi_qty),
                status=STATUS_READY,
                note='Single active listing; target = PP2 only',
            )
        )
    return _finalize(rows, hold_groups, no_listing)


def build_preview_from_kaspi_export(
    kaspi_path,
    *,
    articles: dict[str, int] | None = None,
) -> KaspiStockPreviewResult:
    targets = dict(articles or RAPIDO_PP2_QTY)
    workbook = load_workbook(kaspi_path, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        export_rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not export_rows:
        raise ValueError('kaspi_export_empty')
    headers = [str(value or '').strip() for value in export_rows[0]]
    header_map = {name: index for index, name in enumerate(headers)}
    sku_idx = header_map.get('SKU')
    model_idx = header_map.get('model')
    price_idx = header_map.get('price')
    pp2_idx = header_map.get('PP2')
    if sku_idx is None or model_idx is None:
        raise ValueError(f'kaspi_export_missing_columns:{headers}')

    matched: dict[str, list[dict]] = {article: [] for article in targets}
    for raw in export_rows[1:]:
        values = list(raw or ())
        model = str(values[model_idx] if model_idx < len(values) else '' or '')
        sku = str(values[sku_idx] if sku_idx < len(values) else '' or '').strip()
        hits = [
            article
            for article in targets
            if _row_matches_article(sku, model, article)
        ]
        if not hits:
            continue
        if len(hits) > 1:
            hits = [max(hits, key=lambda item: len(_compact_key(item)))]
        article = hits[0]
        price = values[price_idx] if price_idx is not None and price_idx < len(values) else None
        kaspi_qty_raw = (
            values[pp2_idx] if pp2_idx is not None and pp2_idx < len(values) else None
        )
        kaspi_qty = _parse_export_qty(kaspi_qty_raw)
        price_int = None
        if isinstance(price, (int, float)) and not isinstance(price, bool):
            price_int = int(price)
        matched[article].append(
            {
                'sku': sku,
                'price': price_int,
                'kaspi_qty': kaspi_qty,
            }
        )

    rows: list[KaspiStockPreviewRow] = []
    hold_groups = []
    no_listing = []
    for article, qty in targets.items():
        found = matched.get(article) or []
        status = _status_for_article(
            article,
            len(found),
            declared_no_listing_only=True,
        )
        if status == STATUS_NO_LISTING:
            no_listing.append(article)
            rows.append(
                KaspiStockPreviewRow(
                    article=article,
                    pp2_target=qty,
                    status=STATUS_NO_LISTING,
                    note='No active Kaspi listing',
                )
            )
            continue
        if status == STATUS_HOLD_DUPLICATE:
            hold_groups.append(article)
            items = list(found)
            while len(items) < 2:
                items.append({'sku': '', 'price': None, 'kaspi_qty': None})
            for item in items:
                rows.append(
                    KaspiStockPreviewRow(
                        sku_kaspi=item['sku'],
                        article=article,
                        current_price=item['price'],
                        last_known_kaspi_qty=item['kaspi_qty'],
                        pp2_target=None,
                        delta=None,
                        status=STATUS_HOLD_DUPLICATE,
                        note=(
                            f'PP2 physical qty={qty}; do not copy into each listing'
                        ),
                    )
                )
            continue
        if not found:
            rows.append(
                KaspiStockPreviewRow(
                    article=article,
                    pp2_target=qty,
                    status=STATUS_READY,
                    note='Single listing expected; SKU not found in kaspi_active.xlsx',
                )
            )
            continue
        item = found[0]
        rows.append(
            KaspiStockPreviewRow(
                sku_kaspi=item['sku'],
                article=article,
                current_price=item['price'],
                last_known_kaspi_qty=item['kaspi_qty'],
                pp2_target=qty,
                delta=_delta(qty, item['kaspi_qty']),
                status=STATUS_READY,
                note='Single listing; target = PP2 only. PREVIEW_ONLY.',
            )
        )
    return _finalize(rows, hold_groups, no_listing)


def _parse_export_qty(raw):
    if raw is None or raw == '':
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        if isinstance(raw, float) and not raw.is_integer():
            return None
        qty = int(raw)
        return qty if qty >= 0 else None
    text = str(raw).strip()
    if text.lower() in {'no', 'yes', 'preorder'}:
        return None
    if text.isdigit():
        return int(text)
    return None


def _finalize(rows, hold_groups, no_listing) -> KaspiStockPreviewResult:
    unique_holds = []
    seen = set()
    for article in hold_groups:
        if article not in seen:
            unique_holds.append(article)
            seen.add(article)
    return KaspiStockPreviewResult(
        rows=rows,
        upload_ready=False,
        ready_count=sum(1 for row in rows if row.status == STATUS_READY),
        hold_listing_count=sum(
            1 for row in rows if row.status == STATUS_HOLD_DUPLICATE
        ),
        duplicate_groups=unique_holds,
        no_listing=list(no_listing),
    )


def write_kaspi_stock_preview_xlsx(result: KaspiStockPreviewResult, path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    preview = workbook.active
    preview.title = 'PREVIEW_ONLY'
    preview.append(list(PREVIEW_HEADERS))
    for cell in preview[1]:
        cell.font = Font(bold=True)
    for row in result.rows:
        preview.append(
            [
                row.listing_id if row.listing_id is not None else '',
                row.sku_kaspi,
                row.article,
                row.current_price if row.current_price is not None else '',
                (
                    row.last_known_kaspi_qty
                    if row.last_known_kaspi_qty is not None
                    else ''
                ),
                row.pp2_target if row.pp2_target is not None else '',
                row.delta if row.delta is not None else '',
                row.status,
                row.note,
            ]
        )
    for index, width in enumerate((14, 18, 22, 16, 22, 14, 10, 18, 54), start=1):
        preview.column_dimensions[get_column_letter(index)].width = width

    hold = workbook.create_sheet('HOLD_DUPLICATE')
    hold.append(['article', 'listing_count', 'reason'])
    for cell in hold[1]:
        cell.font = Font(bold=True)
    for article in result.duplicate_groups:
        count = sum(
            1
            for row in result.rows
            if row.article == article and row.status == STATUS_HOLD_DUPLICATE
        )
        hold.append(
            [
                article,
                count,
                'Two active Kaspi listings share one PP2 physical qty',
            ]
        )

    notes = workbook.create_sheet('NOTES')
    notes.append(['key', 'value'])
    notes['A1'].font = Font(bold=True)
    notes['B1'].font = Font(bold=True)
    notes.append(['upload_ready', 'NO'])
    notes.append(['kaspi_writes', 0])
    notes.append(['prices_changed', 0])
    notes.append(['pp2_source', 'Rapido physical inventory 2026-09-18'])
    notes.append(['do_not_sum', 'PP1+PP2 is forbidden; target is PP2 only'])
    notes.append(
        [
            'native_kaspi_format',
            'Kaspi export uses SKU/model/brand/price/PP1-PP5/preorder; '
            'not emitted as upload-ready because HOLD_DUPLICATE listings '
            'must not receive the full PP2 qty.',
        ]
    )
    notes.append(['ready_listings', result.ready_count])
    notes.append(['hold_listings', result.hold_listing_count])
    notes.append(['duplicate_groups', len(result.duplicate_groups)])
    notes.append(['no_listing', ', '.join(result.no_listing)])
    workbook.save(target)
    return target
