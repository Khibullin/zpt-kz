"""Read-only Kaspi listing facts from XLSX/CSV.

Dry-run is the default. Never writes Product, warehouse stock, StockMovement,
publish flags, or Kaspi API. Observed Kaspi qty is never written to PP2.
"""

from __future__ import annotations

import csv
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction
from django.utils import timezone

from catalog.ag_parts_import import cell_text, load_sheet_rows, normalize_header
from catalog.import_ops import sha256_file
from catalog.models import (
    CatalogImportBatch,
    KaspiListingFactSnapshot,
    ProductKaspiListing,
    ProductWarehouseStock,
    SellerProfile,
    Warehouse,
)
from catalog.warehouse_stock_sync import parse_stock_quantity
from catalog.warehouses import WAREHOUSE_CODE_PP2

STATUS_MATCHED_NO_CHANGE = 'MATCHED_NO_CHANGE'
STATUS_WOULD_UPDATE = 'WOULD_UPDATE'
STATUS_UPDATED = 'UPDATED'
STATUS_MISSING_LISTING = 'MISSING_LISTING'
STATUS_AMBIGUOUS_LISTING = 'AMBIGUOUS_LISTING'
STATUS_INVALID_PRICE = 'INVALID_PRICE'
STATUS_INVALID_QUANTITY = 'INVALID_QUANTITY'
STATUS_DUPLICATE_MASTER_SKU = 'DUPLICATE_MASTER_SKU'

RECON_IN_SYNC = 'IN_SYNC'
RECON_KASPI_LOWER = 'KASPI_LOWER'
RECON_KASPI_HIGHER = 'KASPI_HIGHER'
RECON_NO_PP2_BALANCE = 'NO_PP2_BALANCE'

SAFE_APPLY_STATUSES = frozenset({
    STATUS_WOULD_UPDATE,
    STATUS_MATCHED_NO_CHANGE,
})

PROBLEM_STATUSES = frozenset({
    STATUS_MISSING_LISTING,
    STATUS_AMBIGUOUS_LISTING,
    STATUS_INVALID_PRICE,
    STATUS_INVALID_QUANTITY,
    STATUS_DUPLICATE_MASTER_SKU,
})

SKU_ALIASES = (
    'sku',
    'master_sku',
    'kaspi sku',
    'kaspi_sku',
    'kaspi id',
    'kaspi_id',
)
PRICE_ALIASES = (
    'price',
    'цена',
    'цена kaspi',
    'kaspi price',
    'kaspi_price',
)
QUANTITY_ALIASES = (
    'pp2',
    'qty',
    'quantity',
    'остаток',
    'available',
    'avail',
    'kaspi qty',
    'kaspi_qty',
)

SOURCE_CHOICES = {
    KaspiListingFactSnapshot.SOURCE_ACTIVE_XLSX,
    KaspiListingFactSnapshot.SOURCE_MANUAL_IMPORT,
}

PP2_COLUMN_WARNING = (
    'Колонка количества в этом файле — наблюдаемый остаток Kaspi, '
    'не склад Rapido PP2. Значение не пишется в ProductWarehouseStock.'
)


@dataclass
class KaspiFactRow:
    row_number: int
    master_sku: str = ''
    raw_price: object = None
    raw_qty: object = None
    observed_price: int | None = None
    observed_qty: int | None = None
    listing_id: int | None = None
    product_id: int | None = None
    product_article: str = ''
    merchant_sku: str = ''
    pp2_qty: int | None = None
    pp2_has_row: bool = False
    qty_delta: int | None = None
    status: str = ''
    reconciliation: str = ''
    reason: str = ''


@dataclass
class KaspiFactSyncResult:
    rows: list[KaspiFactRow]
    seller_id: int
    seller_name: str
    source_path: str
    source_sha256: str
    apply: bool
    sheet_name: str = ''
    headers: list = field(default_factory=list)
    column_map: dict = field(default_factory=dict)
    quantity_header: str = ''
    products_with_multiple_listings: int = 0
    batch_id: int | None = None


def _header_index(headers, aliases):
    normalized = [normalize_header(header) for header in headers]
    wanted = {normalize_header(alias) for alias in aliases}
    for index, header in enumerate(normalized):
        if header in wanted:
            return index
    return None


def detect_kaspi_fact_columns(
    headers,
    *,
    sku_column='',
    price_column='',
    quantity_column='',
) -> dict[str, int]:
    mapping = {}
    sku_idx = (
        _header_index(headers, (sku_column,))
        if sku_column
        else _header_index(headers, SKU_ALIASES)
    )
    price_idx = (
        _header_index(headers, (price_column,))
        if price_column
        else _header_index(headers, PRICE_ALIASES)
    )
    qty_idx = (
        _header_index(headers, (quantity_column,))
        if quantity_column
        else _header_index(headers, QUANTITY_ALIASES)
    )
    if sku_idx is not None:
        mapping['master_sku'] = sku_idx
    if price_idx is not None:
        mapping['price'] = price_idx
    if qty_idx is not None:
        mapping['quantity'] = qty_idx
    return mapping


def _load_xlsx_tables(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet_names = list(workbook.sheetnames)
    workbook.close()
    tables = []
    for sheet_name in sheet_names:
        headers, _column_map, data_rows, _images = load_sheet_rows(path, sheet_name)
        tables.append((sheet_name, headers, data_rows))
    return tables


def _load_csv_table(path: Path):
    with path.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = csv.reader(handle)
        try:
            header_row = next(reader)
        except StopIteration:
            return path.name, [], []
        headers = [cell_text(value) for value in header_row]
        data_rows = []
        for index, row in enumerate(reader, start=2):
            values = list(row)
            if not any(cell_text(value) for value in values):
                continue
            data_rows.append((index, values))
        return path.name, headers, data_rows


def load_kaspi_fact_rows(
    path: Path,
    *,
    sheet='',
    sku_column='',
    price_column='',
    quantity_column='',
):
    suffix = path.suffix.lower()
    detect_kwargs = {
        'sku_column': sku_column,
        'price_column': price_column,
        'quantity_column': quantity_column,
    }
    if suffix == '.csv':
        sheet_name, headers, data_rows = _load_csv_table(path)
        column_map = detect_kaspi_fact_columns(headers, **detect_kwargs)
        return sheet_name, headers, column_map, data_rows
    if suffix not in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        raise ValueError(f'unsupported_file_type:{suffix or "missing"}')

    if sheet:
        headers, _column_map, data_rows, _images = load_sheet_rows(path, sheet)
        column_map = detect_kaspi_fact_columns(headers, **detect_kwargs)
        return sheet, headers, column_map, data_rows

    chosen = None
    fallback = None
    for sheet_name, headers, data_rows in _load_xlsx_tables(path):
        column_map = detect_kaspi_fact_columns(headers, **detect_kwargs)
        candidate = (sheet_name, headers, column_map, data_rows)
        if fallback is None:
            fallback = candidate
        if {'master_sku', 'price', 'quantity'} <= set(column_map):
            chosen = candidate
            break
    selected = chosen or fallback
    if selected is None:
        raise ValueError('workbook_has_no_sheets')
    return selected


def _require_columns(headers, column_map):
    missing = [
        name for name in ('master_sku', 'price', 'quantity')
        if name not in column_map
    ]
    if missing:
        raise ValueError(
            f'Не найдены колонки {missing}. Заголовки: {headers}'
        )


def _cell_at(values, column_map, field):
    index = column_map.get(field)
    if index is None or index >= len(values):
        return None
    return values[index]


def _pp2_balances(product_ids):
    warehouse = Warehouse.objects.filter(
        code=WAREHOUSE_CODE_PP2,
        is_active=True,
    ).first()
    if warehouse is None or not product_ids:
        return {}
    rows = ProductWarehouseStock.objects.filter(
        product_id__in=product_ids,
        warehouse=warehouse,
    ).values_list('product_id', 'quantity')
    return {product_id: int(quantity) for product_id, quantity in rows}


def _reconcile(observed_qty, pp2_qty, has_row):
    if not has_row:
        return RECON_NO_PP2_BALANCE, None
    delta = observed_qty - pp2_qty
    if delta == 0:
        return RECON_IN_SYNC, delta
    if observed_qty < pp2_qty:
        return RECON_KASPI_LOWER, delta
    return RECON_KASPI_HIGHER, delta


def _cache_differs(listing, price, qty):
    return (
        listing.last_known_our_price != price
        or listing.last_known_kaspi_qty != qty
    )


def classify_kaspi_fact_rows(rows, seller: SellerProfile):
    sku_counts = Counter(
        row.master_sku for row in rows if row.master_sku
    )
    listings_by_sku = defaultdict(list)
    qs = (
        ProductKaspiListing.objects.filter(product__seller_profile=seller)
        .select_related('product')
    )
    for listing in qs:
        listings_by_sku[listing.master_sku].append(listing)

    product_ids = {
        listing.product_id
        for listing_list in listings_by_sku.values()
        for listing in listing_list
    }
    pp2_map = _pp2_balances(product_ids)
    listing_counts_by_product = Counter(
        listing.product_id
        for listing_list in listings_by_sku.values()
        for listing in listing_list
    )

    for row in rows:
        if not row.master_sku:
            row.status = STATUS_MISSING_LISTING
            row.reason = 'empty_master_sku'
            continue
        if sku_counts[row.master_sku] > 1:
            row.status = STATUS_DUPLICATE_MASTER_SKU
            row.reason = 'duplicate_master_sku_in_source'
            continue

        price, price_error = parse_stock_quantity(row.raw_price)
        if price_error:
            row.status = STATUS_INVALID_PRICE
            row.reason = f'invalid_price:{price_error}'
            continue
        qty, qty_error = parse_stock_quantity(row.raw_qty)
        if qty_error:
            row.status = STATUS_INVALID_QUANTITY
            row.reason = f'invalid_quantity:{qty_error}'
            continue
        row.observed_price = price
        row.observed_qty = qty

        matches = listings_by_sku.get(row.master_sku, [])
        if not matches:
            row.status = STATUS_MISSING_LISTING
            row.reason = 'listing_not_found'
            continue
        if len(matches) > 1:
            row.status = STATUS_AMBIGUOUS_LISTING
            row.reason = 'multiple_listings_same_master_sku'
            continue

        listing = matches[0]
        row.listing_id = listing.pk
        row.product_id = listing.product_id
        row.product_article = listing.product.article or ''
        row.merchant_sku = listing.merchant_sku or ''
        if listing.product_id in pp2_map:
            row.pp2_has_row = True
            row.pp2_qty = pp2_map[listing.product_id]
        row.reconciliation, row.qty_delta = _reconcile(
            qty,
            row.pp2_qty if row.pp2_has_row else 0,
            row.pp2_has_row,
        )
        if _cache_differs(listing, price, qty):
            row.status = STATUS_WOULD_UPDATE
            row.reason = 'cache_differs'
        else:
            row.status = STATUS_MATCHED_NO_CHANGE
            row.reason = 'unchanged'

    products_with_multiple = {
        listing.product_id
        for listing_list in listings_by_sku.values()
        for listing in listing_list
        if listing_counts_by_product[listing.product_id] > 1
    }
    matched_products = {
        row.product_id for row in rows if row.product_id
    }
    return len(matched_products & products_with_multiple)


def apply_kaspi_fact_rows(
    rows,
    *,
    seller: SellerProfile,
    source_path: str,
    source_sha256: str,
    source: str,
    started_at,
):
    now = timezone.now()
    listing_ids = [row.listing_id for row in rows if row.listing_id]
    listings = {
        listing.pk: listing
        for listing in ProductKaspiListing.objects.filter(pk__in=listing_ids)
        .select_related('product')
    }
    updated = 0
    snapshots_created = 0
    unchanged = 0
    with transaction.atomic():
        batch = CatalogImportBatch.objects.create(
            seller_profile=seller,
            source=CatalogImportBatch.SOURCE_KASPI_LISTING_FACTS,
            filename=Path(source_path).name,
            file_sha256=source_sha256,
            started_at=started_at,
            finished_at=now,
            mode=CatalogImportBatch.MODE_WRITE,
            source_scope=CatalogImportBatch.SCOPE_PARTIAL,
            status=CatalogImportBatch.STATUS_SUCCESS,
            source_row_count=len(rows),
            source_unique_count=len({
                row.master_sku for row in rows if row.master_sku
            }),
        )
        for row in rows:
            if row.status not in SAFE_APPLY_STATUSES:
                continue
            listing = listings.get(row.listing_id)
            if listing is None:
                continue
            if _cache_differs(listing, row.observed_price, row.observed_qty):
                listing.last_known_our_price = row.observed_price
                listing.last_known_kaspi_qty = row.observed_qty
                listing.last_synced_at = now
                listing.save(update_fields=[
                    'last_known_our_price',
                    'last_known_kaspi_qty',
                    'last_synced_at',
                    'updated_at',
                ])
                row.status = STATUS_UPDATED
                row.reason = 'updated'
                updated += 1
            else:
                row.status = STATUS_MATCHED_NO_CHANGE
                unchanged += 1
            _, created = KaspiListingFactSnapshot.objects.get_or_create(
                listing=listing,
                source_sha256=source_sha256,
                defaults={
                    'observed_price': row.observed_price,
                    'observed_qty': row.observed_qty,
                    'source': source,
                    'source_filename': Path(source_path).name,
                    'source_row': row.row_number,
                    'observed_at': now,
                    'import_batch': batch,
                },
            )
            if created:
                snapshots_created += 1
        skipped = sum(1 for row in rows if row.status in PROBLEM_STATUSES)
        conflict = sum(
            1 for row in rows
            if row.status in {
                STATUS_AMBIGUOUS_LISTING,
                STATUS_DUPLICATE_MASTER_SKU,
            }
        )
        batch.selected_count = updated + unchanged
        batch.created_count = snapshots_created
        batch.updated_count = updated
        batch.unchanged_count = unchanged
        batch.skipped_count = skipped
        batch.conflict_count = conflict
        batch.finished_at = timezone.now()
        batch.save(update_fields=[
            'selected_count',
            'created_count',
            'updated_count',
            'unchanged_count',
            'skipped_count',
            'conflict_count',
            'finished_at',
        ])
    return batch


def summarize_kaspi_fact_rows(rows):
    counts = Counter(row.status for row in rows)
    recon = Counter(
        row.reconciliation for row in rows if row.reconciliation
    )
    return {
        'total_rows': len(rows),
        'matched_no_change': counts[STATUS_MATCHED_NO_CHANGE],
        'would_update': counts[STATUS_WOULD_UPDATE],
        'updated': counts[STATUS_UPDATED],
        'missing_listing': counts[STATUS_MISSING_LISTING],
        'ambiguous_listing': counts[STATUS_AMBIGUOUS_LISTING],
        'invalid_price': counts[STATUS_INVALID_PRICE],
        'invalid_quantity': counts[STATUS_INVALID_QUANTITY],
        'duplicate_master_sku': counts[STATUS_DUPLICATE_MASTER_SKU],
        'in_sync': recon[RECON_IN_SYNC],
        'kaspi_lower': recon[RECON_KASPI_LOWER],
        'kaspi_higher': recon[RECON_KASPI_HIGHER],
        'no_pp2_balance': recon[RECON_NO_PP2_BALANCE],
    }


def sync_kaspi_listing_facts(
    *,
    path,
    seller: SellerProfile,
    apply: bool = False,
    sheet='',
    sku_column='',
    price_column='',
    quantity_column='',
    source=KaspiListingFactSnapshot.SOURCE_ACTIVE_XLSX,
):
    if source not in SOURCE_CHOICES:
        raise ValueError(f'unsupported_source:{source}')
    source_path = Path(path)
    started_at = timezone.now()
    file_digest = sha256_file(source_path)
    sheet_name, headers, column_map, data_rows = load_kaspi_fact_rows(
        source_path,
        sheet=sheet,
        sku_column=sku_column,
        price_column=price_column,
        quantity_column=quantity_column,
    )
    _require_columns(headers, column_map)

    rows = []
    for row_number, values in data_rows:
        raw_sku = _cell_at(values, column_map, 'master_sku')
        rows.append(KaspiFactRow(
            row_number=row_number,
            master_sku=cell_text(raw_sku),
            raw_price=_cell_at(values, column_map, 'price'),
            raw_qty=_cell_at(values, column_map, 'quantity'),
        ))

    multi = classify_kaspi_fact_rows(rows, seller)
    batch = None
    if apply:
        batch = apply_kaspi_fact_rows(
            rows,
            seller=seller,
            source_path=str(source_path),
            source_sha256=file_digest,
            source=source,
            started_at=started_at,
        )

    qty_header = ''
    qty_idx = column_map.get('quantity')
    if qty_idx is not None and qty_idx < len(headers):
        qty_header = cell_text(headers[qty_idx])

    return KaspiFactSyncResult(
        rows=rows,
        seller_id=seller.pk,
        seller_name=seller.name,
        source_path=str(source_path),
        source_sha256=file_digest,
        apply=apply,
        sheet_name=sheet_name,
        headers=list(headers),
        column_map=dict(column_map),
        quantity_header=qty_header,
        products_with_multiple_listings=multi,
        batch_id=batch.pk if batch is not None else None,
    )
