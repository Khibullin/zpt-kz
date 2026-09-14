from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth.models import User
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook, load_workbook

from catalog.kaspi_listing_sync import (
    STATUS_BARCODE_CONFLICT,
    STATUS_CONFLICT_MASTER_SKU,
    STATUS_CONFLICT_PRODUCT,
    STATUS_DUPLICATE_INPUT_ARTICLE,
    STATUS_MATCHED,
    STATUS_MISSING_KASPI_ID,
    STATUS_MISSING_PRODUCT,
    STATUS_MULTIPLE_MASTER_SKU,
    STATUS_WOULD_CREATE_LISTING,
    STATUS_WOULD_UPDATE_LISTING,
    detect_kaspi_link_columns,
    split_master_skus,
    sync_kaspi_product_links,
)
from catalog.models import Product, ProductBarcode, ProductKaspiListing, SellerProfile


def _xlsx(path, headers, rows, sheet='Sheet'):
    workbook = Workbook()
    sheet_obj = workbook.active
    sheet_obj.title = sheet
    sheet_obj.append(headers)
    for row in rows:
        sheet_obj.append(row)
    workbook.save(path)


def _make_seller(username, name, phone):
    user = User.objects.create_user(username=username, password='secret12345')
    return SellerProfile.objects.create(
        user=user,
        name=name,
        phone=phone,
        city='Алматы',
    )


def _make_product(seller, article, **kwargs):
    defaults = {
        'title': f'Title {article}',
        'article': article,
        'seller_name': seller.name,
        'whatsapp_number': seller.phone,
        'seller_profile': seller,
        'city': seller.city,
        'status': 'active',
        'price': 1000,
        'cost_price': 400,
        'stock_qty': 7,
        'publish_to_kaspi': False,
    }
    defaults.update(kwargs)
    return Product.objects.create(**defaults)


class DetectKaspiLinkColumnsTests(TestCase):
    def test_kaspi_id_is_not_article(self):
        mapping = detect_kaspi_link_columns(
            ['Артикул', 'Kaspi ID', 'Merchant SKU', 'Штрихкод']
        )
        self.assertEqual(mapping['article'], 0)
        self.assertEqual(mapping['master_sku'], 1)
        self.assertEqual(mapping['merchant_sku'], 2)
        self.assertEqual(mapping['barcode'], 3)

    def test_sku_becomes_master_when_article_already_mapped(self):
        mapping = detect_kaspi_link_columns(['Артикул', 'SKU'])
        self.assertEqual(mapping['article'], 0)
        self.assertEqual(mapping['master_sku'], 1)

    def test_sku_alone_is_not_article(self):
        mapping = detect_kaspi_link_columns(['SKU', 'kaspi_id'])
        self.assertNotIn('article', mapping)
        self.assertEqual(mapping['master_sku'], 1)

    def test_sku_only_headers_have_no_article(self):
        mapping = detect_kaspi_link_columns(['SKU', 'model', 'brand', 'price'])
        self.assertEqual(mapping, {})


class SyncKaspiProductLinksTests(TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.seller = _make_seller('own-seller', 'AG Parts', '77771111111')
        self.other = _make_seller('other-seller', 'Other Shop', '77772222222')
        self.product = _make_product(self.seller, 'AF-200')
        self.other_product = _make_product(self.other, 'AF-200', title='Other AF-200')

    def _file(self, name, headers, rows):
        path = self.root / name
        _xlsx(path, headers, rows)
        return path

    def _run(self, path, **kwargs):
        stdout = StringIO()
        call_kwargs = {
            'seller_profile_id': self.seller.pk,
        }
        call_kwargs.update(kwargs)
        call_command(
            'sync_kaspi_product_links',
            str(path),
            stdout=stdout,
            **call_kwargs,
        )
        return stdout.getvalue()

    def _sync(self, path, apply=False):
        return sync_kaspi_product_links(
            path=path,
            seller=self.seller,
            apply=apply,
        )

    def test_matches_seller_profile_and_article(self):
        path = self._file(
            'ok.xlsx',
            ['Артикул', 'kaspi_id', 'merchant_sku', 'barcode'],
            [['AF-200', '601717696', 'AF-200', '111']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_WOULD_CREATE_LISTING)
        self.assertEqual(row.product_id, self.product.pk)
        self.assertEqual(row.master_sku, '601717696')

    def test_same_article_other_seller_is_not_match(self):
        path = self._file(
            'other.xlsx',
            ['article', 'kaspi_id'],
            [['AF-200', '601717696']],
        )
        other_result = sync_kaspi_product_links(
            path=path,
            seller=self.other,
            apply=False,
        )
        self.assertEqual(other_result.rows[0].product_id, self.other_product.pk)
        own = self._sync(path)
        self.assertEqual(own.rows[0].product_id, self.product.pk)
        self.assertNotEqual(own.rows[0].product_id, self.other_product.pk)

    def test_missing_product(self):
        path = self._file(
            'missing.xlsx',
            ['артикул', 'kaspi_id'],
            [['NO-SUCH', '601717696']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_MISSING_PRODUCT)
        self.assertEqual(row.reason, 'product_not_found')

    def test_existing_matching_listing(self):
        listing = ProductKaspiListing.objects.create(
            product=self.product,
            master_sku='601717696',
            merchant_sku='AF-200',
            barcode='111',
        )
        ProductBarcode.objects.create(
            product=self.product,
            code='111',
            source=ProductBarcode.SOURCE_KASPI,
            is_primary=True,
        )
        path = self._file(
            'matched.xlsx',
            ['Артикул', 'kaspi_id', 'merchant_sku', 'штрихкод'],
            [['AF-200', '601717696', 'AF-200', '111']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_MATCHED)
        self.assertEqual(row.existing_listing_id, listing.pk)

    def test_would_create_listing(self):
        path = self._file(
            'create.xlsx',
            ['Артикул', 'master_sku'],
            [['AF-200', '601717696']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_WOULD_CREATE_LISTING)
        self.assertIsNone(row.existing_listing_id)

    def test_duplicate_article_in_input(self):
        path = self._file(
            'dup.xlsx',
            ['Артикул', 'kaspi_id'],
            [
                ['AF-200', '111'],
                ['AF-200', '222'],
            ],
        )
        statuses = [row.status for row in self._sync(path).rows]
        self.assertEqual(statuses, [
            STATUS_DUPLICATE_INPUT_ARTICLE,
            STATUS_DUPLICATE_INPUT_ARTICLE,
        ])
        self.assertEqual(ProductKaspiListing.objects.count(), 0)

    def test_one_master_sku_two_products_is_conflict(self):
        second = _make_product(self.seller, 'BF-300')
        path = self._file(
            'master-conflict.xlsx',
            ['Артикул', 'kaspi_id'],
            [
                ['AF-200', 'SHARED-1'],
                ['BF-300', 'SHARED-1'],
            ],
        )
        result = self._sync(path)
        self.assertEqual(
            {row.status for row in result.rows},
            {STATUS_CONFLICT_MASTER_SKU},
        )
        self.assertTrue(
            {row.product_id for row in result.rows} >= {self.product.pk, second.pk}
        )

    def test_existing_master_sku_on_other_product_is_conflict(self):
        second = _make_product(self.seller, 'BF-300')
        ProductKaspiListing.objects.create(
            product=second,
            master_sku='TAKEN',
            merchant_sku='BF-300',
        )
        path = self._file(
            'taken.xlsx',
            ['Артикул', 'kaspi_id'],
            [['AF-200', 'TAKEN']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_CONFLICT_MASTER_SKU)
        self.assertEqual(row.reason, 'master_sku_bound_to_other_product')

    def test_one_product_two_master_skus_in_db_without_match_is_conflict(self):
        ProductKaspiListing.objects.create(
            product=self.product, master_sku='A1', merchant_sku='AF-200',
        )
        ProductKaspiListing.objects.create(
            product=self.product, master_sku='A2', merchant_sku='AF-200-B',
        )
        path = self._file(
            'product-conflict.xlsx',
            ['Артикул', 'kaspi_id'],
            [['AF-200', 'A3']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_CONFLICT_PRODUCT)

    def test_barcode_existing_same_product_is_matched(self):
        ProductKaspiListing.objects.create(
            product=self.product,
            master_sku='601717696',
            merchant_sku='AF-200',
            barcode='111',
        )
        ProductBarcode.objects.create(
            product=self.product,
            code='111',
            source=ProductBarcode.SOURCE_WMS,
            is_primary=True,
        )
        path = self._file(
            'barcode-same.xlsx',
            ['Артикул', 'kaspi_id', 'barcode'],
            [['AF-200', '601717696', '111']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_MATCHED)

    def test_barcode_on_other_product_is_not_conflict(self):
        """Unique constraint is (product, code), not global barcode uniqueness."""
        ProductBarcode.objects.create(
            product=self.other_product,
            code='111',
            source=ProductBarcode.SOURCE_KASPI,
        )
        path = self._file(
            'barcode-other.xlsx',
            ['Артикул', 'kaspi_id', 'barcode'],
            [['AF-200', '601717696', '111']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_WOULD_CREATE_LISTING)
        self.assertNotEqual(row.status, STATUS_BARCODE_CONFLICT)
        self._run(path, apply=True)
        self.assertTrue(
            ProductBarcode.objects.filter(product=self.product, code='111').exists()
        )
        self.assertTrue(
            ProductBarcode.objects.filter(product=self.other_product, code='111').exists()
        )

    def test_default_run_does_not_write_db(self):
        path = self._file(
            'dry.xlsx',
            ['Артикул', 'kaspi_id', 'barcode'],
            [['AF-200', '601717696', '111']],
        )
        output = self._run(path)
        self.assertIn('mode: dry-run', output)
        self.assertEqual(ProductKaspiListing.objects.count(), 0)
        self.assertEqual(ProductBarcode.objects.count(), 0)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 1000)
        self.assertEqual(self.product.cost_price, 400)
        self.assertEqual(self.product.stock_qty, 7)
        self.assertFalse(self.product.publish_to_kaspi)

    def test_apply_writes_only_safe_rows(self):
        conflict_product = _make_product(self.seller, 'CF-400')
        ProductKaspiListing.objects.create(
            product=conflict_product,
            master_sku='TAKEN',
        )
        path = self._file(
            'apply.xlsx',
            ['Артикул', 'kaspi_id', 'merchant_sku', 'barcode'],
            [
                ['AF-200', '601717696', 'AF-200', '111'],
                ['NO-SUCH', '999', '', ''],
                ['CF-400', 'TAKEN', '', ''],
            ],
        )
        output = self._run(path, apply=True)
        self.assertIn('mode: apply', output)
        created = ProductKaspiListing.objects.get(product=self.product)
        self.assertEqual(created.master_sku, '601717696')
        self.assertEqual(created.merchant_sku, 'AF-200')
        self.assertEqual(created.barcode, '111')
        self.assertFalse(created.publish_to_kaspi)
        self.assertEqual(
            ProductBarcode.objects.get(product=self.product, code='111').source,
            ProductBarcode.SOURCE_KASPI,
        )
        self.assertEqual(ProductKaspiListing.objects.filter(product=conflict_product).count(), 1)
        self.assertFalse(Product.objects.filter(article='NO-SUCH').exists())
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 1000)
        self.assertEqual(self.product.cost_price, 400)
        self.assertEqual(self.product.stock_qty, 7)
        self.assertFalse(self.product.publish_to_kaspi)

    def test_conflict_rows_do_not_change_db_on_apply(self):
        listing = ProductKaspiListing.objects.create(
            product=self.product,
            master_sku='KEEP',
            merchant_sku='OLD',
            barcode='000',
        )
        second = _make_product(self.seller, 'BF-300')
        path = self._file(
            'conflicts.xlsx',
            ['Артикул', 'kaspi_id'],
            [
                ['AF-200', 'KEEP'],
                ['BF-300', 'KEEP'],
            ],
        )
        self._run(path, apply=True)
        listing.refresh_from_db()
        self.assertEqual(listing.master_sku, 'KEEP')
        self.assertEqual(listing.merchant_sku, 'OLD')
        self.assertFalse(ProductKaspiListing.objects.filter(product=second).exists())

    def test_would_update_merchant_and_barcode(self):
        listing = ProductKaspiListing.objects.create(
            product=self.product,
            master_sku='601717696',
            merchant_sku='OLD',
            barcode='',
        )
        path = self._file(
            'update.xlsx',
            ['Артикул', 'kaspi_id', 'артикул продавца', 'ean'],
            [['AF-200', '601717696', 'NEW-MERCHANT', '222']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_WOULD_UPDATE_LISTING)
        self._run(path, apply=True)
        listing.refresh_from_db()
        self.assertEqual(listing.merchant_sku, 'NEW-MERCHANT')
        self.assertEqual(listing.barcode, '222')
        self.assertEqual(listing.master_sku, '601717696')
        self.assertTrue(
            ProductBarcode.objects.filter(product=self.product, code='222').exists()
        )

    def test_float_article_and_missing_kaspi_id(self):
        numeric = _make_product(self.seller, '12345')
        path = self._file(
            'float.xlsx',
            ['Артикул', 'kaspi_id'],
            [[12345.0, ''], ['AF-200', '']],
        )
        result = self._sync(path)
        by_article = {row.article: row for row in result.rows}
        self.assertEqual(by_article['12345'].product_id, numeric.pk)
        self.assertEqual(by_article['12345'].status, STATUS_MISSING_KASPI_ID)
        self.assertEqual(by_article['AF-200'].status, STATUS_MISSING_KASPI_ID)

    def test_report_xlsx_written(self):
        path = self._file(
            'report-src.xlsx',
            ['Артикул', 'kaspi_id'],
            [['AF-200', '601717696']],
        )
        report = self.root / 'out.xlsx'
        self._run(path, report=str(report))
        workbook = load_workbook(report)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(
            headers,
            [
                'row_number', 'article', 'product_id', 'product_name',
                'master_sku', 'merchant_sku', 'barcode',
                'existing_listing_id', 'status', 'reason',
            ],
        )
        self.assertEqual(sheet[2][8].value, STATUS_WOULD_CREATE_LISTING)

    def test_missing_seller_raises(self):
        path = self._file(
            'no-seller.xlsx',
            ['Артикул', 'kaspi_id'],
            [['AF-200', '1']],
        )
        with self.assertRaises(CommandError):
            call_command(
                'sync_kaspi_product_links',
                str(path),
                seller_profile_id=999999,
                stdout=StringIO(),
            )

    def test_sku_only_file_does_not_treat_sku_as_article(self):
        path = self._file(
            'kaspi-export.xlsx',
            ['SKU', 'model', 'brand', 'price'],
            [['601717696', 'EXEED фильтр 151000187AA', 'ag-parts', 2200]],
        )
        with self.assertRaises(CommandError) as raised:
            self._run(path)
        self.assertIn('SKU не считается артикулом', str(raised.exception))
        self.assertEqual(ProductKaspiListing.objects.count(), 0)

    def test_pipe_separated_master_skus_are_multiple(self):
        path = self._file(
            'pipe.xlsx',
            ['Артикул', 'kaspi_id'],
            [['AF-200', '123 | 456']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_MULTIPLE_MASTER_SKU)
        self.assertEqual(row.master_skus, ['123', '456'])
        self.assertEqual(row.reason, 'multiple_master_sku: 123, 456')
        self.assertEqual(row.product_id, self.product.pk)

    def test_semicolon_separated_master_skus_are_multiple(self):
        path = self._file(
            'semi.xlsx',
            ['Артикул', 'kaspi_sku'],
            [['AF-200', '123;456']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_MULTIPLE_MASTER_SKU)
        self.assertEqual(row.master_skus, ['123', '456'])

    def test_single_master_sku_is_not_multiple(self):
        self.assertEqual(split_master_skus('123'), ['123'])
        path = self._file(
            'single.xlsx',
            ['Артикул', 'kaspi_id'],
            [['AF-200', '123']],
        )
        row = self._sync(path).rows[0]
        self.assertEqual(row.status, STATUS_WOULD_CREATE_LISTING)
        self.assertEqual(row.master_sku, '123')

    def test_apply_skips_multiple_master_sku(self):
        path = self._file(
            'multi-apply.xlsx',
            ['Артикул', 'kaspi_id'],
            [['AF-200', '123 | 456']],
        )
        output = self._run(path, apply=True)
        self.assertIn('Multiple master SKUs: 1', output)
        self.assertEqual(ProductKaspiListing.objects.count(), 0)
        self.assertEqual(ProductBarcode.objects.count(), 0)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 1000)
        self.assertEqual(self.product.stock_qty, 7)
