"""Reconcile existing PP2 balances to a complete Rapido snapshot.

Exact Product.article match only. Never creates Product, never writes PP1,
never touches Kaspi listings, prices, or Product.stock_qty.

Rows present in the warehouse but missing from the snapshot are left as-is.
Dry-run is the default. Writes require explicit apply=True.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction

from catalog.ag_parts_import import cell_text, extract_article, normalize_header
from catalog.models import Product, ProductWarehouseStock, StockMovement, Warehouse
from catalog.stock_service import StockServiceError, apply_stock_movement, resolve_warehouse
from catalog.warehouse_stock_sync import parse_stock_quantity
from catalog.warehouses import WAREHOUSE_CODE_PP1, WAREHOUSE_CODE_PP2

SOURCE_RAPIDO = 'rapido_inventory'
ALLOWED_WAREHOUSE = WAREHOUSE_CODE_PP2

STATUS_CHANGED = 'CHANGED'
STATUS_UNCHANGED = 'UNCHANGED'
STATUS_UNMATCHED = 'UNMATCHED'
STATUS_ALREADY_APPLIED = 'ALREADY_APPLIED'
STATUS_MISSING_BALANCE = 'MISSING_BALANCE'
STATUS_DUPLICATE_ARTICLE = 'DUPLICATE_ARTICLE'
STATUS_INVALID_ROW = 'INVALID_ROW'
STATUS_AMBIGUOUS_ARTICLE = 'AMBIGUOUS_ARTICLE'
STATUS_ERROR = 'ERROR'

SAFE_APPLY_STATUSES = frozenset({STATUS_CHANGED})

ARTICLE_ALIASES = ('артикул', 'article', 'продукт')
QUANTITY_ALIASES = (
    'qty',
    'quantity',
    'остаток',
    'кол-во',
    'количество',
    'наличие на складе аг',
)


@dataclass
class ReconciliationRow:
    row_number: int
    article: str = ''
    raw_qty: object = None
    qty: int | None = None
    product_id: int | None = None
    product_title: str = ''
    quantity_before: int | None = None
    quantity_after: int | None = None
    quantity_delta: int | None = None
    status: str = ''
    reason: str = ''


@dataclass
class ReconciliationResult:
    rows: list[ReconciliationRow]
    warehouse_code: str
    warehouse_name: str
    reference: str
    source: str
    note: str
    source_path: str
    apply: bool
    sheet_name: str = ''
    headers: list = field(default_factory=list)
    summary: dict = field(default_factory=dict)


def _column_index(headers, aliases):
    wanted = {normalize_header(name) for name in aliases}
    for index, header in enumerate(headers):
        if normalize_header(header) in wanted:
            return index
    return None


def _cell_at(values, index):
    if index is None or index >= len(values):
        return None
    return values[index]


def _load_snapshot_table(path: Path):
    suffix = path.suffix.lower()
    if suffix not in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        raise ValueError(f'unsupported_file_type:{suffix or "missing"}')

    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    try:
        fallback = None
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            headers, article_idx, qty_idx, data_rows = _extract_sheet_table(worksheet)
            candidate = (sheet_name, headers, article_idx, qty_idx, data_rows)
            if fallback is None:
                fallback = candidate
            if article_idx is not None and qty_idx is not None:
                return candidate
    finally:
        workbook.close()
    if fallback is None:
        raise ValueError('workbook_has_no_sheets')
    return fallback


def _extract_sheet_table(worksheet):
    rows = list(worksheet.iter_rows(values_only=True))
    header_index = None
    headers = []
    article_idx = None
    qty_idx = None
    scan_limit = min(len(rows), 30)
    for index in range(scan_limit):
        candidate = [cell_text(value) for value in (rows[index] or ())]
        found_article = _column_index(candidate, ARTICLE_ALIASES)
        found_qty = _column_index(candidate, QUANTITY_ALIASES)
        if found_article is not None and found_qty is not None:
            header_index = index
            headers = candidate
            article_idx = found_article
            qty_idx = found_qty
            break
        if header_index is None and any(candidate):
            headers = candidate
    data_rows = []
    if header_index is None:
        return headers, None, None, data_rows
    for offset, raw in enumerate(rows[header_index + 1 :], start=header_index + 2):
        values = list(raw or ())
        if not any(cell_text(value) for value in values):
            continue
        data_rows.append((offset, values))
    return headers, article_idx, qty_idx, data_rows


def snapshot_rows_from_mapping(mapping: dict[str, int]) -> list[ReconciliationRow]:
    rows = []
    for index, (article, qty) in enumerate(mapping.items(), start=2):
        rows.append(
            ReconciliationRow(
                row_number=index,
                article=str(article).strip(),
                raw_qty=qty,
                qty=int(qty),
            )
        )
    return rows


def snapshot_rows_from_file(path: Path) -> tuple[list[ReconciliationRow], str, list]:
    sheet_name, headers, article_idx, qty_idx, data_rows = _load_snapshot_table(path)
    if article_idx is None:
        raise ValueError(f'article_column_not_found. headers={headers}')
    if qty_idx is None:
        raise ValueError(f'quantity_column_not_found. headers={headers}')
    rows = []
    for row_number, values in data_rows:
        article, _article_key = extract_article(_cell_at(values, article_idx))
        rows.append(
            ReconciliationRow(
                row_number=row_number,
                article=article,
                raw_qty=_cell_at(values, qty_idx),
            )
        )
    return rows, sheet_name, list(headers)


def _find_product(article: str):
    matches = list(Product.objects.filter(article=article).order_by('pk'))
    if not matches:
        return None, 'product_not_found'
    if len(matches) > 1:
        return None, 'ambiguous_article'
    return matches[0], ''


def _existing_reference_movement(*, product, warehouse, source: str, reference: str):
    return StockMovement.objects.filter(
        product=product,
        warehouse=warehouse,
        source=source,
        reference=reference,
    ).exists()


def classify_reconciliation_rows(
    rows: list[ReconciliationRow],
    *,
    warehouse: Warehouse,
    source: str,
    reference: str,
) -> list[ReconciliationRow]:
    article_counts: dict[str, int] = {}
    for row in rows:
        if row.article:
            article_counts[row.article] = article_counts.get(row.article, 0) + 1

    for row in rows:
        if not row.article:
            row.status = STATUS_INVALID_ROW
            row.reason = 'empty_article'
            continue

        qty, qty_error = parse_stock_quantity(row.raw_qty)
        if qty_error != 'empty':
            row.qty = qty
        product, match_error = _find_product(row.article)
        if product is not None:
            row.product_id = product.pk
            row.product_title = product.title

        if article_counts.get(row.article, 0) > 1:
            row.status = STATUS_DUPLICATE_ARTICLE
            row.reason = 'duplicate_source_article'
            continue
        if match_error == 'ambiguous_article':
            row.status = STATUS_AMBIGUOUS_ARTICLE
            row.reason = match_error
            continue
        if product is None:
            row.status = STATUS_UNMATCHED
            row.reason = match_error
            continue
        if qty_error == 'empty':
            row.status = STATUS_ERROR
            row.reason = 'blank_qty'
            continue
        if qty_error:
            row.status = STATUS_ERROR
            row.reason = qty_error
            continue
        row.quantity_after = int(qty)
        existing = ProductWarehouseStock.objects.filter(
            product=product,
            warehouse=warehouse,
        ).first()
        if existing is None:
            row.status = STATUS_MISSING_BALANCE
            row.reason = 'missing_pp2_balance'
            continue
        row.quantity_before = int(existing.quantity)
        row.quantity_delta = row.quantity_after - row.quantity_before
        if _existing_reference_movement(
            product=product,
            warehouse=warehouse,
            source=source,
            reference=reference,
        ):
            row.status = STATUS_ALREADY_APPLIED
            row.reason = 'same_reference_already_applied'
            continue
        if row.quantity_delta == 0:
            row.status = STATUS_UNCHANGED
            row.reason = 'no_change'
            continue
        row.status = STATUS_CHANGED
        row.reason = 'adjust_balance'
    return rows


def warehouse_quantity_total(warehouse: Warehouse) -> int:
    return int(
        sum(
            ProductWarehouseStock.objects.filter(warehouse=warehouse).values_list(
                'quantity', flat=True
            )
        )
    )


def summarize_reconciliation_rows(
    rows: list[ReconciliationRow],
    *,
    existing_pp2_total: int,
    persisted_pp2_total: int | None = None,
) -> dict:
    matched = [
        row
        for row in rows
        if row.product_id is not None
        and row.status
        not in {STATUS_DUPLICATE_ARTICLE, STATUS_AMBIGUOUS_ARTICLE}
    ]
    source_total = 0
    existing_matched_total = 0
    net_delta = 0
    for row in rows:
        if row.qty is not None and row.status not in {
            STATUS_DUPLICATE_ARTICLE,
            STATUS_INVALID_ROW,
            STATUS_ERROR,
        }:
            source_total += int(row.qty)
        if row.quantity_before is not None:
            existing_matched_total += int(row.quantity_before)
        if row.status == STATUS_CHANGED and row.quantity_delta is not None:
            net_delta += int(row.quantity_delta)

    summary = {
        'source_rows': len(rows),
        'matched': len(matched),
        'unmatched': sum(1 for row in rows if row.status == STATUS_UNMATCHED),
        'source_total': source_total,
        'existing_pp2_total': int(existing_pp2_total),
        'existing_matched_total': existing_matched_total,
        'changed': sum(1 for row in rows if row.status == STATUS_CHANGED),
        'unchanged': sum(1 for row in rows if row.status == STATUS_UNCHANGED),
        'already_applied': sum(
            1 for row in rows if row.status == STATUS_ALREADY_APPLIED
        ),
        'missing_balance': sum(
            1 for row in rows if row.status == STATUS_MISSING_BALANCE
        ),
        'errors': sum(
            1
            for row in rows
            if row.status
            in {
                STATUS_ERROR,
                STATUS_DUPLICATE_ARTICLE,
                STATUS_AMBIGUOUS_ARTICLE,
                STATUS_INVALID_ROW,
            }
        ),
        'net_delta': net_delta,
        'result_total': int(existing_pp2_total) + net_delta,
        'persisted_pp2_total': persisted_pp2_total,
    }
    return summary


def apply_reconciliation_rows(
    rows: list[ReconciliationRow],
    *,
    warehouse: Warehouse,
    source: str,
    reference: str,
    note: str,
):
    with transaction.atomic():
        product_ids = [row.product_id for row in rows if row.product_id]
        locked = {
            stock.product_id: stock
            for stock in (
                ProductWarehouseStock.objects.select_for_update()
                .filter(warehouse=warehouse, product_id__in=product_ids)
                .order_by('pk')
            )
        }
        for row in rows:
            if row.status not in SAFE_APPLY_STATUSES:
                continue
            product = Product.objects.get(pk=row.product_id)
            stock = locked.get(row.product_id)
            if stock is None:
                raise ValueError(f'missing_locked_balance:{row.article}')
            if int(stock.quantity) != int(row.quantity_before):
                raise ValueError(f'balance_changed_during_apply:{row.article}')
            stock_qty_before = product.stock_qty
            price_before = product.price
            cost_before = product.cost_price
            status_before = product.status
            apply_stock_movement(
                product=product,
                warehouse=warehouse,
                movement_type=StockMovement.MovementType.ADJUSTMENT,
                quantity_delta=int(row.quantity_delta),
                source=source,
                reference=reference,
                note=note,
            )
            product.refresh_from_db(
                fields=['stock_qty', 'price', 'cost_price', 'status']
            )
            if (
                product.stock_qty != stock_qty_before
                or product.price != price_before
                or product.cost_price != cost_before
                or product.status != status_before
            ):
                raise ValueError('reconciler_changed_product_fields')


def reconcile_warehouse_inventory(
    *,
    warehouse_code: str,
    reference: str,
    note: str,
    apply: bool = False,
    path=None,
    mapping: dict[str, int] | None = None,
    source: str = SOURCE_RAPIDO,
) -> ReconciliationResult:
    if str(warehouse_code or '').strip() != ALLOWED_WAREHOUSE:
        raise ValueError(f'warehouse_must_be_{ALLOWED_WAREHOUSE}')
    try:
        warehouse = resolve_warehouse(warehouse_code)
    except StockServiceError as exc:
        raise ValueError(str(exc)) from exc

    sheet_name = ''
    headers: list = []
    source_path = ''
    if mapping is not None:
        rows = snapshot_rows_from_mapping(mapping)
        source_path = 'snapshot:rapido-2026-09-18'
    else:
        source_file = Path(path)
        if not source_file.exists():
            raise ValueError(f'file_not_found:{source_file}')
        rows, sheet_name, headers = snapshot_rows_from_file(source_file)
        source_path = str(source_file)

    existing_pp2_total = warehouse_quantity_total(warehouse)
    classify_reconciliation_rows(
        rows,
        warehouse=warehouse,
        source=source,
        reference=reference,
    )
    persisted_pp2_total = None
    if apply:
        apply_reconciliation_rows(
            rows,
            warehouse=warehouse,
            source=source,
            reference=reference,
            note=note,
        )
        persisted_pp2_total = warehouse_quantity_total(warehouse)

    return ReconciliationResult(
        rows=rows,
        warehouse_code=warehouse.code,
        warehouse_name=warehouse.name,
        reference=reference,
        source=source,
        note=note,
        source_path=source_path,
        apply=apply,
        sheet_name=sheet_name,
        headers=headers,
        summary=summarize_reconciliation_rows(
            rows,
            existing_pp2_total=existing_pp2_total,
            persisted_pp2_total=persisted_pp2_total,
        ),
    )


def pp1_total() -> int:
    return sum(
        ProductWarehouseStock.objects.filter(
            warehouse__code=WAREHOUSE_CODE_PP1
        ).values_list('quantity', flat=True)
    )


def pp2_total() -> int:
    return sum(
        ProductWarehouseStock.objects.filter(
            warehouse__code=WAREHOUSE_CODE_PP2
        ).values_list('quantity', flat=True)
    )
