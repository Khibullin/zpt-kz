"""Historical Kaspi Sales Report import. Read-only facts, never warehouse."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import transaction
from django.utils import timezone

from catalog.ag_parts_import import cell_text, normalize_header
from catalog.import_ops import sha256_file
from catalog.models import (
    CatalogImportBatch,
    KaspiOrder,
    KaspiSalesOperation,
    Product,
    ProductKaspiListing,
    SellerProfile,
)

STATUS_WOULD_CREATE = 'WOULD_CREATE'
STATUS_ALREADY_IMPORTED = 'ALREADY_IMPORTED'
STATUS_CREATED = 'CREATED'
STATUS_INVALID = 'INVALID'

OP_PURCHASE = KaspiSalesOperation.OperationType.PURCHASE
OP_RETURN = KaspiSalesOperation.OperationType.RETURN

MATCH_LISTING = KaspiSalesOperation.MatchStatus.LISTING_MATCHED
MATCH_PRODUCT = KaspiSalesOperation.MatchStatus.PRODUCT_ONLY
MATCH_UNMATCHED = KaspiSalesOperation.MatchStatus.UNMATCHED
MATCH_AMBIGUOUS = KaspiSalesOperation.MatchStatus.AMBIGUOUS_PRODUCT

OPERATION_TYPE_MAP = {
    'покупка': OP_PURCHASE,
    'возврат': OP_RETURN,
}

QTY_SUFFIX_RE = re.compile(r',\s*(\d+)\s*шт\.?\s*$', re.IGNORECASE)
TOKEN_SAFE_RE = re.compile(r'^[0-9A-Za-z._/-]{3,}$')

HEADER_ALIASES = {
    'order_id': ('номер заказа (id/rrn)', 'номер заказа'),
    'operation_type': ('тип операции',),
    'gross': ('сумма операции (т)', 'сумма операции (тг)', 'сумма операции'),
    'details': ('детали покупки',),
    'operation_date': ('дата операции',),
    'operation_time': ('время',),
    'accounting_date': ('дата учета операции', 'дата учёта операции'),
    'purchase_return_document': (
        '№ документа покупки/возврата',
        'no документа покупки/возврата',
        'номер документа покупки/возврата',
    ),
    'sales_point_id': ('идентификатор точки',),
    'terminal_id': ('id терминала',),
    'payment_type': ('тип оплаты',),
    'payment_type_2': ('тип оплаты 2',),
    'settlement': (
        'сумма к зачислению/ списанию (т)',
        'сумма к зачислению/списанию (т)',
        'сумма к зачислению/ списанию',
    ),
    'commission_amount': ('комиссия за операции (т)', 'комиссия за операции'),
    'commission_ex_vat_amount': (
        'комиссия за операции (т) без ндс',
        'комиссия за операции без ндс',
    ),
    'commission_ex_vat_percent': (
        'комиссия за операции (%) без ндс',
        'комиссия за операции % без ндс',
    ),
    'card_commission_amount': (
        'комиссия за операции по карте (т)',
        'комиссия за операции по карте',
    ),
    'card_commission_percent': (
        'комиссия за операции по карте (%)',
        'комиссия за операции по карте %',
    ),
    'payment_guarantee_amount': (
        'комиссия за обеспечение платежа (т)',
        'комиссия за обеспечение платежа',
    ),
    'payment_guarantee_percent': (
        'комиссия за обеспечение платежа (%)',
        'комиссия за обеспечение платежа %',
    ),
    'kaspi_pay_commission_amount': ('комиссия kaspi pay (т)', 'комиссия kaspi pay'),
    'kaspi_pay_commission_percent': ('комиссия kaspi pay (%)', 'комиссия kaspi pay %'),
    'kaspi_travel_commission_amount': (
        'комиссия kaspi travel (т)',
        'комиссия kaspi travel',
    ),
    'kaspi_travel_commission_percent': (
        'комиссия kaspi travel (%)',
        'комиссия kaspi travel %',
    ),
    'bonus_product_amount': (
        'оплата услуг за акцию бонусы на товар',
        'бонусы на товар',
    ),
    'bonus_review_amount': (
        'оплата услуг за акцию бонусы за отзыв',
        'бонусы за отзыв',
    ),
    'delivery_document': (
        '№ документа списания стоимости за kaspi доставку',
        'no документа списания стоимости за kaspi доставку',
    ),
    'delivery_cost': (
        'стоимость услуги за kaspi доставку',
        'стоимость kaspi доставки',
    ),
    'installment_term': ('срок для кредита на покупки',),
}

REQUIRED_FIELDS = ('order_id', 'operation_type', 'gross', 'details')
MONEY_FIELDS = (
    'gross',
    'settlement',
    'commission_amount',
    'commission_ex_vat_amount',
    'card_commission_amount',
    'payment_guarantee_amount',
    'kaspi_pay_commission_amount',
    'kaspi_travel_commission_amount',
    'bonus_product_amount',
    'bonus_review_amount',
    'delivery_cost',
)
PERCENT_FIELDS = (
    'commission_ex_vat_percent',
    'card_commission_percent',
    'payment_guarantee_percent',
    'kaspi_pay_commission_percent',
    'kaspi_travel_commission_percent',
)


@dataclass
class KaspiSalesRow:
    row_number: int
    raw: dict = field(default_factory=dict)
    external_order_id: str = ''
    operation_type: str = ''
    operation_at: object = None
    accounting_date: object = None
    purchase_return_document: str = ''
    sales_point_id: str = ''
    terminal_id: str = ''
    payment_type: str = ''
    payment_type_2: str = ''
    details: str = ''
    quantity: int | None = None
    gross_amount: Decimal | None = None
    settlement_amount: Decimal | None = None
    commission_amount: Decimal | None = None
    commission_ex_vat_amount: Decimal | None = None
    commission_ex_vat_percent: Decimal | None = None
    card_commission_amount: Decimal | None = None
    card_commission_percent: Decimal | None = None
    payment_guarantee_amount: Decimal | None = None
    payment_guarantee_percent: Decimal | None = None
    kaspi_pay_commission_amount: Decimal | None = None
    kaspi_pay_commission_percent: Decimal | None = None
    kaspi_travel_commission_amount: Decimal | None = None
    kaspi_travel_commission_percent: Decimal | None = None
    bonus_product_amount: Decimal | None = None
    bonus_review_amount: Decimal | None = None
    delivery_document: str = ''
    delivery_cost: Decimal | None = None
    installment_term: str = ''
    unit_gross_amount: Decimal | None = None
    fingerprint: str = ''
    product_id: int | None = None
    listing_id: int | None = None
    match_status: str = ''
    matched_identifier: str = ''
    status: str = ''
    reason: str = ''


@dataclass
class KaspiSalesSyncResult:
    rows: list[KaspiSalesRow]
    seller_id: int
    seller_name: str
    source_path: str
    source_sha256: str
    apply: bool
    sheet_name: str = ''
    header_row: int = 0
    headers: list = field(default_factory=list)
    column_map: dict = field(default_factory=dict)
    batch_id: int | None = None
    summary: dict = field(default_factory=dict)


def normalize_text(value) -> str:
    text = cell_text(value).replace('\xa0', ' ').replace('\u202f', ' ')
    return ' '.join(text.split())


def parse_kaspi_decimal(raw):
    """Return (Decimal|None, error). Empty is None, not 0. Bad text is error."""
    if raw is None:
        return None, ''
    if isinstance(raw, bool):
        return None, 'non_numeric'
    if isinstance(raw, int):
        return Decimal(raw), ''
    if isinstance(raw, Decimal):
        return raw, ''
    if isinstance(raw, float):
        if raw != raw:  # NaN
            return None, 'non_numeric'
        return Decimal(str(raw)), ''
    text = normalize_text(raw)
    if not text:
        return None, ''
    token = text.lower().replace(' ', '')
    if token in {'yes', 'no', 'preorder', 'true', 'false', '-', 'n/a', 'na', 'none'}:
        return None, 'non_numeric'
    negative = text.startswith('-')
    unsigned = text[1:] if negative else text
    unsigned = unsigned.replace(' ', '').replace('\xa0', '')
    if unsigned.count(',') == 1 and unsigned.count('.') == 0:
        unsigned = unsigned.replace(',', '.')
    elif unsigned.count(',') > 1:
        return None, 'non_numeric'
    try:
        amount = Decimal(unsigned)
    except (InvalidOperation, ValueError):
        return None, 'non_numeric'
    if negative:
        amount = -amount
    return amount, ''


def parse_kaspi_date(raw):
    if raw is None or raw == '':
        return None, 'empty'
    if isinstance(raw, datetime):
        return raw.date(), ''
    if isinstance(raw, date):
        return raw, ''
    text = normalize_text(raw)
    if not text:
        return None, 'empty'
    for fmt in ('%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, fmt).date(), ''
        except ValueError:
            continue
    return None, 'invalid_date'


def parse_kaspi_time(raw):
    if raw is None or raw == '':
        return None, 'empty'
    if isinstance(raw, time):
        return raw, ''
    if isinstance(raw, datetime):
        return raw.time(), ''
    text = normalize_text(raw)
    if not text:
        return None, 'empty'
    for fmt in ('%H:%M:%S', '%H:%M'):
        try:
            return datetime.strptime(text, fmt).time(), ''
        except ValueError:
            continue
    return None, 'invalid_time'


def combine_operation_at(date_value, time_value):
    naive = datetime.combine(date_value, time_value)
    current_tz = timezone.get_current_timezone()
    if timezone.is_naive(naive):
        return timezone.make_aware(naive, current_tz)
    return timezone.localtime(naive, current_tz)


def parse_details_quantity(details: str) -> int:
    match = QTY_SUFFIX_RE.search(details or '')
    if not match:
        return 1
    qty = int(match.group(1))
    return qty


def _format_decimal(value) -> str:
    if value is None:
        return ''
    return format(Decimal(value).quantize(Decimal('0.01')), 'f')


def operation_fingerprint(
    *,
    seller_id: int,
    external_order_id: str,
    operation_at,
    operation_type: str,
    purchase_return_document: str,
    details: str,
    gross_amount,
    settlement_amount,
) -> str:
    payload = '|'.join([
        str(seller_id),
        normalize_text(external_order_id),
        operation_at.isoformat(timespec='seconds') if operation_at is not None else '',
        operation_type,
        normalize_text(purchase_return_document),
        normalize_text(details),
        _format_decimal(gross_amount),
        _format_decimal(settlement_amount),
    ])
    return hashlib.sha256(payload.encode('utf-8')).hexdigest()


def detect_sales_columns(headers) -> dict[str, int]:
    normalized = [normalize_header(header).replace('\xa0', ' ') for header in headers]
    mapping = {}
    used = set()
    for field, aliases in HEADER_ALIASES.items():
        wanted = {normalize_header(alias) for alias in aliases}
        for index, header in enumerate(normalized):
            if index in used or not header:
                continue
            if header in wanted:
                mapping[field] = index
                used.add(index)
                break
    if 'payment_type' in mapping:
        # Prefer the first "тип оплаты"; payment_type_2 is a separate header.
        pass
    return mapping


def header_row_is_complete(column_map) -> bool:
    return all(field in column_map for field in REQUIRED_FIELDS)


def _detect_csv_dialect(sample: str):
    try:
        return csv.Sniffer().sniff(sample, delimiters=';,\t')
    except csv.Error:
        semicolon = sample.count(';')
        comma = sample.count(',')
        tab = sample.count('\t')
        delimiter = max(
            ((';', semicolon), (',', comma), ('\t', tab)),
            key=lambda item: item[1],
        )[0]
        dialect = csv.excel
        dialect.delimiter = delimiter
        return dialect


def _decode_csv_bytes(payload: bytes) -> str:
    for encoding in ('utf-8-sig', 'utf-8', 'cp1251'):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode('utf-8', errors='replace')


def _load_csv_rows(path: Path):
    payload = path.read_bytes()
    text = _decode_csv_bytes(payload)
    sample = text[:4096]
    dialect = _detect_csv_dialect(sample)
    reader = csv.reader(io.StringIO(text), dialect)
    rows = []
    for index, values in enumerate(reader, start=1):
        rows.append((index, [cell_text(value) for value in values]))
    return path.name, rows


def _load_xlsx_rows(path: Path, sheet=''):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet:
        worksheet = workbook[sheet]
        sheet_name = sheet
    else:
        worksheet = workbook.active
        sheet_name = worksheet.title
    rows = []
    for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = list(row)
        if not any(cell_text(value) for value in values):
            continue
        rows.append((index, values))
    workbook.close()
    return sheet_name, rows


def find_header_row(rows):
    for row_number, values in rows:
        headers = [cell_text(value) for value in values]
        column_map = detect_sales_columns(headers)
        if header_row_is_complete(column_map):
            return row_number, headers, column_map
    raise ValueError(
        'Не найдена строка заголовков Kaspi Sales Report '
        '(нужны «Номер заказа (ID/RRN)», «Тип операции», '
        '«Сумма операции (т)», «Детали покупки»).'
    )


def _cell(values, column_map, field):
    index = column_map.get(field)
    if index is None or index >= len(values):
        return None
    return values[index]


def build_identifier_index(seller: SellerProfile):
    products = list(
        Product.objects.filter(seller_profile=seller).only('id', 'article')
    )
    listings = list(
        ProductKaspiListing.objects.filter(product__seller_profile=seller)
        .only('id', 'product_id', 'master_sku', 'merchant_sku')
    )
    ident_to_products = defaultdict(set)
    listings_by_product = defaultdict(list)
    for product in products:
        article = normalize_text(product.article)
        if TOKEN_SAFE_RE.match(article):
            ident_to_products[article].add(product.pk)
    for listing in listings:
        listings_by_product[listing.product_id].append(listing)
        merchant = normalize_text(listing.merchant_sku)
        if TOKEN_SAFE_RE.match(merchant):
            ident_to_products[merchant].add(listing.product_id)
    identifiers = sorted(ident_to_products.keys(), key=len, reverse=True)
    return ident_to_products, listings_by_product, identifiers


def identifier_in_details(details: str, identifier: str) -> bool:
    if not identifier:
        return False
    pattern = re.compile(
        r'(?<![0-9A-Za-z._/-])'
        + re.escape(identifier)
        + r'(?![0-9A-Za-z._/-])',
    )
    return pattern.search(details) is not None


def match_product_and_listing(details, ident_to_products, listings_by_product, identifiers):
    hits = [
        ident for ident in identifiers
        if identifier_in_details(details, ident)
    ]
    product_ids = set()
    for ident in hits:
        product_ids.update(ident_to_products[ident])
    if not product_ids:
        return None, None, MATCH_UNMATCHED, ''
    if len(product_ids) > 1:
        return None, None, MATCH_AMBIGUOUS, ''
    product_id = next(iter(product_ids))
    matched_ident = hits[0] if hits else ''
    listings = listings_by_product.get(product_id, [])
    if len(listings) == 1:
        return product_id, listings[0].pk, MATCH_LISTING, matched_ident
    if not listings:
        return product_id, None, MATCH_PRODUCT, matched_ident
    listing_hits = []
    for listing in listings:
        keys = {
            normalize_text(listing.merchant_sku),
            normalize_text(listing.master_sku),
        }
        if any(ident in keys and ident for ident in hits):
            listing_hits.append(listing)
    unique = {item.pk: item for item in listing_hits}
    if len(unique) == 1:
        listing = next(iter(unique.values()))
        return product_id, listing.pk, MATCH_LISTING, matched_ident
    return product_id, None, MATCH_PRODUCT, matched_ident


def classify_sales_row(row: KaspiSalesRow, seller: SellerProfile, ident_index):
    ident_to_products, listings_by_product, identifiers = ident_index
    if not row.external_order_id:
        row.status = STATUS_INVALID
        row.reason = 'missing_order_id'
        return
    if not row.operation_type:
        row.status = STATUS_INVALID
        row.reason = 'invalid_operation_type'
        return
    if row.operation_at is None:
        row.status = STATUS_INVALID
        row.reason = 'invalid_operation_datetime'
        return
    if row.gross_amount is None:
        row.status = STATUS_INVALID
        row.reason = 'invalid_gross_amount'
        return
    if row.quantity is None or row.quantity <= 0:
        row.status = STATUS_INVALID
        row.reason = 'invalid_quantity'
        return
    if row.operation_type == OP_PURCHASE and row.gross_amount <= 0:
        row.status = STATUS_INVALID
        row.reason = 'invalid_sign'
        return
    if row.operation_type == OP_RETURN and row.gross_amount >= 0:
        row.status = STATUS_INVALID
        row.reason = 'invalid_sign'
        return
    row.unit_gross_amount = (
        abs(row.gross_amount) / Decimal(row.quantity)
    ).quantize(Decimal('0.01'))
    row.fingerprint = operation_fingerprint(
        seller_id=seller.pk,
        external_order_id=row.external_order_id,
        operation_at=row.operation_at,
        operation_type=row.operation_type,
        purchase_return_document=row.purchase_return_document,
        details=row.details,
        gross_amount=row.gross_amount,
        settlement_amount=row.settlement_amount,
    )
    product_id, listing_id, match_status, ident = match_product_and_listing(
        row.details,
        ident_to_products,
        listings_by_product,
        identifiers,
    )
    row.product_id = product_id
    row.listing_id = listing_id
    row.match_status = match_status
    row.matched_identifier = ident
    row.status = STATUS_WOULD_CREATE
    row.reason = ''


def parse_sales_rows(data_rows, column_map, _headers, seller):
    ident_index = build_identifier_index(seller)
    rows = []
    for row_number, values in data_rows:
        if not any(normalize_text(value) for value in values):
            continue
        order_id = normalize_text(_cell(values, column_map, 'order_id'))
        op_raw = normalize_text(_cell(values, column_map, 'operation_type'))
        details = normalize_text(_cell(values, column_map, 'details'))
        if (
            not order_id
            and not op_raw
            and not details
            and not normalize_text(_cell(values, column_map, 'gross'))
        ):
            continue
        raw = {}
        for field, index in column_map.items():
            if index < len(values):
                raw[field] = cell_text(values[index]) if not isinstance(
                    values[index], (int, float, Decimal)
                ) else values[index]
        row = KaspiSalesRow(row_number=row_number, raw=_json_safe(raw))
        row.external_order_id = normalize_text(_cell(values, column_map, 'order_id'))
        op_raw = normalize_text(_cell(values, column_map, 'operation_type')).lower()
        row.operation_type = OPERATION_TYPE_MAP.get(op_raw, '')
        date_raw = _cell(values, column_map, 'operation_date')
        time_raw = _cell(values, column_map, 'operation_time')
        if isinstance(date_raw, datetime) and not normalize_text(time_raw):
            aware = date_raw
            if timezone.is_naive(aware):
                aware = timezone.make_aware(aware, timezone.get_current_timezone())
            row.operation_at = aware
        else:
            date_value, date_error = parse_kaspi_date(date_raw)
            time_value, time_error = parse_kaspi_time(time_raw)
            if date_error or time_error:
                row.operation_at = None
            else:
                row.operation_at = combine_operation_at(date_value, time_value)
        acc, acc_error = parse_kaspi_date(_cell(values, column_map, 'accounting_date'))
        row.accounting_date = acc if not acc_error else None
        row.purchase_return_document = normalize_text(
            _cell(values, column_map, 'purchase_return_document')
        )
        row.sales_point_id = normalize_text(_cell(values, column_map, 'sales_point_id'))
        row.terminal_id = normalize_text(_cell(values, column_map, 'terminal_id'))
        row.payment_type = normalize_text(_cell(values, column_map, 'payment_type'))
        row.payment_type_2 = normalize_text(_cell(values, column_map, 'payment_type_2'))
        row.details = normalize_text(_cell(values, column_map, 'details'))
        row.quantity = parse_details_quantity(row.details)
        for field in MONEY_FIELDS + PERCENT_FIELDS:
            attr = {
                'gross': 'gross_amount',
                'settlement': 'settlement_amount',
            }.get(field, field)
            amount, error = parse_kaspi_decimal(_cell(values, column_map, field))
            if field == 'gross' and error:
                row.gross_amount = None
            elif error:
                setattr(row, attr, None)
            else:
                setattr(row, attr, amount)
        row.delivery_document = normalize_text(
            _cell(values, column_map, 'delivery_document')
        )
        row.installment_term = normalize_text(
            _cell(values, column_map, 'installment_term')
        )
        classify_sales_row(row, seller, ident_index)
        rows.append(row)
    return rows


def _json_safe(raw: dict) -> dict:
    safe = {}
    for key, value in raw.items():
        if isinstance(value, Decimal):
            safe[key] = format(value, 'f')
        elif isinstance(value, (datetime,)):
            safe[key] = value.isoformat()
        else:
            safe[key] = cell_text(value) if value is not None else ''
    return safe


def mark_existing_operations(rows, seller: SellerProfile):
    fingerprints = [row.fingerprint for row in rows if row.fingerprint]
    existing = set(
        KaspiSalesOperation.objects.filter(
            seller_profile=seller,
            source_fingerprint__in=fingerprints,
        ).values_list('source_fingerprint', flat=True)
    )
    for row in rows:
        if row.status != STATUS_WOULD_CREATE:
            continue
        if row.fingerprint in existing:
            row.status = STATUS_ALREADY_IMPORTED
            row.reason = 'already_imported'


def summarize_kaspi_sales_rows(rows, seller: SellerProfile):
    existing_orders = set(
        KaspiOrder.objects.filter(seller_profile=seller).values_list(
            'external_order_id', flat=True
        )
    )
    valid = [
        row for row in rows
        if row.status in {STATUS_WOULD_CREATE, STATUS_ALREADY_IMPORTED, STATUS_CREATED}
    ]
    purchases = [row for row in valid if row.operation_type == OP_PURCHASE]
    returns = [row for row in valid if row.operation_type == OP_RETURN]
    would_orders = {
        row.external_order_id
        for row in rows
        if row.status == STATUS_WOULD_CREATE
        and row.external_order_id not in existing_orders
    }
    existing_order_hits = {
        row.external_order_id
        for row in valid
        if row.external_order_id in existing_orders
    }
    def _sum(items, attr):
        total = Decimal('0')
        for item in items:
            value = getattr(item, attr)
            if value is not None:
                total += value
        return total

    return {
        'total_rows': len(rows),
        'purchases': len(purchases),
        'returns': len(returns),
        'valid_rows': len(valid),
        'invalid_rows': sum(1 for row in rows if row.status == STATUS_INVALID),
        'matched_listing': sum(1 for row in valid if row.match_status == MATCH_LISTING),
        'product_only': sum(1 for row in valid if row.match_status == MATCH_PRODUCT),
        'unmatched_product': sum(1 for row in valid if row.match_status == MATCH_UNMATCHED),
        'ambiguous_product': sum(1 for row in valid if row.match_status == MATCH_AMBIGUOUS),
        'would_create_orders': len(would_orders),
        'existing_orders': len(existing_order_hits),
        'would_create_operations': sum(
            1 for row in rows if row.status == STATUS_WOULD_CREATE
        ),
        'already_imported': sum(
            1 for row in rows if row.status == STATUS_ALREADY_IMPORTED
        ),
        'created_operations': sum(1 for row in rows if row.status == STATUS_CREATED),
        'purchase_qty': sum(row.quantity or 0 for row in purchases),
        'return_qty': sum(row.quantity or 0 for row in returns),
        'purchase_gross': _sum(purchases, 'gross_amount'),
        'return_gross': _sum(returns, 'gross_amount'),
        'commission_total': _sum(valid, 'commission_amount'),
        'delivery_total': _sum(valid, 'delivery_cost'),
    }


def apply_kaspi_sales_rows(rows, *, seller, source_path, source_sha256, started_at):
    now = timezone.now()
    created_ops = 0
    created_orders = 0
    unchanged = 0
    skipped = 0
    with transaction.atomic():
        batch = CatalogImportBatch.objects.create(
            seller_profile=seller,
            source=CatalogImportBatch.SOURCE_KASPI_SALES_REPORT,
            filename=Path(source_path).name,
            file_sha256=source_sha256,
            started_at=started_at,
            finished_at=now,
            mode=CatalogImportBatch.MODE_WRITE,
            source_scope=CatalogImportBatch.SCOPE_PARTIAL,
            status=CatalogImportBatch.STATUS_SUCCESS,
            source_row_count=len(rows),
        )
        for row in rows:
            if row.status == STATUS_INVALID:
                skipped += 1
                continue
            if row.status not in {STATUS_WOULD_CREATE, STATUS_ALREADY_IMPORTED}:
                continue
            order, order_created = KaspiOrder.objects.get_or_create(
                seller_profile=seller,
                external_order_id=row.external_order_id,
            )
            if order_created:
                created_orders += 1
            changed = []
            if (
                order.first_operation_at is None
                or row.operation_at < order.first_operation_at
            ):
                order.first_operation_at = row.operation_at
                changed.append('first_operation_at')
            if (
                order.last_operation_at is None
                or row.operation_at > order.last_operation_at
            ):
                order.last_operation_at = row.operation_at
                changed.append('last_operation_at')
            if changed:
                order.save(update_fields=changed + ['updated_at'])
            if row.status == STATUS_ALREADY_IMPORTED:
                unchanged += 1
                continue
            KaspiSalesOperation.objects.create(
                seller_profile=seller,
                order=order,
                product_id=row.product_id,
                listing_id=row.listing_id,
                purchase_return_document=row.purchase_return_document,
                sales_point_id=row.sales_point_id,
                terminal_id=row.terminal_id,
                operation_type=row.operation_type,
                operation_at=row.operation_at,
                accounting_date=row.accounting_date,
                payment_type=row.payment_type,
                payment_type_2=row.payment_type_2,
                gross_amount=row.gross_amount,
                settlement_amount=row.settlement_amount,
                commission_amount=row.commission_amount,
                commission_ex_vat_amount=row.commission_ex_vat_amount,
                commission_ex_vat_percent=row.commission_ex_vat_percent,
                card_commission_amount=row.card_commission_amount,
                card_commission_percent=row.card_commission_percent,
                payment_guarantee_amount=row.payment_guarantee_amount,
                payment_guarantee_percent=row.payment_guarantee_percent,
                kaspi_pay_commission_amount=row.kaspi_pay_commission_amount,
                kaspi_pay_commission_percent=row.kaspi_pay_commission_percent,
                kaspi_travel_commission_amount=row.kaspi_travel_commission_amount,
                kaspi_travel_commission_percent=row.kaspi_travel_commission_percent,
                bonus_product_amount=row.bonus_product_amount,
                bonus_review_amount=row.bonus_review_amount,
                delivery_document=row.delivery_document,
                delivery_cost=row.delivery_cost,
                installment_term=row.installment_term,
                details=row.details,
                quantity=row.quantity,
                unit_gross_amount=row.unit_gross_amount,
                match_status=row.match_status,
                matched_identifier=row.matched_identifier,
                source_filename=Path(source_path).name,
                source_sha256=source_sha256,
                source_row=row.row_number,
                source_fingerprint=row.fingerprint,
                import_batch=batch,
                raw_data=row.raw,
            )
            row.status = STATUS_CREATED
            created_ops += 1
        batch.source_unique_count = len({
            row.external_order_id for row in rows if row.external_order_id
        })
        batch.selected_count = created_ops + unchanged
        batch.created_count = created_ops
        batch.unchanged_count = unchanged
        batch.skipped_count = skipped
        batch.conflict_count = 0
        batch.finished_at = timezone.now()
        batch.save(update_fields=[
            'source_unique_count',
            'selected_count',
            'created_count',
            'unchanged_count',
            'skipped_count',
            'conflict_count',
            'finished_at',
        ])
    return batch


def load_kaspi_sales_table(path: Path, sheet=''):
    suffix = path.suffix.lower()
    if suffix == '.csv':
        sheet_name, rows = _load_csv_rows(path)
    elif suffix in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        sheet_name, rows = _load_xlsx_rows(path, sheet=sheet)
    else:
        raise ValueError(f'unsupported_file_type:{suffix or "missing"}')
    if not rows:
        raise ValueError('empty_sales_report')
    header_row, headers, column_map = find_header_row(rows)
    data_rows = [
        (row_number, values)
        for row_number, values in rows
        if row_number > header_row
    ]
    return sheet_name, header_row, headers, column_map, data_rows


def sync_kaspi_sales_report(
    *,
    path,
    seller: SellerProfile,
    apply: bool = False,
    sheet='',
):
    source_path = Path(path)
    started_at = timezone.now()
    digest = sha256_file(source_path)
    sheet_name, header_row, headers, column_map, data_rows = load_kaspi_sales_table(
        source_path,
        sheet=sheet,
    )
    rows = parse_sales_rows(data_rows, column_map, headers, seller)
    mark_existing_operations(rows, seller)
    summary = summarize_kaspi_sales_rows(rows, seller)
    batch = None
    if apply:
        batch = apply_kaspi_sales_rows(
            rows,
            seller=seller,
            source_path=str(source_path),
            source_sha256=digest,
            started_at=started_at,
        )
    return KaspiSalesSyncResult(
        rows=rows,
        seller_id=seller.pk,
        seller_name=seller.name,
        source_path=str(source_path),
        source_sha256=digest,
        apply=apply,
        sheet_name=sheet_name,
        header_row=header_row,
        headers=list(headers),
        column_map=dict(column_map),
        batch_id=batch.pk if batch is not None else None,
        summary=summary,
    )
