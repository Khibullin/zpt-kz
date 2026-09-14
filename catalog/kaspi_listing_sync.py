"""Match seller products to Kaspi listings from XLSX/CSV.

Dry-run is the default. Never creates Product, never rebinds master SKU
between products, never deletes listings or barcodes, never touches price,
stock, publish flags, or repricer rules.

Match key: SellerProfile + normalize_article(Product.article).
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction

from catalog.ag_parts_import import (
    cell_text,
    extract_article,
    load_sheet_rows,
    normalize_article,
    normalize_header,
)
from catalog.barcode_sync import BARCODE_HEADERS, split_barcode_cell
from catalog.models import Product, ProductBarcode, ProductKaspiListing, SellerProfile

STATUS_MATCHED = 'MATCHED'
STATUS_WOULD_CREATE_LISTING = 'WOULD_CREATE_LISTING'
STATUS_WOULD_UPDATE_LISTING = 'WOULD_UPDATE_LISTING'
STATUS_MISSING_PRODUCT = 'MISSING_PRODUCT'
STATUS_MISSING_KASPI_ID = 'MISSING_KASPI_ID'
STATUS_DUPLICATE_INPUT_ARTICLE = 'DUPLICATE_INPUT_ARTICLE'
STATUS_CONFLICT_MASTER_SKU = 'CONFLICT_MASTER_SKU'
STATUS_CONFLICT_PRODUCT = 'CONFLICT_PRODUCT'
STATUS_CONFLICT_MERCHANT_SKU = 'CONFLICT_MERCHANT_SKU'
STATUS_BARCODE_CONFLICT = 'BARCODE_CONFLICT'
STATUS_MULTIPLE_MASTER_SKU = 'MULTIPLE_MASTER_SKU'
STATUS_INVALID_ROW = 'INVALID_ROW'

SAFE_APPLY_STATUSES = frozenset({
    STATUS_WOULD_CREATE_LISTING,
    STATUS_WOULD_UPDATE_LISTING,
})

PROBLEM_STATUSES = frozenset({
    STATUS_MISSING_PRODUCT,
    STATUS_MISSING_KASPI_ID,
    STATUS_DUPLICATE_INPUT_ARTICLE,
    STATUS_CONFLICT_MASTER_SKU,
    STATUS_CONFLICT_PRODUCT,
    STATUS_CONFLICT_MERCHANT_SKU,
    STATUS_BARCODE_CONFLICT,
    STATUS_MULTIPLE_MASTER_SKU,
    STATUS_INVALID_ROW,
})

REPORT_COLUMNS = (
    'row_number',
    'article',
    'product_id',
    'product_name',
    'master_sku',
    'merchant_sku',
    'barcode',
    'existing_listing_id',
    'status',
    'reason',
)

# Detect master SKU before generic "sku" so Kaspi IDs are not eaten as article.
MASTER_SKU_ALIASES = (
    'master_sku',
    'kaspi master sku',
    'kaspi_sku',
    'kaspi sku',
    'kaspi_id',
    'kaspi id',
    'kaspi product id',
    'kaspi product_id',
    'код товара kaspi',
    'код kaspi',
    'kaspi код товара',
    'kaspi код',
)
ARTICLE_ALIASES_STRICT = (
    'артикул',
    'article',
)
SKU_HEADER = 'sku'
MISSING_ARTICLE_ERROR = (
    'Не найдена колонка артикула (article/артикул). '
    'Голый заголовок SKU не считается артикулом: в выгрузке Kaspi это Kaspi ID. '
    'Добавьте колонку article/артикул или явно укажите --map-sku=article, '
    'только если SKU в этом файле действительно внутренний артикул.'
)
MAP_SKU_ARTICLE = 'article'
MAP_SKU_MASTER = 'master_sku'
MAP_SKU_CHOICES = (MAP_SKU_ARTICLE, MAP_SKU_MASTER)
_MASTER_SKU_SPLIT = re.compile(r'[|;]')
MERCHANT_SKU_ALIASES = (
    'merchant_sku',
    'merchant sku',
    'merchant code',
    'merchant_code',
    'артикул продавца',
    'seller sku',
    'kaspi merchant sku',
    'kaspi merchant_sku',
)


@dataclass
class KaspiLinkRow:
    row_number: int
    article: str = ''
    article_key: str = ''
    master_sku: str = ''
    merchant_sku: str = ''
    barcode: str = ''
    barcodes: list[str] = field(default_factory=list)
    master_skus: list[str] = field(default_factory=list)
    product_id: int | None = None
    product_name: str = ''
    existing_listing_id: int | None = None
    status: str = ''
    reason: str = ''


@dataclass
class KaspiLinkSyncResult:
    rows: list[KaspiLinkRow]
    seller_id: int
    seller_name: str
    source_path: str
    apply: bool
    sheet_name: str = ''


def detect_kaspi_link_columns(headers, map_sku: str = '') -> dict[str, int]:
    """Map logical fields to header indexes. Never guess by column position.

    Bare ``SKU`` is never article unless ``map_sku='article'`` is explicit.
    With an article column already mapped, leftover ``SKU`` is Kaspi master SKU.
    """
    if map_sku and map_sku not in MAP_SKU_CHOICES:
        raise ValueError(f'unsupported_map_sku:{map_sku}')
    normalized = [normalize_header(header) for header in headers]
    used = set()
    mapping = {}

    def take(field, aliases):
        for alias in aliases:
            alias_key = normalize_header(alias)
            for index, header in enumerate(normalized):
                if index in used or not header:
                    continue
                if header == alias_key:
                    mapping[field] = index
                    used.add(index)
                    return

    take('master_sku', MASTER_SKU_ALIASES)
    take('merchant_sku', MERCHANT_SKU_ALIASES)
    take('barcode', BARCODE_HEADERS)
    take('article', ARTICLE_ALIASES_STRICT)

    sku_indexes = [
        index
        for index, header in enumerate(normalized)
        if index not in used and header == SKU_HEADER
    ]
    if sku_indexes:
        sku_index = sku_indexes[0]
        if 'article' in mapping and 'master_sku' not in mapping:
            mapping['master_sku'] = sku_index
            used.add(sku_index)
        elif map_sku == MAP_SKU_ARTICLE and 'article' not in mapping:
            mapping['article'] = sku_index
            used.add(sku_index)
        elif map_sku == MAP_SKU_MASTER and 'master_sku' not in mapping:
            mapping['master_sku'] = sku_index
            used.add(sku_index)

    return mapping


def split_master_skus(raw) -> list[str]:
    """Split a Kaspi master SKU cell. Never pick the first ID silently."""
    text = cell_text(raw)
    if not text:
        return []
    parts = []
    seen = set()
    for piece in _MASTER_SKU_SPLIT.split(text):
        sku = cell_text(piece)
        if not sku or sku in seen:
            continue
        seen.add(sku)
        parts.append(sku)
    return parts


def _cell_at(values, column_map, field):
    index = column_map.get(field)
    if index is None or index >= len(values):
        return ''
    return values[index]


def _parse_barcodes(raw) -> list[str]:
    codes = []
    seen = set()
    for part in split_barcode_cell(cell_text(raw)):
        code = cell_text(part)
        if not code or code in seen:
            continue
        seen.add(code)
        codes.append(code)
    return codes


def _load_xlsx_tables(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet_names = list(workbook.sheetnames)
    workbook.close()
    tables = []
    for sheet_name in sheet_names:
        headers, _column_map, data_rows, _images = load_sheet_rows(path, sheet_name)
        tables.append((sheet_name, headers, data_rows))
    return tables


def _load_csv_table(path: Path):
    with path.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = csv.reader(handle)
        try:
            header_row = next(reader)
        except StopIteration:
            return path.name, [], []
        headers = [cell_text(value) for value in header_row]
        data_rows = []
        for index, row in enumerate(reader, start=2):
            values = list(row)
            if not any(cell_text(value) for value in values):
                continue
            data_rows.append((index, values))
        return path.name, headers, data_rows


def load_kaspi_link_rows(path: Path, map_sku: str = ''):
    suffix = path.suffix.lower()
    if suffix == '.csv':
        sheet_name, headers, data_rows = _load_csv_table(path)
        column_map = detect_kaspi_link_columns(headers, map_sku=map_sku)
        return sheet_name, headers, column_map, data_rows
    if suffix not in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        raise ValueError(f'unsupported_file_type:{suffix or "missing"}')

    chosen = None
    fallback = None
    for sheet_name, headers, data_rows in _load_xlsx_tables(path):
        column_map = detect_kaspi_link_columns(headers, map_sku=map_sku)
        candidate = (sheet_name, headers, column_map, data_rows)
        if fallback is None:
            fallback = candidate
        if 'article' in column_map:
            chosen = candidate
            break
    selected = chosen or fallback
    if selected is None:
        raise ValueError('workbook_has_no_sheets')
    return selected


def _require_article_column(headers, column_map):
    if 'article' in column_map:
        return
    sku_present = any(normalize_header(header) == SKU_HEADER for header in headers)
    if sku_present:
        raise ValueError(MISSING_ARTICLE_ERROR)
    raise ValueError(f'Не найдена колонка артикула. Заголовки: {headers}')


def _parse_input_rows(headers, column_map, data_rows) -> list[KaspiLinkRow]:
    rows = []
    for row_number, values in data_rows:
        raw_article = _cell_at(values, column_map, 'article')
        article, article_key = extract_article(raw_article)
        barcodes = _parse_barcodes(_cell_at(values, column_map, 'barcode'))
        master_skus = split_master_skus(_cell_at(values, column_map, 'master_sku'))
        row = KaspiLinkRow(
            row_number=row_number,
            article=article or cell_text(raw_article),
            article_key=article_key,
            master_sku=', '.join(master_skus),
            master_skus=master_skus,
            merchant_sku=cell_text(_cell_at(values, column_map, 'merchant_sku')),
            barcode=barcodes[0] if barcodes else '',
            barcodes=barcodes,
        )
        if len(master_skus) > 1:
            row.status = STATUS_MULTIPLE_MASTER_SKU
            row.reason = f"multiple_master_sku: {', '.join(master_skus)}"
        elif len(master_skus) == 1:
            row.master_sku = master_skus[0]
        rows.append(row)
    return rows


def _build_product_indexes(seller: SellerProfile):
    exact = {}
    by_key = {}
    queryset = Product.objects.filter(seller_profile=seller).exclude(article='')
    for product in queryset:
        exact.setdefault(product.article, []).append(product)
        key = normalize_article(product.article)
        if key:
            by_key.setdefault(key, []).append(product)
    return exact, by_key


def _find_seller_product(row: KaspiLinkRow, exact, by_key):
    exact_matches = exact.get(row.article) or []
    if len(exact_matches) > 1:
        return None, 'ambiguous_article_for_seller'
    if len(exact_matches) == 1:
        return exact_matches[0], None

    keyed = list({item.pk: item for item in (by_key.get(row.article_key) or [])}.values())
    if len(keyed) > 1:
        return None, 'ambiguous_normalized_article_for_seller'
    if len(keyed) == 1:
        return keyed[0], None
    return None, 'product_not_found'


def _listing_master_index(seller: SellerProfile):
    by_product = {}
    by_master = {}
    queryset = ProductKaspiListing.objects.select_related('product').filter(
        product__seller_profile=seller,
    )
    for listing in queryset:
        by_product.setdefault(listing.product_id, []).append(listing)
        by_master.setdefault(listing.master_sku, []).append(listing)
    return by_product, by_master


def _barcode_index(product_ids):
    by_product = {}
    if not product_ids:
        return by_product
    for barcode in ProductBarcode.objects.filter(product_id__in=product_ids):
        by_product.setdefault(barcode.product_id, set()).add(barcode.code)
    return by_product


def _classify_linked_row(
    row: KaspiLinkRow,
    product: Product,
    listings,
    existing_codes,
    master_owners: dict[str, set[int]],
):
    row.product_id = product.pk
    row.product_name = product.title
    listings = list(listings or [])

    other_owners = {
        product_id
        for product_id in master_owners.get(row.master_sku, set())
        if product_id != product.pk
    }
    if other_owners:
        row.status = STATUS_CONFLICT_MASTER_SKU
        row.reason = 'master_sku_bound_to_other_product'
        return

    matched_listing = next(
        (item for item in listings if item.master_sku == row.master_sku),
        None,
    )
    if matched_listing is None and len(listings) == 1:
        matched_listing = listings[0]
        if matched_listing.master_sku != row.master_sku:
            other_for_new = master_owners.get(row.master_sku, set()) - {product.pk}
            if other_for_new:
                row.status = STATUS_CONFLICT_MASTER_SKU
                row.reason = 'master_sku_bound_to_other_product'
                row.existing_listing_id = matched_listing.pk
                return
    if matched_listing is None and len(listings) > 1:
        row.status = STATUS_CONFLICT_PRODUCT
        row.reason = 'product_has_incompatible_listings'
        return

    barcode_would_add = bool(
        row.barcodes and any(code not in existing_codes for code in row.barcodes)
    )

    if matched_listing is None:
        row.status = STATUS_WOULD_CREATE_LISTING
        row.reason = 'create_listing'
        return

    row.existing_listing_id = matched_listing.pk
    listing_changes = []
    if matched_listing.master_sku != row.master_sku:
        listing_changes.append('master_sku')
    if row.merchant_sku and matched_listing.merchant_sku != row.merchant_sku:
        listing_changes.append('merchant_sku')
    if row.barcode and matched_listing.barcode != row.barcode:
        listing_changes.append('listing_barcode')
    if barcode_would_add:
        listing_changes.append('product_barcode')

    if not listing_changes:
        row.status = STATUS_MATCHED
        row.reason = 'listing_matches'
        return

    row.status = STATUS_WOULD_UPDATE_LISTING
    row.reason = 'update:' + ','.join(listing_changes)


def classify_kaspi_link_rows(rows: list[KaspiLinkRow], seller: SellerProfile):
    exact, by_key = _build_product_indexes(seller)
    listings_by_product, listings_by_master = _listing_master_index(seller)

    article_counts = {}
    master_to_keys = {}
    for row in rows:
        if row.article_key:
            article_counts[row.article_key] = article_counts.get(row.article_key, 0) + 1
        if (
            row.article_key
            and row.status != STATUS_MULTIPLE_MASTER_SKU
            and len(row.master_skus) == 1
        ):
            master_to_keys.setdefault(row.master_skus[0], set()).add(row.article_key)

    product_ids = set()
    for row in rows:
        if not row.article_key:
            row.status = STATUS_INVALID_ROW
            row.reason = 'empty_article'
            continue
        if article_counts.get(row.article_key, 0) > 1:
            row.status = STATUS_DUPLICATE_INPUT_ARTICLE
            row.reason = 'duplicate_input_article'
            continue
        if row.status == STATUS_MULTIPLE_MASTER_SKU:
            product, error = _find_seller_product(row, exact, by_key)
            if error is None:
                row.product_id = product.pk
                row.product_name = product.title
                product_ids.add(product.pk)
            continue
        product, error = _find_seller_product(row, exact, by_key)
        if error == 'product_not_found':
            row.status = STATUS_MISSING_PRODUCT
            row.reason = 'product_not_found'
            continue
        if error:
            row.status = STATUS_CONFLICT_PRODUCT
            row.reason = error
            continue
        row.product_id = product.pk
        row.product_name = product.title
        product_ids.add(product.pk)
        if not row.master_sku:
            row.status = STATUS_MISSING_KASPI_ID
            row.reason = 'missing_kaspi_id'
            continue
        if len(master_to_keys.get(row.master_sku, set())) > 1:
            row.status = STATUS_CONFLICT_MASTER_SKU
            row.reason = 'master_sku_maps_to_multiple_articles'
            continue

    barcode_by_product = _barcode_index(product_ids)
    master_owners = {
        master_sku: {item.product_id for item in listings}
        for master_sku, listings in listings_by_master.items()
    }
    for row in rows:
        if row.status:
            continue
        product = None
        if row.product_id:
            exact_hits = exact.get(row.article) or []
            product = next((item for item in exact_hits if item.pk == row.product_id), None)
            if product is None:
                keyed = by_key.get(row.article_key) or []
                product = next((item for item in keyed if item.pk == row.product_id), None)
        if product is None:
            row.status = STATUS_INVALID_ROW
            row.reason = 'empty_article' if not row.article_key else 'product_not_found'
            continue
        _classify_linked_row(
            row,
            product,
            listings_by_product.get(product.pk),
            barcode_by_product.get(product.pk, set()),
            master_owners,
        )
    return rows


def _ensure_kaspi_barcodes(product: Product, codes: list[str]):
    if not codes:
        return
    has_primary = product.barcodes.filter(is_primary=True).exists()
    for code in codes:
        _obj, created = ProductBarcode.objects.get_or_create(
            product=product,
            code=code,
            defaults={
                'source': ProductBarcode.SOURCE_KASPI,
                'is_primary': not has_primary,
            },
        )
        if created:
            has_primary = True


def apply_kaspi_link_rows(rows: list[KaspiLinkRow]):
    """Write only WOULD_CREATE / WOULD_UPDATE rows. Conflict/missing skipped."""
    with transaction.atomic():
        for row in rows:
            if row.status not in SAFE_APPLY_STATUSES:
                continue
            product = Product.objects.get(pk=row.product_id)
            if row.status == STATUS_WOULD_CREATE_LISTING:
                listing = ProductKaspiListing.objects.create(
                    product=product,
                    master_sku=row.master_sku,
                    merchant_sku=row.merchant_sku,
                    barcode=row.barcode,
                )
                row.existing_listing_id = listing.pk
            else:
                listing = ProductKaspiListing.objects.select_for_update().get(
                    pk=row.existing_listing_id,
                )
                update_fields = []
                if row.master_sku and listing.master_sku != row.master_sku:
                    listing.master_sku = row.master_sku
                    update_fields.append('master_sku')
                if row.merchant_sku and listing.merchant_sku != row.merchant_sku:
                    listing.merchant_sku = row.merchant_sku
                    update_fields.append('merchant_sku')
                if row.barcode and listing.barcode != row.barcode:
                    listing.barcode = row.barcode
                    update_fields.append('barcode')
                if update_fields:
                    update_fields.append('updated_at')
                    listing.save(update_fields=update_fields)
            _ensure_kaspi_barcodes(product, row.barcodes)


def summarize_kaspi_link_rows(rows: list[KaspiLinkRow]) -> dict[str, int]:
    counts = {
        'total_rows': len(rows),
        'valid_rows': 0,
        'matched': 0,
        'would_create_listings': 0,
        'would_update_listings': 0,
        'missing_products': 0,
        'missing_kaspi_ids': 0,
        'duplicate_input_articles': 0,
        'master_sku_conflicts': 0,
        'product_conflicts': 0,
        'merchant_sku_conflicts': 0,
        'barcode_conflicts': 0,
        'multiple_master_skus': 0,
        'invalid_rows': 0,
    }
    for row in rows:
        if row.status != STATUS_INVALID_ROW:
            counts['valid_rows'] += 1
        if row.status == STATUS_MATCHED:
            counts['matched'] += 1
        elif row.status == STATUS_WOULD_CREATE_LISTING:
            counts['would_create_listings'] += 1
        elif row.status == STATUS_WOULD_UPDATE_LISTING:
            counts['would_update_listings'] += 1
        elif row.status == STATUS_MISSING_PRODUCT:
            counts['missing_products'] += 1
        elif row.status == STATUS_MISSING_KASPI_ID:
            counts['missing_kaspi_ids'] += 1
        elif row.status == STATUS_DUPLICATE_INPUT_ARTICLE:
            counts['duplicate_input_articles'] += 1
        elif row.status == STATUS_CONFLICT_MASTER_SKU:
            counts['master_sku_conflicts'] += 1
        elif row.status == STATUS_CONFLICT_PRODUCT:
            counts['product_conflicts'] += 1
        elif row.status == STATUS_CONFLICT_MERCHANT_SKU:
            counts['merchant_sku_conflicts'] += 1
        elif row.status == STATUS_BARCODE_CONFLICT:
            counts['barcode_conflicts'] += 1
        elif row.status == STATUS_MULTIPLE_MASTER_SKU:
            counts['multiple_master_skus'] += 1
        elif row.status == STATUS_INVALID_ROW:
            counts['invalid_rows'] += 1
    return counts


def write_kaspi_link_report(rows: list[KaspiLinkRow], report_path: Path):
    path = Path(report_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix in {'.xlsx', '.xlsm'}:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'kaspi_links'
        sheet.append(list(REPORT_COLUMNS))
        for row in rows:
            sheet.append([
                row.row_number,
                row.article,
                row.product_id or '',
                row.product_name,
                row.master_sku,
                row.merchant_sku,
                row.barcode,
                row.existing_listing_id or '',
                row.status,
                row.reason,
            ])
        workbook.save(path)
        return path
    if not suffix:
        path = path.with_suffix('.csv')
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(REPORT_COLUMNS))
        writer.writeheader()
        for row in rows:
            writer.writerow({
                'row_number': row.row_number,
                'article': row.article,
                'product_id': row.product_id or '',
                'product_name': row.product_name,
                'master_sku': row.master_sku,
                'merchant_sku': row.merchant_sku,
                'barcode': row.barcode,
                'existing_listing_id': row.existing_listing_id or '',
                'status': row.status,
                'reason': row.reason,
            })
    return path


def inspect_kaspi_link_file(path, map_sku: str = ''):
    """Parse XLSX/CSV without Product lookup. Used for file-level audits."""
    source = Path(path)
    sheet_name, headers, column_map, data_rows = load_kaspi_link_rows(
        source,
        map_sku=map_sku,
    )
    _require_article_column(headers, column_map)
    rows = _parse_input_rows(headers, column_map, data_rows)
    return sheet_name, headers, column_map, rows


def sync_kaspi_product_links(
    *,
    path,
    seller: SellerProfile,
    apply: bool = False,
    map_sku: str = '',
):
    source = Path(path)
    sheet_name, headers, column_map, data_rows = load_kaspi_link_rows(
        source,
        map_sku=map_sku,
    )
    _require_article_column(headers, column_map)
    rows = _parse_input_rows(headers, column_map, data_rows)
    classify_kaspi_link_rows(rows, seller)
    if apply:
        apply_kaspi_link_rows(rows)
    return KaspiLinkSyncResult(
        rows=rows,
        seller_id=seller.pk,
        seller_name=seller.name,
        source_path=str(source),
        apply=apply,
        sheet_name=sheet_name,
    )
