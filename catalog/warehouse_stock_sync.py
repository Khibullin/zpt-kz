"""Import PP1/PP2 warehouse balances from XLSX/CSV.

Dry-run is the default. Never creates Product, never writes Product.stock_qty,
never sends stock to Kaspi, never touches price or repricer.

Match key: SellerProfile + normalize_article(Product.article).
Kaspi available stock is PP2 only; this importer never sums PP1+PP2.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction

from catalog.ag_parts_import import (
    cell_text,
    load_sheet_rows,
    normalize_article,
    normalize_header,
)
from catalog.models import Product, ProductWarehouseStock, SellerProfile, Warehouse
from catalog.stock_service import set_stock_quantity
from catalog.warehouses import (
    DEFAULT_WAREHOUSES,
    WAREHOUSE_CODE_PP1,
    WAREHOUSE_CODE_PP2,
)

STATUS_MATCHED_NO_CHANGE = 'MATCHED_NO_CHANGE'
STATUS_WOULD_CREATE_BALANCE = 'WOULD_CREATE_BALANCE'
STATUS_WOULD_ADJUST = 'WOULD_ADJUST'
STATUS_MISSING_PRODUCT = 'MISSING_PRODUCT'
STATUS_INVALID_QUANTITY = 'INVALID_QUANTITY'
STATUS_DUPLICATE_ARTICLE = 'DUPLICATE_ARTICLE'
STATUS_INVALID_ROW = 'INVALID_ROW'

SAFE_APPLY_STATUSES = frozenset({
    STATUS_WOULD_CREATE_BALANCE,
    STATUS_WOULD_ADJUST,
})

PROBLEM_STATUSES = frozenset({
    STATUS_MISSING_PRODUCT,
    STATUS_INVALID_QUANTITY,
    STATUS_DUPLICATE_ARTICLE,
    STATUS_INVALID_ROW,
})

REPORT_COLUMNS = (
    'row_number',
    'article',
    'product_id',
    'old_pp1',
    'new_pp1',
    'delta_pp1',
    'old_pp2',
    'new_pp2',
    'delta_pp2',
    'status',
    'reason',
)

ARTICLE_ALIASES = (
    'артикул',
    'article',
    'продукт',
)
QUANTITY_ALIASES = (
    'qty',
    'quantity',
    'остаток',
    'кол-во',
    'количество',
    'stock',
)
PP1_ALIASES = ('pp1', 'pp 1')
PP2_ALIASES = ('pp2', 'pp 2')
NON_NUMERIC_TOKENS = frozenset({
    'no',
    'yes',
    'preorder',
    'true',
    'false',
    'none',
    'null',
    '-',
    'n/a',
    'na',
})
KASPI_PP_WARNING = (
    'Запись идёт в склады PP1 (Основной склад) и/или PP2 (Fulfillment). '
    'Если источник — выгрузка Kaspi merchant, PP1/PP2 там пункты выдачи, '
    'а не склады ZPT. Нечисловые значения (no, yes, preorder, пусто) не '
    'превращаются в 0, а дают INVALID_QUANTITY. Для --apply в PP1/PP2 '
    'нужен флаг --confirm-pp-source.'
)
SOURCE_IMPORT = 'warehouse_import'
MISSING_ARTICLE_ERROR = (
    'Не найдена колонка артикула (article/артикул/продукт). '
    'Голый SKU не считается артикулом.'
)
MISSING_QTY_ERROR = (
    'Не найдены колонки остатков. Нужны PP1 и PP2 вместе, '
    'либо --warehouse-code и колонка количества.'
)


@dataclass
class WarehouseQtyPlan:
    warehouse_code: str
    raw: object = None
    new_qty: int | None = None
    old_qty: int | None = None
    delta: int | None = None
    has_row: bool = False
    error: str = ''
    active: bool = False


@dataclass
class WarehouseStockRow:
    row_number: int
    article: str = ''
    article_key: str = ''
    product_id: int | None = None
    product_name: str = ''
    pp1: WarehouseQtyPlan = field(
        default_factory=lambda: WarehouseQtyPlan(WAREHOUSE_CODE_PP1)
    )
    pp2: WarehouseQtyPlan = field(
        default_factory=lambda: WarehouseQtyPlan(WAREHOUSE_CODE_PP2)
    )
    extra: list[WarehouseQtyPlan] = field(default_factory=list)
    status: str = ''
    reason: str = ''

    def plans(self) -> list[WarehouseQtyPlan]:
        plans = []
        if self.pp1.new_qty is not None or self.pp1.error or self.pp1.raw not in (None, ''):
            plans.append(self.pp1)
        if self.pp2.new_qty is not None or self.pp2.error or self.pp2.raw not in (None, ''):
            plans.append(self.pp2)
        plans.extend(self.extra)
        if plans:
            return plans
        return [item for item in (self.pp1, self.pp2) if item.warehouse_code]


@dataclass
class WarehouseStockSyncResult:
    rows: list[WarehouseStockRow]
    seller_id: int
    seller_name: str
    source_path: str
    apply: bool
    sheet_name: str = ''
    headers: list = field(default_factory=list)
    dual_pp_columns: bool = False
    warehouse_codes: list[str] = field(default_factory=list)
    samples: dict = field(default_factory=dict)


def parse_stock_quantity(raw):
    """Return (int_qty, error_code). Only unambiguous non-negative integers."""
    if raw is None:
        return None, 'empty'
    if isinstance(raw, bool):
        return None, 'non_numeric'
    if isinstance(raw, int):
        if raw < 0:
            return None, 'negative'
        return raw, ''
    if isinstance(raw, float):
        if raw != raw:  # NaN
            return None, 'non_numeric'
        if not raw.is_integer():
            return None, 'not_integer'
        qty = int(raw)
        if qty < 0:
            return None, 'negative'
        return qty, ''
    text = cell_text(raw)
    if not text:
        return None, 'empty'
    token = text.lower()
    if token in NON_NUMERIC_TOKENS:
        return None, 'non_numeric'
    if text.startswith('-') and text[1:].isdigit():
        return None, 'negative'
    if text.isdigit():
        return int(text), ''
    return None, 'non_numeric'


def _header_index(headers, aliases):
    normalized = [normalize_header(header) for header in headers]
    wanted = {normalize_header(alias) for alias in aliases}
    for index, header in enumerate(normalized):
        if header in wanted:
            return index
    return None


def detect_warehouse_stock_columns(
    headers,
    *,
    warehouse_code='',
    quantity_column='',
) -> dict[str, int]:
    mapping = {}
    article_idx = _header_index(headers, ARTICLE_ALIASES)
    if article_idx is not None:
        mapping['article'] = article_idx
    pp1_idx = _header_index(headers, PP1_ALIASES)
    pp2_idx = _header_index(headers, PP2_ALIASES)
    if pp1_idx is not None:
        mapping['pp1'] = pp1_idx
    if pp2_idx is not None:
        mapping['pp2'] = pp2_idx

    if warehouse_code:
        if quantity_column:
            qty_idx = _header_index(headers, (quantity_column,))
        else:
            qty_idx = _header_index(headers, QUANTITY_ALIASES)
            if qty_idx is None and warehouse_code == WAREHOUSE_CODE_PP1:
                qty_idx = mapping.get('pp1')
            if qty_idx is None and warehouse_code == WAREHOUSE_CODE_PP2:
                qty_idx = mapping.get('pp2')
        if qty_idx is not None:
            mapping['quantity'] = qty_idx
            mapping['warehouse_code'] = warehouse_code
    return mapping


def _load_xlsx_tables(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet_names = list(workbook.sheetnames)
    workbook.close()
    tables = []
    for sheet_name in sheet_names:
        headers, _column_map, data_rows, _images = load_sheet_rows(path, sheet_name)
        tables.append((sheet_name, headers, data_rows))
    return tables


def _load_csv_table(path: Path):
    with path.open('r', encoding='utf-8-sig', newline='') as handle:
        reader = csv.reader(handle)
        try:
            header_row = next(reader)
        except StopIteration:
            return path.name, [], []
        headers = [cell_text(value) for value in header_row]
        data_rows = []
        for index, row in enumerate(reader, start=2):
            values = list(row)
            if not any(cell_text(value) for value in values):
                continue
            data_rows.append((index, values))
        return path.name, headers, data_rows


def load_warehouse_stock_rows(path: Path, *, warehouse_code='', quantity_column=''):
    suffix = path.suffix.lower()
    if suffix == '.csv':
        sheet_name, headers, data_rows = _load_csv_table(path)
        column_map = detect_warehouse_stock_columns(
            headers,
            warehouse_code=warehouse_code,
            quantity_column=quantity_column,
        )
        return sheet_name, headers, column_map, data_rows
    if suffix not in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
        raise ValueError(f'unsupported_file_type:{suffix or "missing"}')

    chosen = None
    fallback = None
    for sheet_name, headers, data_rows in _load_xlsx_tables(path):
        column_map = detect_warehouse_stock_columns(
            headers,
            warehouse_code=warehouse_code,
            quantity_column=quantity_column,
        )
        candidate = (sheet_name, headers, column_map, data_rows)
        if fallback is None:
            fallback = candidate
        if 'article' in column_map:
            chosen = candidate
            break
    selected = chosen or fallback
    if selected is None:
        raise ValueError('workbook_has_no_sheets')
    return selected


def _require_columns(headers, column_map, *, warehouse_code=''):
    if 'article' not in column_map:
        raise ValueError(MISSING_ARTICLE_ERROR)
    dual = 'pp1' in column_map and 'pp2' in column_map
    single = 'quantity' in column_map and bool(warehouse_code)
    if warehouse_code and dual and 'quantity' not in column_map:
        raise ValueError(
            'Файл содержит PP1 и PP2. Не указывайте --warehouse-code: '
            'оба склада импортируются одним запуском.'
        )
    if not dual and not single:
        raise ValueError(f'{MISSING_QTY_ERROR} Заголовки: {headers}')
    return dual


def _cell_at(values, column_map, field):
    index = column_map.get(field)
    if index is None or index >= len(values):
        return None
    return values[index]


def _sample_values(data_rows, column_map, field, limit=8):
    samples = []
    seen = set()
    for _row_number, values in data_rows:
        raw = _cell_at(values, column_map, field)
        text = cell_text(raw) if raw is not None else ''
        if text in seen:
            continue
        seen.add(text)
        samples.append(text if text else '(empty)')
        if len(samples) >= limit:
            break
    return samples


def _fill_plan(plan: WarehouseQtyPlan, raw, product, warehouses):
    plan.raw = raw
    qty, error = parse_stock_quantity(raw)
    if error:
        plan.error = error
        plan.new_qty = None
        plan.delta = None
        return
    plan.new_qty = qty
    warehouse = warehouses.get(plan.warehouse_code)
    if warehouse is None:
        plan.error = f'warehouse_missing:{plan.warehouse_code}'
        return
    if product is None:
        plan.old_qty = None
        plan.delta = None
        plan.has_row = False
        return
    existing = ProductWarehouseStock.objects.filter(
        product=product,
        warehouse=warehouse,
    ).first()
    plan.has_row = existing is not None
    plan.old_qty = int(existing.quantity) if existing is not None else 0
    plan.delta = plan.new_qty - plan.old_qty


def _build_product_indexes(seller: SellerProfile):
    exact = {}
    by_key = {}
    queryset = Product.objects.filter(seller_profile=seller).exclude(article='')
    for product in queryset:
        exact.setdefault(product.article, []).append(product)
        key = normalize_article(product.article)
        if key:
            by_key.setdefault(key, []).append(product)
    return exact, by_key


def _find_seller_product(article, article_key, exact, by_key):
    exact_matches = exact.get(article) or []
    if len(exact_matches) > 1:
        return None, 'ambiguous_article_for_seller'
    if len(exact_matches) == 1:
        return exact_matches[0], None
    keyed = list({item.pk: item for item in (by_key.get(article_key) or [])}.values())
    if len(keyed) > 1:
        return None, 'ambiguous_normalized_article_for_seller'
    if len(keyed) == 1:
        return keyed[0], None
    return None, 'product_not_found'


def _row_status(plans: list[WarehouseQtyPlan]) -> tuple[str, str]:
    if any(plan.error for plan in plans):
        errors = [plan.error for plan in plans if plan.error]
        return STATUS_INVALID_QUANTITY, ','.join(errors)
    changing = [
        plan for plan in plans
        if (not plan.has_row) or plan.delta
    ]
    if not changing:
        return STATUS_MATCHED_NO_CHANGE, 'no_change'
    creating = [plan for plan in changing if not plan.has_row]
    adjusting = [plan for plan in changing if plan.has_row]
    if creating and not adjusting:
        return STATUS_WOULD_CREATE_BALANCE, 'create_balance'
    return STATUS_WOULD_ADJUST, 'adjust_balance'


def classify_warehouse_stock_rows(rows: list[WarehouseStockRow], seller: SellerProfile):
    exact, by_key = _build_product_indexes(seller)
    warehouses = {item.code: item for item in Warehouse.objects.filter(is_active=True)}
    article_counts = {}
    for row in rows:
        if row.article_key:
            article_counts[row.article_key] = article_counts.get(row.article_key, 0) + 1

    for row in rows:
        if not row.article_key:
            row.status = STATUS_INVALID_ROW
            row.reason = 'empty_article'
            continue
        product, error = _find_seller_product(
            row.article,
            row.article_key,
            exact,
            by_key,
        )
        if product is not None:
            row.product_id = product.pk
            row.product_name = product.title
        plans = [
            plan for plan in (row.pp1, row.pp2, *row.extra)
            if plan.active
        ]
        for plan in plans:
            _fill_plan(plan, plan.raw, product, warehouses)
        if article_counts.get(row.article_key, 0) > 1:
            row.status = STATUS_DUPLICATE_ARTICLE
            row.reason = 'duplicate_input_article'
            continue
        if any(plan.error for plan in plans):
            row.status, row.reason = _row_status(plans)
            continue
        if product is None:
            row.status = STATUS_MISSING_PRODUCT
            row.reason = error or 'product_not_found'
            continue
        row.status, row.reason = _row_status(plans)
    return rows


def apply_warehouse_stock_rows(rows: list[WarehouseStockRow], *, source_path=''):
    reference = Path(source_path).name if source_path else ''
    with transaction.atomic():
        for row in rows:
            if row.status not in SAFE_APPLY_STATUSES:
                continue
            product = Product.objects.get(pk=row.product_id)
            stock_qty_before = product.stock_qty
            for plan in (row.pp1, row.pp2, *row.extra):
                if plan.new_qty is None or plan.error:
                    continue
                if plan.has_row and not plan.delta:
                    continue
                set_stock_quantity(
                    product=product,
                    warehouse=plan.warehouse_code,
                    new_quantity=plan.new_qty,
                    source=SOURCE_IMPORT,
                    reference=reference,
                    note=f'row={row.row_number} article={row.article}',
                )
            product.refresh_from_db(fields=['stock_qty'])
            if product.stock_qty != stock_qty_before:
                raise ValueError('importer_changed_product_stock_qty')


def summarize_warehouse_stock_rows(rows: list[WarehouseStockRow]) -> dict[str, int]:
    counts = {
        'total_rows': len(rows),
        'matched_no_change': 0,
        'would_create_balance': 0,
        'would_adjust': 0,
        'missing_products': 0,
        'invalid_quantity': 0,
        'duplicate_article': 0,
        'invalid_rows': 0,
    }
    for row in rows:
        if row.status == STATUS_MATCHED_NO_CHANGE:
            counts['matched_no_change'] += 1
        elif row.status == STATUS_WOULD_CREATE_BALANCE:
            counts['would_create_balance'] += 1
        elif row.status == STATUS_WOULD_ADJUST:
            counts['would_adjust'] += 1
        elif row.status == STATUS_MISSING_PRODUCT:
            counts['missing_products'] += 1
        elif row.status == STATUS_INVALID_QUANTITY:
            counts['invalid_quantity'] += 1
        elif row.status == STATUS_DUPLICATE_ARTICLE:
            counts['duplicate_article'] += 1
        elif row.status == STATUS_INVALID_ROW:
            counts['invalid_rows'] += 1
    return counts


def write_warehouse_stock_report(rows: list[WarehouseStockRow], report_path: Path):
    path = Path(report_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()
    if suffix in {'.xlsx', '.xlsm'}:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'warehouse_stocks'
        sheet.append(list(REPORT_COLUMNS))
        for row in rows:
            sheet.append([
                row.row_number,
                row.article,
                row.product_id or '',
                row.pp1.old_qty if row.pp1.old_qty is not None else '',
                row.pp1.new_qty if row.pp1.new_qty is not None else '',
                row.pp1.delta if row.pp1.delta is not None else '',
                row.pp2.old_qty if row.pp2.old_qty is not None else '',
                row.pp2.new_qty if row.pp2.new_qty is not None else '',
                row.pp2.delta if row.pp2.delta is not None else '',
                row.status,
                row.reason,
            ])
        workbook.save(path)
        return path
    if not suffix:
        path = path.with_suffix('.csv')
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(REPORT_COLUMNS))
        writer.writeheader()
        for row in rows:
            writer.writerow({
                'row_number': row.row_number,
                'article': row.article,
                'product_id': row.product_id or '',
                'old_pp1': row.pp1.old_qty if row.pp1.old_qty is not None else '',
                'new_pp1': row.pp1.new_qty if row.pp1.new_qty is not None else '',
                'delta_pp1': row.pp1.delta if row.pp1.delta is not None else '',
                'old_pp2': row.pp2.old_qty if row.pp2.old_qty is not None else '',
                'new_pp2': row.pp2.new_qty if row.pp2.new_qty is not None else '',
                'delta_pp2': row.pp2.delta if row.pp2.delta is not None else '',
                'status': row.status,
                'reason': row.reason,
            })
    return path


def _parse_input_rows(column_map, data_rows, *, warehouse_code='', dual=False):
    rows = []
    for row_number, values in data_rows:
        raw_article = _cell_at(values, column_map, 'article')
        article = cell_text(raw_article)
        row = WarehouseStockRow(
            row_number=row_number,
            article=article,
            article_key=normalize_article(article),
        )
        if dual:
            row.pp1.raw = _cell_at(values, column_map, 'pp1')
            row.pp1.active = True
            row.pp2.raw = _cell_at(values, column_map, 'pp2')
            row.pp2.active = True
        elif warehouse_code:
            plan = WarehouseQtyPlan(warehouse_code)
            plan.raw = _cell_at(values, column_map, 'quantity')
            plan.active = True
            if warehouse_code == WAREHOUSE_CODE_PP1:
                row.pp1 = plan
            elif warehouse_code == WAREHOUSE_CODE_PP2:
                row.pp2 = plan
            else:
                row.extra.append(plan)
        rows.append(row)
    return rows


def sync_warehouse_stocks(
    *,
    path,
    seller: SellerProfile,
    apply: bool = False,
    warehouse_code='',
    quantity_column='',
    confirm_pp_source: bool = False,
):
    source = Path(path)
    sheet_name, headers, column_map, data_rows = load_warehouse_stock_rows(
        source,
        warehouse_code=warehouse_code,
        quantity_column=quantity_column,
    )
    dual = _require_columns(headers, column_map, warehouse_code=warehouse_code)
    writes_pp = dual or warehouse_code in {
        WAREHOUSE_CODE_PP1,
        WAREHOUSE_CODE_PP2,
    }
    if apply and writes_pp and not confirm_pp_source:
        raise ValueError(
            'PP1/PP2 apply requires --confirm-pp-source. '
            + KASPI_PP_WARNING
        )
    if warehouse_code and not Warehouse.objects.filter(
        code=warehouse_code, is_active=True
    ).exists():
        raise ValueError(f'warehouse_not_found:{warehouse_code}')

    rows = _parse_input_rows(
        column_map,
        data_rows,
        warehouse_code=warehouse_code,
        dual=dual,
    )
    classify_warehouse_stock_rows(rows, seller)
    if apply:
        apply_warehouse_stock_rows(rows, source_path=str(source))

    samples = {}
    if 'pp1' in column_map:
        samples['PP1'] = _sample_values(data_rows, column_map, 'pp1')
    if 'pp2' in column_map:
        samples['PP2'] = _sample_values(data_rows, column_map, 'pp2')
    if 'quantity' in column_map:
        samples['quantity'] = _sample_values(data_rows, column_map, 'quantity')

    codes = []
    if dual:
        codes = [WAREHOUSE_CODE_PP1, WAREHOUSE_CODE_PP2]
    elif warehouse_code:
        codes = [warehouse_code]

    return WarehouseStockSyncResult(
        rows=rows,
        seller_id=seller.pk,
        seller_name=seller.name,
        source_path=str(source),
        apply=apply,
        sheet_name=sheet_name,
        headers=list(headers),
        dual_pp_columns=dual,
        warehouse_codes=codes,
        samples=samples,
    )


def ensure_default_warehouses():
    """Idempotent PP1/PP2 seed used by tests and data migration."""
    created = []
    for code, name in DEFAULT_WAREHOUSES:
        warehouse, was_created = Warehouse.objects.get_or_create(
            code=code,
            defaults={'name': name, 'is_active': True},
        )
        if was_created:
            created.append(warehouse)
    return created
