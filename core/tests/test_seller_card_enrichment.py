from datetime import time
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from core.admin import import_seller_row, parse_sellers_xlsx
from core.models import (
    Seller,
    SellerExternalCard,
    SellerExtraPhone,
    SellerSocialLink,
    SellerWorkInterval,
)
from core.services.seller_card_display import build_seller_card, route_links
from core.services.seller_card_enrichment import (
    analyze_organizations,
    apply_plans,
    enrich_files,
    names_compatible,
    parse_mapsgrabs_xlsx,
    parse_organization_row,
    parse_schedule,
)
from core.services.seller_card_links import assess_whatsapp_link, clean_http_url, validate_card_url


def _seller(**kwargs):
    payload = {
        'name': 'Авторазбор',
        'whatsapp': '77010000001',
        'transport_type': 'car',
        'city': 'Актобе',
    }
    payload.update(kwargs)
    return Seller.objects.create(**payload)


def _org(**kwargs):
    row = {
        'source': '2ГИС',
        'external_id': '7001',
        'name': 'Barys Auto',
        'address': 'улица Абая, 10',
        'city': 'Алматы',
        'coordinates': '43.250000, 76.950000',
        'phones': '+77011112233, +77012223344',
        'websites': 'http://barys.example',
        'socials': 'https://instagram.com/barys_auto, https://wa.me/77011112233',
        'schedule': 'Пн: 09:00–13:00, 14:00–18:00; Вт: 09:00–18:00',
        'card_url': 'https://2gis.ru/firm/7001',
    }
    row.update(kwargs)
    return parse_organization_row(row, exported_at='2026-10-14T12:27:32.699Z', file_name='test.xlsx', row_number=2)


def _mapsgrabs_bytes(rows, source='2ГИС', exported_at='2026-10-17T04:28:15.568Z'):
    workbook = Workbook()
    meta = workbook.active
    meta.title = 'О выгрузке'
    meta.append(['Параметр', 'Значение'])
    meta.append(['Источник', source])
    meta.append(['Дата выгрузки', exported_at])
    sheet = workbook.create_sheet('Организации')
    sheet.append([
        'Источник', 'ID организации', 'Название', 'Адрес', 'Город', 'Координаты',
        'Контакты', 'Сайты', 'Соцсети', 'Режим работы', 'Ссылка на карточку',
        'Рейтинг', 'Отзывов',
    ])
    for row in rows:
        sheet.append(row)
    workbook.create_sheet('Отзывы').append(['Оценка', 'Текст'])
    payload = BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return payload


class ScheduleAndLinkTests(TestCase):
    def test_schedule_formats(self):
        daily, error = parse_schedule(
            'Пн: 09:00–13:00, 14:00–18:00; Вс: выходной'
        )
        self.assertEqual(error, '')
        self.assertEqual(daily[0]['opens'], '09:00')
        self.assertEqual(daily[1]['closes'], '18:00')
        self.assertTrue(daily[-1]['is_closed'])

        compact, error = parse_schedule('пн-пт 09:00–18:00; сб 10:00–18:00')
        self.assertEqual(error, '')
        self.assertEqual(len(compact), 6)
        self.assertNotIn(6, [item['weekday'] for item in compact])

        every, error = parse_schedule('ежедневно, 09:00–18:00')
        self.assertEqual(error, '')
        self.assertEqual(len(every), 7)

        empty, error = parse_schedule('')
        self.assertEqual(empty, [])
        self.assertEqual(error, '')

        full_day, error = parse_schedule('Пн: 00:00–24:00')
        self.assertEqual(error, '')
        self.assertEqual(full_day[0]['closes'], '23:59')
        overnight, error = parse_schedule('Пн: 10:00–01:00')
        self.assertEqual(error, '')
        self.assertEqual(overnight[0]['opens'], '10:00')
        self.assertEqual(overnight[0]['closes'], '01:00')

    def test_rejects_unsafe_and_mismatched_links(self):
        url, error = clean_http_url('javascript:alert(1)')
        self.assertEqual(url, '')
        self.assertIn('небезопасная', error)
        url, error = clean_http_url('viber://chat/?number=77716710101')
        self.assertEqual(url, '')

        cleaned, error = validate_card_url('two_gis', 'https://yandex.ru/maps/org/1')
        self.assertEqual(cleaned, '')
        self.assertIn('не соответствует', error)
        cleaned, error = validate_card_url('yandex', 'https://yandex.ru/maps/?text=авторазбор')
        self.assertEqual(cleaned, '')

        stored, digits, problem = assess_whatsapp_link(
            'https://wa.me/277075562423',
            {'77075562423'},
        )
        self.assertEqual(stored, '')
        self.assertEqual(digits, '')
        self.assertIn('подозрительный', problem)

    def test_mapsgrabs_row_keeps_contacts_and_skips_reviews(self):
        org = _org(
            websites='http://www.tiktok.com/@shah.korea777?is_from_webapp=1, http://70000001052211627.2gis.biz/',
            socials='https://instagram.com/korea.zapchasti, https://wa.me/77716710101, viber://chat/?number=77716710101',
            phones='+77716710101',
            card_url='https://2gis.ru/firm/70000001058677613',
            external_id='70000001058677613',
        )
        kinds = {item['kind'] for item in org['socials']}
        self.assertIn('tiktok', kinds)
        self.assertIn('instagram', kinds)
        self.assertIn('whatsapp', kinds)
        self.assertEqual(org['website'], '')
        self.assertTrue(any('viber' in item['url'] or 'небезопасная' in item['reason'] for item in org['invalid']))
        self.assertTrue(any(item['reason'].startswith('ссылка ведёт на карту') for item in org['invalid']))
        self.assertNotIn('рейтинг', org)


class MatchingTests(TestCase):
    def test_unique_phone_matches_and_shared_phone_does_not_merge(self):
        first = _seller(name='Филиал 1', whatsapp='77020000001', phone2='77020000009')
        second = _seller(name='Филиал 2', whatsapp='77020000002', phone2='77020000009')
        unique = _org(phones='+77020000001', name='Совсем другое имя', address='другая улица, 99', external_id='1')
        shared = _org(phones='+77020000009', name='Ещё одно имя', address='третья улица, 3', external_id='2')
        report = analyze_organizations([unique, shared])
        self.assertEqual(report['matched'][0]['seller_id'], first.id)
        self.assertEqual(len(report['ambiguous']), 1)
        self.assertIn('филиал', report['ambiguous'][0]['reasons'][0])
        self.assertNotEqual(report['matched'][0]['seller_id'], second.id)

    def test_similar_name_without_address_is_not_enough(self):
        _seller(name='Barys Auto', whatsapp='77030000001', city='Алматы', address='улица Абая, 10')
        org = _org(phones='', name='Barys Auto', address='проспект Достык, 20', external_id='3')
        report = analyze_organizations([org])
        self.assertEqual(report['matched'], [])
        self.assertEqual(len(report['unmatched']), 1)
        self.assertFalse(names_compatible('Авто', 'Авторазбор Север'))

    def test_name_and_address_match(self):
        seller = _seller(
            name='Barys Auto, центр авторазбора',
            whatsapp='77030000008',
            city='Алматы',
            address='',
            market_location='ул. Абая, 10',
        )
        org = _org(phones='', external_id='4')
        report = analyze_organizations([org])
        self.assertEqual(report['matched'][0]['seller_id'], seller.id)

    def test_external_id_in_notes_matches_without_replacing_other_source(self):
        seller = _seller(
            name='Другое название',
            whatsapp='77040000001',
            notes='Карточка https://2gis.ru/firm/7001 и телефон магазина',
        )
        yandex = _org(
            source='Яндекс Карты',
            external_id='555',
            phones='+77040000001',
            card_url='https://yandex.ru/maps/org/555',
            socials='',
            websites='',
        )
        report = analyze_organizations([_org(phones=''), yandex])
        self.assertEqual({row['seller_id'] for row in report['matched']}, {seller.id})
        apply_plans(report['plans'])
        seller.refresh_from_db()
        sources = set(seller.external_cards.values_list('source', 'card_url'))
        self.assertIn(('two_gis', 'https://2gis.ru/firm/7001'), sources)
        self.assertIn(('yandex', 'https://yandex.ru/maps/org/555'), sources)
        self.assertEqual(seller.notes, 'Карточка https://2gis.ru/firm/7001 и телефон магазина')
        self.assertEqual(seller.whatsapp, '77040000001')


class ApplyRulesTests(TestCase):
    def test_conflict_keeps_filled_values_and_primary_phone(self):
        seller = _seller(
            name='Motor center',
            whatsapp='8 702 470 05 55',
            phone2='',
            address='просп. Кенеса Нокина, 18Л',
            website='https://old.example',
            notes='Исходное примечание',
            receive_requests=True,
            is_active=True,
            is_paused=False,
        )
        org = _org(
            source='Яндекс Карты',
            external_id='1208969566',
            name='Motor center',
            address='Актобе, просп. Кенеса Нокина, 99',
            city='Актобе',
            coordinates='50.306148, 57.206707',
            phones='+77024700555, +77764700555',
            websites='',
            socials='https://instagram.com/motor_center_kz',
            schedule='пн-пт 09:00–18:00',
            card_url='https://yandex.ru/maps/org/1208969566',
        )
        report = analyze_organizations([org])
        self.assertEqual(report['matched'][0]['seller_id'], seller.id)
        self.assertTrue(any(item['field'] == 'адрес' for item in report['conflicts']))
        apply_plans(report['plans'])
        seller.refresh_from_db()
        self.assertEqual(seller.whatsapp, '8 702 470 05 55')
        self.assertEqual(seller.address, 'просп. Кенеса Нокина, 18Л')
        self.assertEqual(seller.website, 'https://old.example')
        self.assertEqual(seller.notes, 'Исходное примечание')
        self.assertTrue(seller.receive_requests)
        self.assertTrue(seller.is_active)
        self.assertFalse(seller.is_paused)
        self.assertEqual(seller.phone2, '77764700555')
        self.assertEqual(seller.latitude is not None, True)
        self.assertEqual(SellerWorkInterval.objects.filter(seller=seller).count(), 5)

    def test_empty_values_do_not_clear_and_reimport_has_no_duplicates(self):
        seller = _seller(
            name='Barys Auto',
            whatsapp='77011112233',
            city='Алматы',
            website='https://keep.example',
        )
        org = _org(websites='', phones='+77011112233, +77012223344')
        report = analyze_organizations([org])
        apply_plans(report['plans'])
        apply_plans(analyze_organizations([org])['plans'])
        seller.refresh_from_db()
        self.assertEqual(seller.website, 'https://keep.example')
        self.assertEqual(seller.whatsapp, '77011112233')
        self.assertEqual(seller.extra_phones.count(), 0)
        self.assertEqual(seller.phone2, '77012223344')
        self.assertEqual(seller.external_cards.count(), 1)
        self.assertEqual(seller.social_links.filter(kind='instagram').count(), 1)
        self.assertEqual(seller.work_hours.count(), 3)

    def test_does_not_create_sellers_for_unknown_rows(self):
        before = Seller.objects.count()
        report = analyze_organizations([_org()])
        apply_plans(report['plans'])
        self.assertEqual(Seller.objects.count(), before)
        self.assertEqual(report['plans'], [])

    def test_legacy_seller_import_keeps_card_fields(self):
        seller = _seller(whatsapp='77015550011', address='улица Абая, 10', receive_requests=True)
        SellerExternalCard.objects.create(
            seller=seller,
            source='two_gis',
            external_id='7001',
            card_url='https://2gis.ru/firm/7001',
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['seller_name', 'whatsapp', 'transport_type', 'city'])
        sheet.append(['Новое имя', '77015550011', 'car', 'Алматы'])
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        rows, errors = parse_sellers_xlsx(payload)
        self.assertEqual(errors, [])
        import_seller_row(rows[0])
        seller.refresh_from_db()
        self.assertEqual(seller.address, 'улица Абая, 10')
        self.assertEqual(seller.external_cards.count(), 1)
        self.assertFalse(seller.receive_requests)


class CardPageTests(TestCase):
    def test_full_and_partial_cards(self):
        full = _seller(
            name='Полная карточка',
            is_dismantler=True,
            whatsapp='77011112233',
            phone2='77012223344',
            city='Алматы',
            address='улица Абая, дом 10, очень длинный адрес который не должен разъезжать страницу',
            market_location='Бутик 4',
            website='https://barys.example',
            notes='Оригинальное описание',
            latitude='43.250000',
            longitude='76.950000',
        )
        SellerExtraPhone.objects.create(seller=full, phone='77013334455')
        SellerSocialLink.objects.create(seller=full, kind='instagram', url='https://instagram.com/barys_auto')
        SellerSocialLink.objects.create(seller=full, kind='telegram', url='https://t.me/barys_auto')
        SellerSocialLink.objects.create(seller=full, kind='tiktok', url='https://tiktok.com/@barys')
        SellerSocialLink.objects.create(seller=full, kind='vk', url='https://vk.com/barys')
        SellerSocialLink.objects.create(seller=full, kind='youtube', url='https://youtube.com/@barys')
        SellerWorkInterval.objects.create(seller=full, weekday=0, opens=time(9, 0), closes=time(13, 0), position=0)
        SellerWorkInterval.objects.create(seller=full, weekday=0, opens=time(14, 0), closes=time(18, 0), position=1)
        SellerWorkInterval.objects.create(seller=full, weekday=6, is_closed=True, position=0)
        SellerExternalCard.objects.create(
            seller=full, source='two_gis', external_id='7001', card_url='https://2gis.ru/firm/7001',
        )
        SellerExternalCard.objects.create(
            seller=full, source='yandex', external_id='555', card_url='https://yandex.ru/maps/org/555',
        )
        SellerExternalCard.objects.create(
            seller=full, source='google', external_id='g1', card_url='https://www.google.com/maps/place/barys',
        )
        partial = _seller(
            name='Короткая карточка',
            is_dismantler=True,
            whatsapp='77019990022',
            city='',
            market_location='',
            notes='',
        )

        full_url = reverse('parts_seller_detail_public', args=[full.id])
        partial_url = reverse('parts_seller_detail_public', args=[partial.id])
        full_body = self.client.get(full_url).content.decode()
        partial_body = self.client.get(partial_url).content.decode()

        self.assertIn('Построить маршрут', full_body)
        self.assertIn('https://2gis.kz/directions/points/|76.950000,43.250000', full_body)
        self.assertIn('https://www.google.com/maps/dir/?api=1&amp;destination=43.250000,76.950000', full_body)
        self.assertIn('https://yandex.kz/maps/?rtext=~43.250000,76.950000&amp;rtt=auto', full_body)
        self.assertIn('Посмотреть в 2ГИС', full_body)
        self.assertIn('Посмотреть в Яндексе', full_body)
        self.assertIn('Посмотреть в Google Картах', full_body)
        self.assertIn('https://2gis.ru/firm/7001', full_body)
        self.assertNotIn('?text=', full_body)
        self.assertNotIn('/search/', full_body)
        self.assertIn('Режим работы', full_body)
        self.assertIn('09:00–13:00, 14:00–18:00', full_body)
        self.assertIn('выходной', full_body)
        self.assertNotIn('Открыто сейчас', full_body)
        self.assertIn('tel:+77012223344', full_body)
        self.assertIn('tel:+77013334455', full_body)
        self.assertIn('Instagram', full_body)
        self.assertIn('Telegram', full_body)
        self.assertIn('TikTok', full_body)
        self.assertIn('YouTube', full_body)
        self.assertIn('Сайт организации', full_body)
        self.assertIn('очень длинный адрес', full_body)
        self.assertNotIn('Оригинальное описание', full_body)
        self.assertIn('dismantlers.css?v=20261017a', full_body)

        self.assertNotIn('Построить маршрут', partial_body)
        self.assertNotIn('Режим работы', partial_body)
        self.assertNotIn('Посмотреть в', partial_body)
        self.assertNotIn('Не указано', partial_body)
        self.assertNotIn('Не указан', partial_body)
        self.assertNotIn('Режим работы', partial_body)
        self.assertIn('tel:+77019990022', partial_body)

        routes = route_links(full.latitude, full.longitude)
        self.assertEqual(routes[0]['label'], '2ГИС')
        self.assertEqual(len(build_seller_card(partial)['routes']), 0)
        self.assertEqual(len(build_seller_card(partial)['sources']), 0)

    def test_catalog_shows_address_without_empty_placeholder(self):
        _seller(
            name='С адресом',
            is_dismantler=True,
            address='улица Абая, 10',
            city='Алматы',
            market_location='',
        )
        body = self.client.get(reverse('dismantlers_catalog')).content.decode()
        self.assertIn('улица Абая, 10', body)
        self.assertNotIn('Не указано', body)


class WorkbookCommandTests(TestCase):
    def test_preview_of_workbook_does_not_write(self):
        payload = _mapsgrabs_bytes([
            [
                'Яндекс Карты', '1208969566', 'Motor center',
                'Актобе, просп. Кенеса Нокина, 18Л', 'Актобе', '50.306148, 57.206707',
                '+77024700555, +77764700555', 'http://motor-center.kz/',
                'https://wa.me/77024700555, https://wa.me/277075562423, https://instagram.com/motor_center_kz',
                'пн-пт 09:00–18:00; сб 10:00–18:00',
                'https://yandex.ru/maps/org/1208969566',
                4.3, 18,
            ],
        ], source='Яндекс Карты')
        path = self._write(payload)
        _seller(name='Motor center', whatsapp='77024700555', city='Актобе', address='')
        report = enrich_files([path], apply=False)
        self.assertEqual(len(report['matched']), 1)
        self.assertIsNone(report['applied'])
        self.assertEqual(SellerExternalCard.objects.count(), 0)
        self.assertTrue(report['suspicious_whatsapp'])
        self.assertEqual(Seller.objects.get().address, '')

    def test_admin_preview_does_not_create_sellers(self):
        User.objects.create_superuser('card-admin', 'card@example.com', 'pass-word-1')
        self.client.login(username='card-admin', password='pass-word-1')
        payload = _mapsgrabs_bytes([
            [
                '2ГИС', '42', 'Новый разбор', 'улица Тест, 1', 'Семей',
                '50.400000, 80.200000', '+77015550999', '', '', '',
                'https://2gis.ru/firm/42', '', '',
            ],
        ])
        upload = SimpleUploadedFile(
            'family.xlsx',
            payload.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response = self.client.post(
            reverse('admin:core_seller_enrich_cards'),
            {'action': 'preview', 'file': upload},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Seller.objects.count(), 0)
        self.assertContains(response, 'Без совпадения')

    def _write(self, payload):
        from pathlib import Path
        from tempfile import TemporaryDirectory

        folder = TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        path = Path(folder.name) / 'export.xlsx'
        path.write_bytes(payload.getvalue())
        return str(path)
