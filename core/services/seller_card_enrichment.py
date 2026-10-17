"""Match map exports to existing sellers and fill empty card fields.

The enrichment never creates sellers and never changes request dispatch,
activity, or publication flags. Filled values are kept when a new value
conflicts. Empty import cells do not clear stored data.
"""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from urllib.parse import urlparse

from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from openpyxl import load_workbook

from core.models import (
    Seller,
    SellerExternalCard,
    SellerExtraPhone,
    SellerSocialLink,
    SellerWorkInterval,
)
from core.phone_utils import normalize_phone_for_whatsapp
from core.services.seller_card_links import (
    assess_whatsapp_link,
    classify_social,
    clean_http_url,
    is_map_host,
    validate_card_url,
)

SOURCE_LABELS = {
    'two_gis': '2ГИС',
    'yandex': 'Яндекс Карты',
    'google': 'Google Карты',
}

HEADER_FIELDS = {
    'источник': 'source',
    'id организации': 'external_id',
    'название': 'name',
    'адрес': 'address',
    'город': 'city',
    'координаты': 'coordinates',
    'контакты': 'phones',
    'сайты': 'websites',
    'соцсети': 'socials',
    'режим работы': 'schedule',
    'ссылка на карточку': 'card_url',
}

GENERIC_NAME_TOKENS = {
    'авто', 'авторазбор', 'авторазборка', 'авторазборы', 'автосервис',
    'автозапчасти', 'запчасти', 'запчасть', 'магазин', 'центр', 'разбор',
    'разборка', 'сервис', 'сто', 'и', 'на', 'для', 'shop', 'auto', 'motor',
    'motors', 'kz', 'the',
}

CITY_ALIASES = {
    'астана': 'астана',
    'нурсултан': 'астана',
    'нур султан': 'астана',
    'актобе': 'актобе',
    'актюбинск': 'актобе',
    'алматы': 'алматы',
    'алма ата': 'алматы',
    'шымкент': 'шымкент',
    'семей': 'семей',
    'семипалатинск': 'семей',
    'кокшетау': 'кокшетау',
}

DAY_INDEX = {
    'пн': 0, 'понедельник': 0,
    'вт': 1, 'вторник': 1,
    'ср': 2, 'среда': 2,
    'чт': 3, 'четверг': 3,
    'пт': 4, 'пятница': 4,
    'сб': 5, 'суббота': 5,
    'вс': 6, 'воскресенье': 6,
}
DAY_TOKEN = (
    'понедельник|вторник|среда|четверг|пятница|суббота|воскресенье|'
    'пн|вт|ср|чт|пт|сб|вс'
)
SEGMENT_RE = re.compile(
    rf'^(?P<days>ежедневно|круглосуточно|(?:{DAY_TOKEN})(?:\s*[-–—]\s*(?:{DAY_TOKEN}))?)\s*[:,]?\s*(?P<rest>.*)$',
    re.IGNORECASE,
)
INTERVAL_RE = re.compile(r'(\d{1,2}:\d{2})\s*[-–—]\s*(\d{1,2}:\d{2})')
URL_RE = re.compile(r'https?://[^\s<>"\']+', re.IGNORECASE)
PHONE_RE = re.compile(r'(?:\+\d|\b[78]\d)[\d\-\s()]{8,}\d')
FIRM_RE = re.compile(r'/(?:firm|org)/(?:[^/\d][^/]*/)?(\d+)')
ADDRESS_LINE_RE = re.compile(r'^\s*адрес\s*[:\-]\s*(.+)$', re.IGNORECASE)


def _is_hidden_export_row(record: dict) -> bool:
    source = _text(record.get('source')).lower()
    if 'тариф' in source or 'первые 50' in source:
        return True
    for field in ('source', 'external_id', 'name', 'address'):
        text = _text(record.get(field))
        if text and set(text) <= {'*'}:
            return True
    return False


def _text(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_text(value: object) -> str:
    text = _text(value).lower().replace('ё', 'е')
    text = text.replace('нур-султан', 'нур султан').replace('алма-ата', 'алма ата')
    text = re.sub(r'[^\w\s]', ' ', text, flags=re.UNICODE)
    return re.sub(r'\s+', ' ', text).strip()


def city_key(value: object) -> str:
    text = normalize_text(value)
    return CITY_ALIASES.get(text, text)


def parse_source(value: object) -> str:
    text = normalize_text(value)
    if '2гис' in text or '2gis' in text:
        return 'two_gis'
    if 'яндекс' in text or 'yandex' in text:
        return 'yandex'
    if 'google' in text or 'гугл' in text:
        return 'google'
    return ''


def split_list(value: object) -> list[str]:
    text = _text(value)
    if not text or text.lower() in {'none', 'null'}:
        return []
    return [part.strip() for part in re.split(r'\s*,\s*', text) if part.strip()]


def parse_coordinates(value: object) -> tuple[str, str, str]:
    text = _text(value)
    if not text:
        return '', '', ''
    parts = [part.strip() for part in text.split(',')]
    if len(parts) != 2:
        return '', '', 'некорректные координаты'
    try:
        lat = Decimal(parts[0])
        lon = Decimal(parts[1])
    except (InvalidOperation, ValueError):
        return '', '', 'некорректные координаты'
    if not (Decimal('-90') <= lat <= Decimal('90') and Decimal('-180') <= lon <= Decimal('180')):
        return '', '', 'координаты вне допустимого диапазона'
    if lat == 0 and lon == 0:
        return '', '', 'пустые координаты'
    return format(lat, 'f'), format(lon, 'f'), ''


def _hhmm(value: str) -> str:
    hour, minute = (int(part) for part in value.split(':'))
    if hour == 24 and minute == 0:
        return '23:59'
    if not (0 <= hour <= 23 and 0 <= minute <= 59):
        return ''
    return f'{hour:02d}:{minute:02d}'


def _day_span(token: str) -> list[int]:
    cleaned = token.strip().lower().replace('ё', 'е')
    parts = re.split(r'\s*[-–—]\s*', cleaned)
    if len(parts) == 1:
        index = DAY_INDEX.get(parts[0])
        return [] if index is None else [index]
    start = DAY_INDEX.get(parts[0])
    end = DAY_INDEX.get(parts[1]) if len(parts) == 2 else None
    if start is None or end is None:
        return []
    if start <= end:
        return list(range(start, end + 1))
    return list(range(start, 7)) + list(range(0, end + 1))


def parse_schedule(value: object) -> tuple[list[dict], str]:
    text = _text(value)
    if not text:
        return [], ''
    lowered = text.lower().replace('ё', 'е')
    if lowered in {'круглосуточно', '24 часа', '24/7'}:
        return [
            {'weekday': day, 'opens': '00:00', 'closes': '23:59', 'is_closed': False}
            for day in range(7)
        ], ''
    intervals = []
    unknown = []
    for raw_segment in re.split(r'\s*;\s*', text):
        segment = raw_segment.strip()
        if not segment:
            continue
        match = SEGMENT_RE.match(segment.lower().replace('ё', 'е'))
        if not match:
            unknown.append(segment)
            continue
        days_token = match.group('days')
        rest = match.group('rest') or ''
        days = list(range(7)) if days_token in {'ежедневно', 'круглосуточно'} else _day_span(days_token)
        if not days:
            unknown.append(segment)
            continue
        if 'выходн' in rest or rest.strip() in {'закрыто', 'закрыт'}:
            for day in days:
                intervals.append({
                    'weekday': day,
                    'opens': '',
                    'closes': '',
                    'is_closed': True,
                })
            continue
        found = []
        for opens_raw, closes_raw in INTERVAL_RE.findall(rest):
            opens = _hhmm(opens_raw)
            closes = _hhmm(closes_raw)
            if not opens or not closes or opens == closes:
                unknown.append(segment)
                found = []
                break
            found.append((opens, closes))
        if not found:
            if segment not in unknown:
                unknown.append(segment)
            continue
        for day in days:
            for opens, closes in found:
                intervals.append({
                    'weekday': day,
                    'opens': opens,
                    'closes': closes,
                    'is_closed': False,
                })
    if unknown and not intervals:
        return [], 'режим работы не разобран'
    if unknown:
        return intervals, 'часть режима работы не разобрана'
    return intervals, ''


def schedule_key(intervals: list[dict]) -> tuple:
    rows = []
    for item in intervals:
        rows.append((
            int(item['weekday']),
            item.get('opens') or '',
            item.get('closes') or '',
            bool(item.get('is_closed')),
        ))
    return tuple(sorted(rows))


def _phones_from_text(value: object) -> list[str]:
    phones = []
    for match in PHONE_RE.findall(_text(value)):
        normalized = normalize_phone_for_whatsapp(match)
        if normalized and normalized not in phones:
            phones.append(normalized)
    return phones


def _classify_links(raw_items: list[str], contact_phones: set[str], bucket: str) -> dict:
    socials = []
    websites = []
    invalid = []
    suspicious = []
    seen_socials = set()
    for raw in raw_items:
        cleaned, error = clean_http_url(raw)
        if error:
            invalid.append({'url': _text(raw), 'reason': error, 'bucket': bucket})
            continue
        if not cleaned:
            continue
        kind = classify_social(cleaned)
        if kind == 'whatsapp':
            url, digits, problem = assess_whatsapp_link(raw, contact_phones)
            if problem:
                suspicious.append({'url': _text(raw), 'reason': problem})
                continue
            key = ('whatsapp', url)
            if key not in seen_socials:
                seen_socials.add(key)
                socials.append({'kind': 'whatsapp', 'url': url, 'phone': digits})
            continue
        if kind:
            key = (kind, cleaned)
            if key not in seen_socials:
                seen_socials.add(key)
                socials.append({'kind': kind, 'url': cleaned})
            continue
        if is_map_host(cleaned):
            invalid.append({
                'url': cleaned,
                'reason': 'ссылка ведёт на карту, а не на сайт организации',
                'bucket': bucket,
            })
            continue
        if bucket == 'socials':
            invalid.append({
                'url': cleaned,
                'reason': 'домен не относится к поддерживаемой соцсети',
                'bucket': bucket,
            })
            continue
        websites.append(cleaned)
    return {
        'socials': socials,
        'websites': websites,
        'invalid': invalid,
        'suspicious': suspicious,
    }


def parse_organization_row(row: dict, *, exported_at: str, file_name: str, row_number: int) -> dict:
    source = parse_source(row.get('source'))
    external_id = _text(row.get('external_id'))
    phones = []
    rejected_phones = []
    for raw_phone in split_list(row.get('phones')):
        normalized = normalize_phone_for_whatsapp(raw_phone)
        if not normalized:
            rejected_phones.append(raw_phone)
            continue
        if normalized not in phones:
            phones.append(normalized)
    contact_phones = set(phones)
    social_pack = _classify_links(split_list(row.get('socials')), contact_phones, 'socials')
    site_pack = _classify_links(split_list(row.get('websites')), contact_phones, 'websites')
    websites = []
    for url in site_pack['websites']:
        if url not in websites:
            websites.append(url)
    socials = []
    seen = set()
    for item in social_pack['socials'] + site_pack['socials']:
        key = (item['kind'], item['url'])
        if key in seen:
            continue
        seen.add(key)
        socials.append({'kind': item['kind'], 'url': item['url']})
    latitude, longitude, coordinate_error = parse_coordinates(row.get('coordinates'))
    schedule, schedule_error = parse_schedule(row.get('schedule'))
    card_url, card_error = ('', '')
    if _text(row.get('card_url')):
        card_url, card_error = validate_card_url(source, row.get('card_url'))
    invalid = social_pack['invalid'] + site_pack['invalid']
    if coordinate_error:
        invalid.append({'url': _text(row.get('coordinates')), 'reason': coordinate_error, 'bucket': 'coordinates'})
    if schedule_error:
        invalid.append({'url': _text(row.get('schedule')), 'reason': schedule_error, 'bucket': 'schedule'})
    if card_error:
        invalid.append({'url': _text(row.get('card_url')), 'reason': card_error, 'bucket': 'card'})
    for raw_phone in rejected_phones:
        invalid.append({'url': raw_phone, 'reason': 'телефон не прошёл нормализацию', 'bucket': 'phone'})
    website = websites[0] if len(websites) == 1 else ''
    website_conflict = websites if len(websites) > 1 else []
    return {
        'source': source,
        'external_id': external_id,
        'name': _text(row.get('name')),
        'address': _text(row.get('address')),
        'city': '' if _text(row.get('city')).lower() == 'none' else _text(row.get('city')),
        'latitude': latitude,
        'longitude': longitude,
        'phones': phones,
        'website': website,
        'website_options': website_conflict,
        'socials': socials,
        'schedule': schedule,
        'card_url': card_url,
        'exported_at': exported_at,
        'file_name': file_name,
        'row_number': row_number,
        'invalid': invalid,
        'suspicious_whatsapp': social_pack['suspicious'] + site_pack['suspicious'],
    }


def _sheet_headers(header_row) -> dict[str, int]:
    headers = {}
    for index, value in enumerate(header_row):
        key = HEADER_FIELDS.get(normalize_text(value))
        if key:
            headers[key] = index
    return headers


def _meta_value(workbook, name: str) -> str:
    if 'О выгрузке' not in workbook.sheetnames:
        return ''
    wanted = normalize_text(name)
    for key, value in workbook['О выгрузке'].iter_rows(values_only=True):
        if normalize_text(key) == wanted:
            return _text(value)
    return ''


def parse_mapsgrabs_xlsx(file_obj, file_name: str = '') -> tuple[list[dict], list[str]]:
    """Read one MapsGrabs workbook. Reviews, ratings, and media are ignored."""
    if hasattr(file_obj, 'read'):
        payload = file_obj.read()
        if hasattr(file_obj, 'seek'):
            file_obj.seek(0)
        workbook = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    else:
        workbook = load_workbook(file_obj, data_only=True, read_only=True)
    try:
        if 'Организации' not in workbook.sheetnames:
            return [], ['В файле нет листа «Организации»']
        sheet = workbook['Организации']
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            return [], ['Лист «Организации» пуст']
        headers = _sheet_headers(header)
        required = ('source', 'external_id', 'name', 'card_url')
        missing = [name for name in required if name not in headers]
        if missing:
            return [], ['Не найдены колонки выгрузки: ' + ', '.join(missing)]
        exported_at = _meta_value(workbook, 'Дата выгрузки')
        parsed = []
        errors = []
        hidden_rows = 0
        for offset, row in enumerate(rows, start=2):
            if not row or not any(row):
                continue
            record = {
                field: row[index] if index < len(row) else ''
                for field, index in headers.items()
            }
            if _is_hidden_export_row(record):
                hidden_rows += 1
                continue
            org = parse_organization_row(
                record,
                exported_at=exported_at,
                file_name=file_name,
                row_number=offset,
            )
            if not org['source'] or not org['external_id']:
                errors.append(f'Строка {offset}: нет источника или ID организации')
                continue
            parsed.append(org)
        if hidden_rows:
            errors.append(
                f'Скрыто тарифом выгрузки: {hidden_rows} строк без открытых контактов'
            )
        return parsed, errors
    finally:
        workbook.close()


def is_mapsgrabs_workbook(file_obj) -> bool:
    payload = file_obj.read()
    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        return 'Организации' in workbook.sheetnames
    finally:
        workbook.close()


def _merge_orgs(orgs: list[dict]) -> tuple[list[dict], list[dict]]:
    grouped = defaultdict(list)
    for org in orgs:
        grouped[(org['source'], org['external_id'])].append(org)
    merged = []
    conflicts = []
    for (source, external_id), items in grouped.items():
        base = dict(items[-1])
        base['phones'] = []
        base['socials'] = []
        base['invalid'] = []
        base['suspicious_whatsapp'] = []
        base['row_numbers'] = []
        websites = []
        schedules = []
        addresses = []
        coordinates = []
        card_urls = []
        for item in items:
            base['row_numbers'].append(item['row_number'])
            base['invalid'].extend(item['invalid'])
            base['suspicious_whatsapp'].extend(item['suspicious_whatsapp'])
            for phone in item['phones']:
                if phone not in base['phones']:
                    base['phones'].append(phone)
            for social in item['socials']:
                if social not in base['socials']:
                    base['socials'].append(social)
            if item['address']:
                addresses.append(item['address'])
            if item['latitude'] and item['longitude']:
                coordinates.append((item['latitude'], item['longitude']))
            if item['card_url']:
                card_urls.append(item['card_url'])
            if item['website']:
                websites.append(item['website'])
            websites.extend(item.get('website_options') or [])
            if item['schedule']:
                schedules.append(schedule_key(item['schedule']))
        base['address'] = _agreed(addresses, conflicts, source, external_id, base['name'], 'адрес')
        agreed_coordinates = _agreed(coordinates, conflicts, source, external_id, base['name'], 'координаты')
        if agreed_coordinates:
            base['latitude'], base['longitude'] = agreed_coordinates
        else:
            base['latitude'] = ''
            base['longitude'] = ''
        base['card_url'] = _agreed(card_urls, conflicts, source, external_id, base['name'], 'ссылка на карточку')
        unique_sites = []
        for site in websites:
            if site not in unique_sites:
                unique_sites.append(site)
        base['website'] = unique_sites[0] if len(unique_sites) == 1 else ''
        if len(unique_sites) > 1:
            conflicts.append({
                'source': source,
                'external_id': external_id,
                'name': base['name'],
                'field': 'сайт',
                'values': unique_sites,
            })
        unique_schedules = []
        for key in schedules:
            if key not in unique_schedules:
                unique_schedules.append(key)
        if len(unique_schedules) == 1:
            base['schedule'] = [
                {'weekday': day, 'opens': opens, 'closes': closes, 'is_closed': closed}
                for day, opens, closes, closed in unique_schedules[0]
            ]
        elif len(unique_schedules) > 1:
            base['schedule'] = []
            conflicts.append({
                'source': source,
                'external_id': external_id,
                'name': base['name'],
                'field': 'режим работы',
                'values': ['разные расписания в повторных строках'],
            })
        merged.append(base)
    return merged, conflicts


def _agreed(values, conflicts, source, external_id, name, field):
    unique = []
    for value in values:
        if value not in unique:
            unique.append(value)
    if len(unique) == 1:
        return unique[0]
    if len(unique) > 1:
        conflicts.append({
            'source': source,
            'external_id': external_id,
            'name': name,
            'field': field,
            'values': [str(value) for value in unique],
        })
    return '' if not unique else ''


def name_tokens(value: object) -> set[str]:
    tokens = {
        token for token in normalize_text(value).split()
        if len(token) >= 3 and token not in GENERIC_NAME_TOKENS
    }
    return tokens


def _tokens_fit(shorter: list[str], longer: list[str]) -> bool:
    if not shorter or len(shorter) > len(longer):
        return False
    for start in range(len(longer) - len(shorter) + 1):
        if longer[start:start + len(shorter)] == shorter:
            return True
    return False


def names_compatible(left: object, right: object) -> bool:
    a = normalize_text(left)
    b = normalize_text(right)
    if not a or not b:
        return False
    if a == b:
        return True
    a_tokens = a.split()
    b_tokens = b.split()
    if _tokens_fit(a_tokens, b_tokens) or _tokens_fit(b_tokens, a_tokens):
        return True
    left_tokens = name_tokens(left)
    right_tokens = name_tokens(right)
    if not left_tokens or not right_tokens:
        return False
    shared = left_tokens & right_tokens
    smaller = min(len(left_tokens), len(right_tokens))
    return len(shared) / smaller >= 0.5


def _house_numbers(value: str) -> list[tuple[str, str]]:
    found = []
    for match in re.finditer(r'(\d+)\s*([a-zа-я])?', normalize_text(value)):
        found.append((match.group(1), match.group(2) or ''))
    return found


def addresses_compatible(left: object, right: object) -> bool:
    a = normalize_text(left)
    b = normalize_text(right)
    if not a or not b:
        return False
    if a == b:
        return True
    houses_a = _house_numbers(a)
    houses_b = _house_numbers(b)
    if not houses_a or not houses_b:
        return False
    house_match = False
    for number_a, suffix_a in houses_a:
        for number_b, suffix_b in houses_b:
            if number_a == number_b and (not suffix_a or not suffix_b or suffix_a == suffix_b):
                house_match = True
    if not house_match:
        return False
    stop = {'улица', 'ул', 'проспект', 'просп', 'пр', 'город', 'область', 'район', 'дом', 'мкр'}
    words_a = {word for word in a.split() if len(word) >= 4 and not word.isdigit()} - stop
    words_b = {word for word in b.split() if len(word) >= 4 and not word.isdigit()} - stop
    return bool(words_a & words_b)


def cities_compatible(left: object, right: object) -> bool:
    a = city_key(left)
    b = city_key(right)
    if not a or not b:
        return True
    return a == b


def _address_candidates(seller_snapshot: dict) -> list[str]:
    values = []
    for value in (seller_snapshot.get('address'), seller_snapshot.get('market_location')):
        text = _text(value)
        if text and not text.lower().startswith(('http://', 'https://')):
            values.append(text)
    for line in _text(seller_snapshot.get('notes')).splitlines():
        match = ADDRESS_LINE_RE.match(line)
        if match:
            values.append(match.group(1).strip())
    return values


def _note_urls(notes: str) -> list[str]:
    return URL_RE.findall(notes or '')


def _external_ids_from_text(notes: str) -> list[tuple[str, str]]:
    found = []
    for url in _note_urls(notes):
        source = ''
        host = (urlparse(url).hostname or '').lower()
        if '2gis.' in host:
            source = 'two_gis'
        elif 'yandex.' in host:
            source = 'yandex'
        elif 'google.' in host or host.endswith('goo.gl'):
            source = 'google'
        match = FIRM_RE.search(urlparse(url).path)
        if source and match:
            found.append((source, match.group(1)))
    return found


def seller_snapshots() -> list[dict]:
    sellers = Seller.objects.prefetch_related(
        'extra_phones',
        'work_hours',
        'social_links',
        'external_cards',
    )
    snapshots = []
    for seller in sellers:
        phones = []
        for raw in (seller.whatsapp, seller.phone2):
            normalized = normalize_phone_for_whatsapp(raw)
            if normalized and normalized not in phones:
                phones.append(normalized)
        for extra in seller.extra_phones.all():
            if extra.phone not in phones:
                phones.append(extra.phone)
        for phone in _phones_from_text(seller.notes):
            if phone not in phones:
                phones.append(phone)
        external = [
            (card.source, card.external_id, card.card_url, card.exported_at)
            for card in seller.external_cards.all()
        ]
        for source, external_id in _external_ids_from_text(seller.notes):
            if (source, external_id) not in {(item[0], item[1]) for item in external}:
                external.append((source, external_id, '', None))
        schedule = [
            {
                'weekday': interval.weekday,
                'opens': interval.opens.strftime('%H:%M') if interval.opens else '',
                'closes': interval.closes.strftime('%H:%M') if interval.closes else '',
                'is_closed': interval.is_closed,
            }
            for interval in seller.work_hours.all()
        ]
        snapshots.append({
            'id': seller.id,
            'name': seller.name,
            'city': seller.city,
            'address': seller.address,
            'market_location': seller.market_location,
            'notes': seller.notes,
            'whatsapp': seller.whatsapp,
            'phone2': seller.phone2,
            'latitude': '' if seller.latitude is None else format(seller.latitude, 'f'),
            'longitude': '' if seller.longitude is None else format(seller.longitude, 'f'),
            'website': seller.website or '',
            'phones': phones,
            'socials': [{'kind': link.kind, 'url': link.url} for link in seller.social_links.all()],
            'external': external,
            'schedule': schedule,
            'protected': {
                'is_active': seller.is_active,
                'is_paused': seller.is_paused,
                'receive_requests': seller.receive_requests,
            },
        })
    return snapshots


def _index_sellers(snapshots: list[dict]) -> dict:
    phone_owners = defaultdict(set)
    external_owners = defaultdict(set)
    for seller in snapshots:
        for phone in seller['phones']:
            phone_owners[phone].add(seller['id'])
        for source, external_id, _url, _exported in seller['external']:
            external_owners[(source, external_id)].add(seller['id'])
    return {
        'by_id': {seller['id']: seller for seller in snapshots},
        'phone_owners': phone_owners,
        'external_owners': external_owners,
    }


def match_organization(org: dict, index: dict) -> dict:
    candidates = {}
    shared_phones = []

    def add(seller_id, reason):
        candidates.setdefault(seller_id, [])
        if reason not in candidates[seller_id]:
            candidates[seller_id].append(reason)

    for phone in org['phones']:
        owners = index['phone_owners'].get(phone, set())
        if len(owners) == 1:
            add(next(iter(owners)), f'телефон {phone}')
        elif len(owners) > 1:
            shared_phones.append(phone)
    owners = index['external_owners'].get((org['source'], org['external_id']), set())
    if len(owners) == 1:
        add(next(iter(owners)), f'внешний ID {org["external_id"]}')
    address_matches = []
    for seller in index['by_id'].values():
        if not names_compatible(org['name'], seller['name']):
            continue
        if not cities_compatible(org['city'], seller['city']):
            continue
        if any(addresses_compatible(org['address'], candidate) for candidate in _address_candidates(seller)):
            address_matches.append(seller['id'])
    if len(address_matches) == 1:
        add(address_matches[0], 'название и адрес')
    elif len(address_matches) > 1:
        for seller_id in address_matches:
            add(seller_id, 'название и адрес совпали у нескольких продавцов')
    if len(candidates) == 1:
        seller_id = next(iter(candidates))
        return {
            'status': 'matched',
            'seller_id': seller_id,
            'reasons': candidates[seller_id],
            'shared_phones': shared_phones,
        }
    if len(candidates) > 1:
        return {
            'status': 'ambiguous',
            'seller_id': None,
            'reasons': [
                f'#{seller_id}: {", ".join(reasons)}'
                for seller_id, reasons in candidates.items()
            ],
            'shared_phones': shared_phones,
        }
    if shared_phones:
        return {
            'status': 'ambiguous',
            'seller_id': None,
            'reasons': ['общий телефон принадлежит нескольким филиалам'],
            'shared_phones': shared_phones,
        }
    return {
        'status': 'unmatched',
        'seller_id': None,
        'reasons': ['нет однозначного совпадения по телефону, внешнему ID или названию с адресом'],
        'shared_phones': [],
    }


def _same_coordinate(left: str, right: str) -> bool:
    if not left or not right:
        return False
    return abs(Decimal(left) - Decimal(right)) <= Decimal('0.0001')


def _note_proposals(seller: dict) -> dict:
    phones = [phone for phone in _phones_from_text(seller['notes']) if phone not in {normalize_phone_for_whatsapp(seller['whatsapp'])}]
    pack = _classify_links(_note_urls(seller['notes']), set(seller['phones']), 'socials')
    websites = pack['websites']
    addresses = []
    for line in _text(seller['notes']).splitlines():
        match = ADDRESS_LINE_RE.match(line)
        if match:
            addresses.append(match.group(1).strip())
    return {
        'phones': phones,
        'socials': [{'kind': item['kind'], 'url': item['url']} for item in pack['socials']],
        'website': websites[0] if len(websites) == 1 else '',
        'website_options': websites if len(websites) > 1 else [],
        'address': addresses[0] if len(addresses) == 1 else '',
        'address_options': addresses if len(addresses) > 1 else [],
        'invalid': pack['invalid'],
        'suspicious_whatsapp': pack['suspicious'],
    }


def plan_seller_updates(seller: dict, orgs: list[dict]) -> dict:
    fills = []
    conflicts = []
    notes = _note_proposals(seller)
    addresses = [org['address'] for org in orgs if org.get('address')]
    if notes['address']:
        addresses.append(notes['address'])
    addresses.extend(notes['address_options'])
    coordinates = [
        (org['latitude'], org['longitude'])
        for org in orgs
        if org.get('latitude') and org.get('longitude')
    ]
    websites = [org['website'] for org in orgs if org.get('website')]
    websites.extend(org.get('website_options') or [] for org in orgs)
    flat_websites = []
    for item in websites:
        if isinstance(item, list):
            flat_websites.extend(item)
        elif item:
            flat_websites.append(item)
    if notes['website']:
        flat_websites.append(notes['website'])
    flat_websites.extend(notes['website_options'])
    schedules = [schedule_key(org['schedule']) for org in orgs if org.get('schedule')]
    phones = []
    socials = []
    cards = []
    for org in orgs:
        for phone in org.get('phones') or []:
            if phone not in phones:
                phones.append(phone)
        for social in org.get('socials') or []:
            if social not in socials:
                socials.append(social)
        if org.get('source') and org.get('external_id') and org.get('card_url'):
            cards.append(org)
    for url in _note_urls(seller['notes']):
        for source in ('two_gis', 'yandex', 'google'):
            cleaned, error = validate_card_url(source, url)
            if error or not cleaned:
                continue
            match = FIRM_RE.search(urlparse(cleaned).path)
            if not match:
                continue
            cards.append({
                'source': source,
                'external_id': match.group(1),
                'card_url': cleaned,
                'exported_at': '',
            })
            break
    for phone in notes['phones']:
        if phone not in phones:
            phones.append(phone)
    for social in notes['socials']:
        if social not in socials:
            socials.append(social)

    address = _single(addresses, conflicts, seller, 'адрес')
    if address:
        if seller['address']:
            if not addresses_compatible(seller['address'], address):
                conflicts.append(_conflict(seller, 'адрес', seller['address'], address))
            else:
                address = ''
        else:
            fills.append({'field': 'address', 'value': address[:500]})
    coordinate = _single(coordinates, conflicts, seller, 'координаты')
    if coordinate:
        lat, lon = coordinate
        if seller['latitude'] or seller['longitude']:
            if not (_same_coordinate(seller['latitude'], lat) and _same_coordinate(seller['longitude'], lon)):
                conflicts.append(_conflict(seller, 'координаты', f'{seller["latitude"]}, {seller["longitude"]}', f'{lat}, {lon}'))
        else:
            fills.append({'field': 'latitude', 'value': lat})
            fills.append({'field': 'longitude', 'value': lon})
    website = _single(flat_websites, conflicts, seller, 'сайт')
    if website:
        if seller['website']:
            if seller['website'].rstrip('/') != website.rstrip('/'):
                conflicts.append(_conflict(seller, 'сайт', seller['website'], website))
        else:
            fills.append({'field': 'website', 'value': website})
    schedule = _single(schedules, conflicts, seller, 'режим работы')
    if schedule:
        if seller['schedule']:
            if schedule_key(seller['schedule']) != schedule:
                conflicts.append(_conflict(seller, 'режим работы', 'заполнен', 'другое расписание'))
            schedule = None
        else:
            fills.append({
                'field': 'schedule',
                'value': [
                    {'weekday': day, 'opens': opens, 'closes': closes, 'is_closed': closed}
                    for day, opens, closes, closed in schedule
                ],
            })
    primary = normalize_phone_for_whatsapp(seller['whatsapp']) or ''
    secondary = normalize_phone_for_whatsapp(seller['phone2']) or ''
    known_phones = set(seller['phones'])
    new_phones = [phone for phone in phones if phone not in known_phones and phone != primary]
    phone2_value = ''
    extra_phones = []
    if new_phones:
        if not _text(seller['phone2']):
            phone2_value = new_phones[0]
            extra_phones = new_phones[1:]
            fills.append({'field': 'phone2', 'value': phone2_value})
        else:
            extra_phones = [phone for phone in new_phones if phone != secondary]
        for phone in extra_phones:
            fills.append({'field': 'extra_phone', 'value': phone})
    known_socials = {(item['kind'], item['url']) for item in seller['socials']}
    for social in socials:
        key = (social['kind'], social['url'])
        if key not in known_socials:
            fills.append({'field': 'social', 'value': social})
            known_socials.add(key)
    known_cards = {(item[0], item[1]): item for item in seller['external']}
    cards_by_source = {item[0]: item for item in seller['external']}
    for org in cards:
        key = (org['source'], org['external_id'])
        existing = known_cards.get(key)
        same_source = cards_by_source.get(org['source'])
        if existing:
            _url = existing[2]
            if _url and _url.rstrip('/') != org['card_url'].rstrip('/'):
                conflicts.append(_conflict(seller, f'ссылка {SOURCE_LABELS.get(org["source"], org["source"])}', _url, org['card_url']))
            elif not _url:
                fills.append({'field': 'card', 'value': _card_payload(org), 'update_url': True})
            continue
        if same_source and same_source[1] != org['external_id']:
            conflicts.append(_conflict(
                seller,
                f'источник {SOURCE_LABELS.get(org["source"], org["source"])}',
                same_source[1],
                org['external_id'],
            ))
            continue
        fills.append({'field': 'card', 'value': _card_payload(org)})
        cards_by_source[org['source']] = (org['source'], org['external_id'], org['card_url'], None)
    return {'fills': fills, 'conflicts': conflicts}


def _single(values, conflicts, seller, field):
    unique = []
    for value in values:
        if value and value not in unique:
            unique.append(value)
    if len(unique) > 1:
        conflicts.append(_conflict(seller, field, '', 'несколько разных значений'))
        return None if field == 'режим работы' else ''
    return unique[0] if unique else (None if field == 'режим работы' else '')


def _conflict(seller, field, current, proposed):
    return {
        'seller_id': seller['id'],
        'seller_name': seller['name'],
        'field': field,
        'current': current,
        'proposed': proposed,
    }


def _card_payload(org: dict) -> dict:
    return {
        'source': org['source'],
        'external_id': org['external_id'],
        'card_url': org['card_url'],
        'exported_at': org.get('exported_at') or '',
    }


def analyze_organizations(orgs: list[dict], snapshots: list[dict] | None = None) -> dict:
    merged, merge_conflicts = _merge_orgs(orgs)
    snapshots = seller_snapshots() if snapshots is None else snapshots
    index = _index_sellers(snapshots)
    matched = []
    unmatched = []
    ambiguous = []
    invalid = []
    suspicious = []
    by_seller = defaultdict(list)
    for org in merged:
        for item in org['invalid']:
            invalid.append({**item, 'name': org['name'], 'external_id': org['external_id'], 'source': org['source']})
        for item in org['suspicious_whatsapp']:
            suspicious.append({**item, 'name': org['name'], 'external_id': org['external_id'], 'source': org['source']})
        result = match_organization(org, index)
        row = {
            'name': org['name'],
            'city': org['city'],
            'source': org['source'],
            'external_id': org['external_id'],
            'file_name': org.get('file_name', ''),
            'reasons': result['reasons'],
        }
        if result['status'] == 'matched':
            row['seller_id'] = result['seller_id']
            row['seller_name'] = index['by_id'][result['seller_id']]['name']
            matched.append(row)
            by_seller[result['seller_id']].append(org)
        elif result['status'] == 'ambiguous':
            ambiguous.append(row)
        else:
            unmatched.append(row)
    plans = []
    conflicts = list(merge_conflicts)
    for seller_id, seller_orgs in by_seller.items():
        seller = index['by_id'][seller_id]
        plan = plan_seller_updates(seller, seller_orgs)
        conflicts.extend(plan['conflicts'])
        note_pack = _note_proposals(seller)
        for item in note_pack['invalid']:
            invalid.append({**item, 'name': seller['name'], 'external_id': '', 'source': 'notes'})
        for item in note_pack['suspicious_whatsapp']:
            suspicious.append({**item, 'name': seller['name'], 'external_id': '', 'source': 'notes'})
        if plan['fills']:
            plans.append({
                'seller_id': seller_id,
                'seller_name': seller['name'],
                'fills': plan['fills'],
                'protected': seller['protected'],
            })
    return {
        'organizations': len(merged),
        'matched': matched,
        'unmatched': unmatched,
        'ambiguous': ambiguous,
        'conflicts': conflicts,
        'invalid_links': invalid,
        'suspicious_whatsapp': suspicious,
        'plans': plans,
    }


def _parse_exported_at(value: str):
    if not value:
        return None
    parsed = parse_datetime(value.replace('Z', '+00:00'))
    if parsed is None:
        try:
            parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
        except ValueError:
            return None
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed


def _time_value(value: str):
    if not value:
        return None
    hour, minute = (int(part) for part in value.split(':'))
    return time(hour, minute)


@transaction.atomic
def apply_plans(plans: list[dict]) -> dict:
    """Write planned fills. Protected seller flags are read back and not saved."""
    enriched = 0
    added = defaultdict(int)
    examples = []
    for plan in plans:
        seller = Seller.objects.select_for_update().get(pk=plan['seller_id'])
        before = (
            seller.is_active,
            seller.is_paused,
            seller.receive_requests,
            seller.whatsapp,
            seller.notes,
        )
        update_fields = []
        wrote = False
        for fill in plan['fills']:
            field = fill['field']
            value = fill['value']
            if field == 'address' and not seller.address:
                seller.address = value
                update_fields.append('address')
                added['address'] += 1
                wrote = True
            elif field == 'latitude' and seller.latitude is None:
                seller.latitude = Decimal(value)
                update_fields.append('latitude')
                wrote = True
            elif field == 'longitude' and seller.longitude is None:
                seller.longitude = Decimal(value)
                update_fields.append('longitude')
                wrote = True
            elif field == 'website' and not seller.website:
                seller.website = value
                update_fields.append('website')
                added['website'] += 1
                wrote = True
            elif field == 'phone2' and not (seller.phone2 or '').strip():
                seller.phone2 = value[:20]
                update_fields.append('phone2')
                added['phone2'] += 1
                wrote = True
            elif field == 'extra_phone':
                _, created = SellerExtraPhone.objects.get_or_create(seller=seller, phone=value)
                if created:
                    added['extra_phone'] += 1
                    wrote = True
            elif field == 'social':
                _, created = SellerSocialLink.objects.get_or_create(
                    seller=seller,
                    kind=value['kind'],
                    url=value['url'],
                )
                if created:
                    added['social'] += 1
                    wrote = True
            elif field == 'schedule' and not seller.work_hours.exists():
                positions = defaultdict(int)
                for interval in value:
                    weekday = int(interval['weekday'])
                    position = positions[weekday]
                    positions[weekday] += 1
                    SellerWorkInterval.objects.create(
                        seller=seller,
                        weekday=weekday,
                        opens=_time_value(interval.get('opens') or ''),
                        closes=_time_value(interval.get('closes') or ''),
                        is_closed=bool(interval.get('is_closed')),
                        position=position,
                    )
                added['schedule'] += 1
                wrote = True
            elif field == 'card':
                exported_at = _parse_exported_at(value.get('exported_at') or '')
                card = SellerExternalCard.objects.filter(
                    source=value['source'],
                    external_id=value['external_id'],
                ).first()
                if card is None:
                    SellerExternalCard.objects.create(
                        seller=seller,
                        source=value['source'],
                        external_id=value['external_id'],
                        card_url=value['card_url'],
                        exported_at=exported_at,
                    )
                    added['card'] += 1
                    wrote = True
                elif card.seller_id == seller.id and not card.card_url:
                    card.card_url = value['card_url']
                    card.exported_at = exported_at or card.exported_at
                    card.save(update_fields=['card_url', 'exported_at'])
                    added['card'] += 1
                    wrote = True
                elif card.seller_id == seller.id and exported_at and (card.exported_at is None or exported_at > card.exported_at):
                    if card.card_url.rstrip('/') == value['card_url'].rstrip('/'):
                        card.exported_at = exported_at
                        card.save(update_fields=['exported_at'])
        if update_fields:
            seller.save(update_fields=sorted(set(update_fields)))
        seller.refresh_from_db()
        after = (
            seller.is_active,
            seller.is_paused,
            seller.receive_requests,
            seller.whatsapp,
            seller.notes,
        )
        if before != after:
            raise RuntimeError(f'Обогащение изменило защищённые поля продавца #{seller.id}')
        if wrote:
            enriched += 1
            if len(examples) < 8:
                examples.append({'id': seller.id, 'name': seller.name, 'fills': [item['field'] for item in plan['fills']]})
    return {'enriched_sellers': enriched, 'added': dict(added), 'examples': examples}


def enrich_files(paths: list[str], *, apply: bool = False) -> dict:
    orgs = []
    file_errors = []
    for path in paths:
        parsed, errors = parse_mapsgrabs_xlsx(path, file_name=path)
        orgs.extend(parsed)
        for error in errors:
            file_errors.append(f'{path}: {error}')
    report = analyze_organizations(orgs)
    report['file_errors'] = file_errors
    report['source_rows'] = len(orgs)
    if apply:
        report['applied'] = apply_plans(report['plans'])
    else:
        report['applied'] = None
    return report
